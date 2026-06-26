"""Accounts & Finance — Chart of Accounts, expenses, vouchers, GST helpers."""

from __future__ import annotations

import json
import sqlite3
import xml.etree.ElementTree as ET
from calendar import monthrange
from datetime import datetime, timedelta
from io import BytesIO
from typing import Any

ACCOUNT_TYPES = ("Asset", "Liability", "Income", "Expense")
EXPENSE_TYPES = ("Expense", "Purchase")
PAYMENT_SOURCES = ("Petty Cash", "Bank", "Cash", "Credit")
PAYMENT_STATUSES = ("Draft", "Unpaid", "Partially Paid", "Paid")


def _normalize_account_type(value: Any) -> str:
    raw = (str(value).strip() if value is not None else "") or "Expense"
    lowered = raw.lower()
    for acct_type in ACCOUNT_TYPES:
        if lowered == acct_type.lower():
            return acct_type
    aliases = {
        "assets": "Asset",
        "liabilities": "Liability",
        "incomes": "Income",
        "expenses": "Expense",
    }
    return aliases.get(lowered, raw if raw in ACCOUNT_TYPES else "Expense")


GST_RATES = (0, 5, 12, 18, 28)
TAX_TYPES = ("CGST_SGST", "IGST")
TDS_SECTIONS = ("194C", "194J", "194H", "194I", "194Q", "Other")
FILING_STATUSES = ("Pending", "Filed", "Overdue", "Not Applicable")
COA_CASH = "A001"
COA_BANK = "A002"
COA_PETTY = "A003"
COA_CREDITORS = "L001"
COA_GST = "L002"
COA_INPUT_GST = "A007"
COA_TDS = "L003"
COA_PF = "L004"
COA_ESI = "L005"

DEFAULT_CHART_OF_ACCOUNTS = [
    ("A001", "Cash", "Asset", "Current Assets", None, 10, 0, 0, 0),
    ("A002", "Bank", "Asset", "Current Assets", None, 20, 0, 0, 0),
    ("A003", "Petty Cash", "Asset", "Current Assets", None, 30, 0, 0, 0),
    ("A004", "Fixed Assets", "Asset", "Fixed Assets", None, 40, 0, 0, 0),
    ("A005", "Security Deposits", "Asset", "Deposits", None, 50, 1, 0, 0),
    ("A006", "Receivables", "Asset", "Receivables", None, 60, 1, 0, 0),
    ("A007", "Input GST", "Asset", "Statutory", None, 65, 0, 0, 0),
    ("L001", "Creditors", "Liability", "Payables", None, 110, 0, 1, 0),
    ("L002", "GST Payable", "Liability", "Statutory", None, 120, 0, 0, 0),
    ("L003", "TDS Payable", "Liability", "Statutory", None, 130, 0, 0, 0),
    ("L004", "PF Payable", "Liability", "Statutory", None, 140, 0, 0, 0),
    ("L005", "ESI Payable", "Liability", "Statutory", None, 150, 0, 0, 0),
    ("I001", "Contract Revenue", "Income", "Operating", None, 210, 1, 0, 0),
    ("I002", "Client Billing", "Income", "Operating", None, 220, 1, 0, 0),
    ("I003", "Other Income", "Income", "Other", None, 230, 0, 0, 0),
    ("E001", "Material Purchase", "Expense", "Direct", None, 310, 1, 1, 1),
    ("E002", "Labour", "Expense", "Direct", None, 320, 1, 0, 0),
    ("E003", "Equipment", "Expense", "Direct", None, 330, 1, 0, 0),
    ("E004", "Fuel", "Expense", "Direct", None, 340, 1, 0, 1),
    ("E005", "Site", "Expense", "Site Overheads", None, 350, 1, 0, 0),
    ("E006", "Office", "Expense", "Admin", None, 360, 0, 0, 1),
    ("E007", "Travel", "Expense", "Admin", None, 370, 0, 0, 1),
    ("E008", "Salary", "Expense", "Personnel", None, 380, 0, 0, 0),
    ("E009", "Professional Charges", "Expense", "Admin", None, 390, 0, 1, 1),
]


def _table_exists(db, table: str) -> bool:
    row = db.execute(
        "SELECT name FROM sqlite_master WHERE type='table' AND name=?",
        (table,),
    ).fetchone()
    return row is not None


def _ensure_column(db, table: str, column: str, col_type: str) -> None:
    if not _table_exists(db, table):
        return
    cols = [row[1] for row in db.execute(f"PRAGMA table_info({table})").fetchall()]
    if column not in cols:
        db.execute(f"ALTER TABLE {table} ADD COLUMN {column} {col_type}")


def _now_ts() -> str:
    return datetime.now().strftime("%Y-%m-%d %H:%M:%S")


def _safe_float(value: Any, default: float = 0.0) -> float:
    try:
        if value is None or value == "":
            return default
        return float(value)
    except (TypeError, ValueError):
        return default


def calc_gst_line(
    quantity: float,
    rate: float,
    gst_percent: float,
    tax_type: str = "CGST_SGST",
) -> dict[str, float]:
    """Compute taxable value and tax split for one expense line."""
    amount = round(quantity * rate, 2)
    taxable = round(amount, 2)
    gst_pct = _safe_float(gst_percent)
    tax_total = round(taxable * gst_pct / 100.0, 2)
    if tax_type == "IGST":
        cgst = 0.0
        sgst = 0.0
        igst = tax_total
    else:
        half = round(tax_total / 2.0, 2)
        cgst = half
        sgst = round(tax_total - half, 2)
        igst = 0.0
    line_total = round(taxable + cgst + sgst + igst, 2)
    return {
        "amount": amount,
        "taxable_value": taxable,
        "cgst": cgst,
        "sgst": sgst,
        "igst": igst,
        "line_total": line_total,
    }


def ensure_accounts_schema(db) -> None:
    """Idempotent accounts schema for Phase 1 structure foundation."""
    db.execute("""
        CREATE TABLE IF NOT EXISTS chart_of_accounts(
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            code TEXT UNIQUE NOT NULL,
            name TEXT NOT NULL,
            account_type TEXT NOT NULL,
            category TEXT,
            subcategory TEXT,
            parent_id INTEGER,
            is_active INTEGER DEFAULT 1,
            sort_order INTEGER DEFAULT 0,
            requires_project INTEGER DEFAULT 0,
            requires_vendor INTEGER DEFAULT 0,
            default_gst_applicable INTEGER DEFAULT 0,
            created_at TEXT,
            FOREIGN KEY(parent_id) REFERENCES chart_of_accounts(id)
        )
    """)
    for column, col_type in (
        ("code", "TEXT"),
        ("name", "TEXT"),
        ("account_type", "TEXT"),
        ("category", "TEXT"),
        ("subcategory", "TEXT"),
        ("parent_id", "INTEGER"),
        ("is_active", "INTEGER DEFAULT 1"),
        ("sort_order", "INTEGER DEFAULT 0"),
        ("requires_project", "INTEGER DEFAULT 0"),
        ("requires_vendor", "INTEGER DEFAULT 0"),
        ("default_gst_applicable", "INTEGER DEFAULT 0"),
        ("created_at", "TEXT"),
    ):
        _ensure_column(db, "chart_of_accounts", column, col_type)

    db.execute("UPDATE chart_of_accounts SET is_active=1 WHERE is_active IS NULL")
    db.execute(
        "UPDATE chart_of_accounts SET account_type='Expense' "
        "WHERE account_type IS NULL OR TRIM(account_type)=''"
    )

    db.execute("""
        CREATE TABLE IF NOT EXISTS account_expenses(
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            entry_date TEXT,
            project_id INTEGER,
            chart_account_id INTEGER NOT NULL,
            vendor_name TEXT,
            invoice_number TEXT,
            invoice_date TEXT,
            expense_type TEXT DEFAULT 'Expense',
            payment_source TEXT DEFAULT 'Bank',
            petty_cash_request_id INTEGER,
            payment_status TEXT DEFAULT 'Draft',
            subtotal REAL DEFAULT 0,
            total_cgst REAL DEFAULT 0,
            total_sgst REAL DEFAULT 0,
            total_igst REAL DEFAULT 0,
            grand_total REAL DEFAULT 0,
            remarks TEXT,
            created_by TEXT,
            created_at TEXT,
            modified_by TEXT,
            modified_at TEXT,
            approval_status TEXT DEFAULT 'Pending Checker',
            FOREIGN KEY(project_id) REFERENCES projects(id),
            FOREIGN KEY(chart_account_id) REFERENCES chart_of_accounts(id),
            FOREIGN KEY(petty_cash_request_id) REFERENCES petty_cash_requests(id)
        )
    """)
    for column, col_type in (
        ("entry_date", "TEXT"),
        ("project_id", "INTEGER"),
        ("chart_account_id", "INTEGER"),
        ("vendor_name", "TEXT"),
        ("invoice_number", "TEXT"),
        ("invoice_date", "TEXT"),
        ("expense_type", "TEXT DEFAULT 'Expense'"),
        ("payment_source", "TEXT DEFAULT 'Bank'"),
        ("petty_cash_request_id", "INTEGER"),
        ("payment_status", "TEXT DEFAULT 'Draft'"),
        ("subtotal", "REAL DEFAULT 0"),
        ("total_cgst", "REAL DEFAULT 0"),
        ("total_sgst", "REAL DEFAULT 0"),
        ("total_igst", "REAL DEFAULT 0"),
        ("grand_total", "REAL DEFAULT 0"),
        ("remarks", "TEXT"),
        ("created_by", "TEXT"),
        ("created_at", "TEXT"),
        ("modified_by", "TEXT"),
        ("modified_at", "TEXT"),
        ("approval_status", "TEXT DEFAULT 'Pending Checker'"),
    ):
        _ensure_column(db, "account_expenses", column, col_type)

    db.execute("""
        CREATE TABLE IF NOT EXISTS account_expense_lines(
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            expense_id INTEGER NOT NULL,
            item_name TEXT,
            quantity REAL DEFAULT 0,
            unit TEXT,
            rate REAL DEFAULT 0,
            amount REAL DEFAULT 0,
            gst_percent REAL DEFAULT 0,
            tax_type TEXT DEFAULT 'CGST_SGST',
            taxable_value REAL DEFAULT 0,
            cgst REAL DEFAULT 0,
            sgst REAL DEFAULT 0,
            igst REAL DEFAULT 0,
            line_total REAL DEFAULT 0,
            sort_order INTEGER DEFAULT 0,
            FOREIGN KEY(expense_id) REFERENCES account_expenses(id) ON DELETE CASCADE
        )
    """)
    for column, col_type in (
        ("expense_id", "INTEGER"),
        ("item_name", "TEXT"),
        ("quantity", "REAL DEFAULT 0"),
        ("unit", "TEXT"),
        ("rate", "REAL DEFAULT 0"),
        ("amount", "REAL DEFAULT 0"),
        ("gst_percent", "REAL DEFAULT 0"),
        ("tax_type", "TEXT DEFAULT 'CGST_SGST'"),
        ("taxable_value", "REAL DEFAULT 0"),
        ("cgst", "REAL DEFAULT 0"),
        ("sgst", "REAL DEFAULT 0"),
        ("igst", "REAL DEFAULT 0"),
        ("line_total", "REAL DEFAULT 0"),
        ("sort_order", "INTEGER DEFAULT 0"),
    ):
        _ensure_column(db, "account_expense_lines", column, col_type)

    db.execute("""
        CREATE TABLE IF NOT EXISTS payment_vouchers(
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            voucher_number TEXT,
            voucher_date TEXT,
            vendor_name TEXT,
            project_id INTEGER,
            chart_account_id INTEGER,
            amount REAL DEFAULT 0,
            payment_mode TEXT,
            bank_name TEXT,
            utr_number TEXT,
            remarks TEXT,
            payment_status TEXT DEFAULT 'Draft',
            petty_cash_request_id INTEGER,
            attachment_filename TEXT,
            created_by TEXT,
            created_at TEXT,
            modified_by TEXT,
            modified_at TEXT,
            approval_status TEXT DEFAULT 'Pending Checker',
            FOREIGN KEY(project_id) REFERENCES projects(id),
            FOREIGN KEY(chart_account_id) REFERENCES chart_of_accounts(id),
            FOREIGN KEY(petty_cash_request_id) REFERENCES petty_cash_requests(id)
        )
    """)
    for column, col_type in (
        ("voucher_number", "TEXT"),
        ("voucher_date", "TEXT"),
        ("vendor_name", "TEXT"),
        ("project_id", "INTEGER"),
        ("chart_account_id", "INTEGER"),
        ("amount", "REAL DEFAULT 0"),
        ("payment_mode", "TEXT"),
        ("bank_name", "TEXT"),
        ("utr_number", "TEXT"),
        ("remarks", "TEXT"),
        ("payment_status", "TEXT DEFAULT 'Draft'"),
        ("petty_cash_request_id", "INTEGER"),
        ("attachment_filename", "TEXT"),
        ("created_by", "TEXT"),
        ("created_at", "TEXT"),
        ("modified_by", "TEXT"),
        ("modified_at", "TEXT"),
        ("approval_status", "TEXT DEFAULT 'Pending Checker'"),
    ):
        _ensure_column(db, "payment_vouchers", column, col_type)

    db.execute("""
        CREATE TABLE IF NOT EXISTS receipt_vouchers(
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            voucher_number TEXT,
            receipt_date TEXT,
            client_name TEXT,
            project_id INTEGER,
            invoice_ref TEXT,
            chart_account_id INTEGER,
            amount REAL DEFAULT 0,
            bank_name TEXT,
            utr_number TEXT,
            remarks TEXT,
            created_by TEXT,
            created_at TEXT,
            modified_by TEXT,
            modified_at TEXT,
            approval_status TEXT DEFAULT 'Pending Checker',
            FOREIGN KEY(project_id) REFERENCES projects(id),
            FOREIGN KEY(chart_account_id) REFERENCES chart_of_accounts(id)
        )
    """)
    for column, col_type in (
        ("voucher_number", "TEXT"),
        ("receipt_date", "TEXT"),
        ("client_name", "TEXT"),
        ("project_id", "INTEGER"),
        ("invoice_ref", "TEXT"),
        ("chart_account_id", "INTEGER"),
        ("amount", "REAL DEFAULT 0"),
        ("bank_name", "TEXT"),
        ("utr_number", "TEXT"),
        ("remarks", "TEXT"),
        ("created_by", "TEXT"),
        ("created_at", "TEXT"),
        ("modified_by", "TEXT"),
        ("modified_at", "TEXT"),
        ("approval_status", "TEXT DEFAULT 'Pending Checker'"),
    ):
        _ensure_column(db, "receipt_vouchers", column, col_type)

    for column, col_type in (
        ("tds_applicable", "INTEGER DEFAULT 0"),
        ("tds_section", "TEXT"),
        ("tds_rate", "REAL DEFAULT 0"),
        ("tds_amount", "REAL DEFAULT 0"),
        ("net_payable", "REAL DEFAULT 0"),
        ("amount_paid", "REAL DEFAULT 0"),
        ("attachment_filename", "TEXT"),
        ("journal_entry_id", "INTEGER"),
    ):
        _ensure_column(db, "account_expenses", column, col_type)

    for column, col_type in (
        ("tds_applicable", "INTEGER DEFAULT 0"),
        ("tds_section", "TEXT"),
        ("tds_rate", "REAL DEFAULT 0"),
        ("tds_amount", "REAL DEFAULT 0"),
        ("net_payable", "REAL DEFAULT 0"),
        ("gross_amount", "REAL DEFAULT 0"),
        ("journal_entry_id", "INTEGER"),
    ):
        _ensure_column(db, "payment_vouchers", column, col_type)

    _ensure_column(db, "receipt_vouchers", "journal_entry_id", "INTEGER")
    _ensure_column(db, "receipt_vouchers", "attachment_filename", "TEXT")
    for column, col_type in (
        ("payment_mode", "TEXT DEFAULT 'Bank'"),
        ("taxable_value", "REAL DEFAULT 0"),
        ("gst_percent", "REAL DEFAULT 0"),
        ("tax_type", "TEXT DEFAULT 'CGST_SGST'"),
        ("total_cgst", "REAL DEFAULT 0"),
        ("total_sgst", "REAL DEFAULT 0"),
        ("total_igst", "REAL DEFAULT 0"),
        ("grand_total", "REAL DEFAULT 0"),
    ):
        _ensure_column(db, "receipt_vouchers", column, col_type)

    if _table_exists(db, "staff"):
        for column, col_type in (
            ("pf_applicable", "INTEGER DEFAULT 0"),
            ("esi_applicable", "INTEGER DEFAULT 0"),
            ("pf_rate", "REAL DEFAULT 12"),
            ("esi_rate", "REAL DEFAULT 0.75"),
        ):
            _ensure_column(db, "staff", column, col_type)

    db.execute("""
        CREATE TABLE IF NOT EXISTS journal_entries(
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            entry_number TEXT,
            entry_date TEXT,
            entry_type TEXT,
            reference_type TEXT,
            reference_id INTEGER,
            narration TEXT,
            created_by TEXT,
            created_at TEXT,
            is_void INTEGER DEFAULT 0
        )
    """)
    db.execute("""
        CREATE TABLE IF NOT EXISTS journal_entry_lines(
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            journal_entry_id INTEGER NOT NULL,
            chart_account_id INTEGER NOT NULL,
            project_id INTEGER,
            party_name TEXT,
            debit REAL DEFAULT 0,
            credit REAL DEFAULT 0,
            line_narration TEXT,
            sort_order INTEGER DEFAULT 0,
            FOREIGN KEY(journal_entry_id) REFERENCES journal_entries(id) ON DELETE CASCADE,
            FOREIGN KEY(chart_account_id) REFERENCES chart_of_accounts(id)
        )
    """)
    for column, col_type in (
        ("entry_number", "TEXT"),
        ("entry_date", "TEXT"),
        ("entry_type", "TEXT"),
        ("reference_type", "TEXT"),
        ("reference_id", "INTEGER"),
        ("narration", "TEXT"),
        ("created_by", "TEXT"),
        ("created_at", "TEXT"),
        ("is_void", "INTEGER DEFAULT 0"),
    ):
        _ensure_column(db, "journal_entries", column, col_type)
    for column, col_type in (
        ("journal_entry_id", "INTEGER"),
        ("chart_account_id", "INTEGER"),
        ("project_id", "INTEGER"),
        ("party_name", "TEXT"),
        ("debit", "REAL DEFAULT 0"),
        ("credit", "REAL DEFAULT 0"),
        ("line_narration", "TEXT"),
        ("sort_order", "INTEGER DEFAULT 0"),
    ):
        _ensure_column(db, "journal_entry_lines", column, col_type)
    db.execute("""
        CREATE TABLE IF NOT EXISTS payment_allocations(
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            payment_voucher_id INTEGER NOT NULL,
            expense_id INTEGER,
            allocated_amount REAL DEFAULT 0,
            FOREIGN KEY(payment_voucher_id) REFERENCES payment_vouchers(id) ON DELETE CASCADE,
            FOREIGN KEY(expense_id) REFERENCES account_expenses(id)
        )
    """)
    db.execute("""
        CREATE TABLE IF NOT EXISTS account_attachments(
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            entity_type TEXT NOT NULL,
            entity_id INTEGER NOT NULL,
            filename TEXT NOT NULL,
            uploaded_by TEXT,
            uploaded_at TEXT
        )
    """)
    db.execute("""
        CREATE TABLE IF NOT EXISTS gst_filing_periods(
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            period_month INTEGER,
            period_year INTEGER,
            period_label TEXT,
            due_date TEXT,
            filing_status TEXT DEFAULT 'Pending',
            filed_date TEXT,
            remarks TEXT
        )
    """)
    db.execute("""
        CREATE TABLE IF NOT EXISTS pf_esi_register(
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            register_type TEXT NOT NULL,
            staff_id INTEGER,
            staff_name TEXT,
            period_month INTEGER,
            period_year INTEGER,
            employee_contribution REAL DEFAULT 0,
            employer_contribution REAL DEFAULT 0,
            total_amount REAL DEFAULT 0,
            filing_status TEXT DEFAULT 'Pending',
            due_date TEXT,
            filed_date TEXT,
            remarks TEXT,
            created_at TEXT
        )
    """)
    db.execute("""
        CREATE TABLE IF NOT EXISTS tds_register(
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            source_type TEXT,
            source_id INTEGER,
            entry_date TEXT,
            vendor_name TEXT,
            tds_section TEXT,
            tds_rate REAL DEFAULT 0,
            tds_amount REAL DEFAULT 0,
            gross_amount REAL DEFAULT 0,
            payment_status TEXT DEFAULT 'Deducted',
            filing_status TEXT DEFAULT 'Pending',
            due_date TEXT,
            filed_date TEXT,
            remarks TEXT
        )
    """)
    for column, col_type in (
        ("source_type", "TEXT"),
        ("source_id", "INTEGER"),
        ("entry_date", "TEXT"),
        ("vendor_name", "TEXT"),
        ("tds_section", "TEXT"),
        ("tds_rate", "REAL DEFAULT 0"),
        ("tds_amount", "REAL DEFAULT 0"),
        ("gross_amount", "REAL DEFAULT 0"),
        ("payment_status", "TEXT DEFAULT 'Deducted'"),
        ("filing_status", "TEXT DEFAULT 'Pending'"),
        ("due_date", "TEXT"),
        ("filed_date", "TEXT"),
        ("remarks", "TEXT"),
    ):
        _ensure_column(db, "tds_register", column, col_type)

    seed_chart_of_accounts(db)
    ensure_extra_chart_heads(db)
    seed_gst_filing_periods(db)


def _chart_row_by_code(db, code: str) -> dict | None:
    """Lookup chart head by code regardless of active flag (for idempotent seeding)."""
    row = db.execute(
        "SELECT * FROM chart_of_accounts WHERE code=?",
        (code,),
    ).fetchone()
    return dict(row) if row else None


def seed_chart_of_accounts(db) -> None:
    """Insert default CoA heads missing by code (idempotent for empty or partial tables)."""
    if not _table_exists(db, "chart_of_accounts"):
        return
    now = _now_ts()
    for code, name, acct_type, category, subcategory, sort_order, req_proj, req_vendor, gst in DEFAULT_CHART_OF_ACCOUNTS:
        existing = _chart_row_by_code(db, code)
        if existing:
            if not int(existing.get("is_active") or 0):
                db.execute(
                    "UPDATE chart_of_accounts SET is_active=1 WHERE id=?",
                    (existing["id"],),
                )
            continue
        db.execute(
            "INSERT INTO chart_of_accounts("
            "code, name, account_type, category, subcategory, parent_id, is_active, "
            "sort_order, requires_project, requires_vendor, default_gst_applicable, created_at"
            ") VALUES(?,?,?,?,?,?,1,?,?,?,?,?)",
            (code, name, acct_type, category, subcategory, None, sort_order, req_proj, req_vendor, gst, now),
        )


def ensure_extra_chart_heads(db) -> None:
    """Insert statutory heads added after initial CoA seed (safe for existing DBs)."""
    extras = [
        ("A007", "Input GST", "Asset", "Statutory", None, 65, 0, 0, 0),
    ]
    now = _now_ts()
    for code, name, acct_type, category, subcategory, sort_order, req_proj, req_vendor, gst in extras:
        existing = _chart_row_by_code(db, code)
        if existing:
            if not int(existing.get("is_active") or 0):
                db.execute(
                    "UPDATE chart_of_accounts SET is_active=1 WHERE id=?",
                    (existing["id"],),
                )
            continue
        db.execute(
            "INSERT INTO chart_of_accounts("
            "code, name, account_type, category, subcategory, parent_id, is_active, "
            "sort_order, requires_project, requires_vendor, default_gst_applicable, created_at"
            ") VALUES(?,?,?,?,?,?,1,?,?,?,?,?)",
            (code, name, acct_type, category, subcategory, None, sort_order, req_proj, req_vendor, gst, now),
        )


def get_chart_account_by_code(db, code: str) -> dict | None:
    row = db.execute(
        "SELECT * FROM chart_of_accounts WHERE code=? AND is_active=1",
        (code,),
    ).fetchone()
    return dict(row) if row else None


def _calc_tds(amount: float, applicable: bool, rate: float) -> dict[str, float]:
    gross = round(_safe_float(amount), 2)
    if not applicable or rate <= 0:
        return {"gross_amount": gross, "tds_amount": 0.0, "net_payable": gross}
    tds = round(gross * rate / 100.0, 2)
    return {"gross_amount": gross, "tds_amount": tds, "net_payable": round(gross - tds, 2)}


def seed_gst_filing_periods(db) -> None:
    if not _table_exists(db, "gst_filing_periods"):
        return
    count = db.execute("SELECT COUNT(*) AS c FROM gst_filing_periods").fetchone()["c"]
    if int(count or 0) > 0:
        return
    today = datetime.now()
    for offset in range(0, 6):
        month = today.month - offset
        year = today.year
        while month <= 0:
            month += 12
            year -= 1
        due_month = month + 1
        due_year = year
        if due_month > 12:
            due_month = 1
            due_year += 1
        due_day = min(20, monthrange(due_year, due_month)[1])
        due_date = f"{due_year:04d}-{due_month:02d}-{due_day:02d}"
        label = datetime(year, month, 1).strftime("%b %Y")
        db.execute(
            "INSERT INTO gst_filing_periods(period_month, period_year, period_label, due_date, filing_status) "
            "VALUES(?,?,?,?, 'Pending')",
            (month, year, label, due_date),
        )


def generate_voucher_number(db, prefix: str) -> str:
    today = datetime.now().strftime("%Y%m%d")
    key = f"{prefix}-{today}"
    row = db.execute(
        "SELECT COUNT(*) AS c FROM payment_vouchers WHERE voucher_number LIKE ?",
        (f"{key}-%",),
    ).fetchone()
    seq = int(row["c"] or 0) + 1
    if prefix == "RV":
        row = db.execute(
            "SELECT COUNT(*) AS c FROM receipt_vouchers WHERE voucher_number LIKE ?",
            (f"{key}-%",),
        ).fetchone()
        seq = max(seq, int(row["c"] or 0) + 1)
    return f"{key}-{seq:03d}"


def list_chart_of_accounts(db, active_only: bool = True) -> list[dict]:
    where = "WHERE COALESCE(is_active, 1)=1" if active_only else ""
    rows = db.execute(
        f"SELECT * FROM chart_of_accounts {where} ORDER BY account_type, sort_order, code"
    ).fetchall()
    result: list[dict] = []
    for row in rows:
        item = dict(row)
        item["account_type"] = _normalize_account_type(item.get("account_type"))
        result.append(item)
    return result


def list_expense_chart_heads(db) -> list[dict]:
    return [h for h in list_chart_of_accounts(db) if h.get("account_type") == "Expense"]


def list_income_chart_heads(db) -> list[dict]:
    return [h for h in list_chart_of_accounts(db) if h.get("account_type") == "Income"]


def chart_accounts_grouped(db) -> dict[str, list[dict]]:
    grouped: dict[str, list[dict]] = {t: [] for t in ACCOUNT_TYPES}
    for row in list_chart_of_accounts(db, active_only=False):
        acct_type = _normalize_account_type(row.get("account_type"))
        grouped.setdefault(acct_type, []).append(row)
    return grouped


def get_chart_account(db, account_id: int) -> dict | None:
    row = db.execute("SELECT * FROM chart_of_accounts WHERE id=?", (account_id,)).fetchone()
    return dict(row) if row else None


def save_chart_account(db, data: dict, account_id: int | None = None) -> int:
    now = _now_ts()
    code = (data.get("code") or "").strip()
    name = (data.get("name") or "").strip()
    if not code or not name:
        raise ValueError("Code and name are required.")
    acct_type = (data.get("account_type") or "Expense").strip()
    if acct_type not in ACCOUNT_TYPES:
        raise ValueError("Invalid account type.")
    values = (
        code,
        name,
        acct_type,
        (data.get("category") or "").strip(),
        (data.get("subcategory") or "").strip(),
        data.get("parent_id") or None,
        1 if str(data.get("is_active", "1")) in ("1", "true", "True", "on") else 0,
        int(data.get("sort_order") or 0),
        1 if str(data.get("requires_project", "")) in ("1", "true", "True", "on") else 0,
        1 if str(data.get("requires_vendor", "")) in ("1", "true", "True", "on") else 0,
        1 if str(data.get("default_gst_applicable", "")) in ("1", "true", "True", "on") else 0,
    )
    if account_id:
        db.execute(
            "UPDATE chart_of_accounts SET code=?, name=?, account_type=?, category=?, subcategory=?, "
            "parent_id=?, is_active=?, sort_order=?, requires_project=?, requires_vendor=?, "
            "default_gst_applicable=? WHERE id=?",
            values + (account_id,),
        )
        return account_id
    db.execute(
        "INSERT INTO chart_of_accounts("
        "code, name, account_type, category, subcategory, parent_id, is_active, sort_order, "
        "requires_project, requires_vendor, default_gst_applicable, created_at"
        ") VALUES(?,?,?,?,?,?,?,?,?,?,?,?)",
        values + (now,),
    )
    return db.execute("SELECT last_insert_rowid()").fetchone()[0]


def parse_expense_lines_from_form(form) -> tuple[list[dict], str | None]:
    names = form.getlist("item_name[]")
    quantities = form.getlist("quantity[]")
    units = form.getlist("unit[]")
    rates = form.getlist("rate[]")
    gst_percents = form.getlist("gst_percent[]")
    tax_types = form.getlist("tax_type[]")
    row_count = max(len(names), len(quantities), len(rates))
    lines: list[dict] = []
    for idx in range(row_count):
        item_name = (names[idx] if idx < len(names) else "").strip()
        qty_raw = quantities[idx] if idx < len(quantities) else ""
        rate_raw = rates[idx] if idx < len(rates) else ""
        unit = (units[idx] if idx < len(units) else "").strip()
        gst_raw = gst_percents[idx] if idx < len(gst_percents) else "0"
        tax_type = (tax_types[idx] if idx < len(tax_types) else "CGST_SGST").strip()
        if tax_type not in TAX_TYPES:
            tax_type = "CGST_SGST"
        if not item_name and not str(qty_raw).strip() and not str(rate_raw).strip():
            continue
        if not item_name:
            return [], "Each line must have an item name."
        qty = _safe_float(qty_raw)
        rate = _safe_float(rate_raw)
        gst_percent = _safe_float(gst_raw)
        calc = calc_gst_line(qty, rate, gst_percent, tax_type)
        lines.append({
            "item_name": item_name,
            "quantity": qty,
            "unit": unit,
            "rate": rate,
            "amount": calc["amount"],
            "gst_percent": gst_percent,
            "tax_type": tax_type,
            "taxable_value": calc["taxable_value"],
            "cgst": calc["cgst"],
            "sgst": calc["sgst"],
            "igst": calc["igst"],
            "line_total": calc["line_total"],
            "sort_order": len(lines),
        })
    if not lines:
        return [], "Add at least one expense line item."
    return lines, None


def _persist_expense_lines(db, expense_id: int, lines: list[dict]) -> None:
    db.execute("DELETE FROM account_expense_lines WHERE expense_id=?", (expense_id,))
    for line in lines:
        db.execute(
            "INSERT INTO account_expense_lines("
            "expense_id, item_name, quantity, unit, rate, amount, gst_percent, tax_type, "
            "taxable_value, cgst, sgst, igst, line_total, sort_order"
            ") VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?)",
            (
                expense_id,
                line["item_name"],
                line["quantity"],
                line.get("unit") or "",
                line["rate"],
                line["amount"],
                line["gst_percent"],
                line["tax_type"],
                line["taxable_value"],
                line["cgst"],
                line["sgst"],
                line["igst"],
                line["line_total"],
                line.get("sort_order", 0),
            ),
        )


def _totals_from_lines(lines: list[dict]) -> dict[str, float]:
    subtotal = round(sum(l["taxable_value"] for l in lines), 2)
    total_cgst = round(sum(l["cgst"] for l in lines), 2)
    total_sgst = round(sum(l["sgst"] for l in lines), 2)
    total_igst = round(sum(l["igst"] for l in lines), 2)
    grand_total = round(subtotal + total_cgst + total_sgst + total_igst, 2)
    return {
        "subtotal": subtotal,
        "total_cgst": total_cgst,
        "total_sgst": total_sgst,
        "total_igst": total_igst,
        "grand_total": grand_total,
    }


def save_account_expense(db, form, username: str, expense_id: int | None = None) -> int:
    lines, err = parse_expense_lines_from_form(form)
    if err:
        raise ValueError(err)
    totals = _totals_from_lines(lines)
    now = _now_ts()
    chart_account_id = form.get("chart_account_id")
    if not chart_account_id:
        raise ValueError("Select a head of account.")
    head = get_chart_account(db, int(chart_account_id))
    if not head:
        raise ValueError("Invalid head of account.")
    if head.get("requires_project") and not form.get("project_id"):
        raise ValueError("Project is required for this account head.")
    if head.get("requires_vendor") and not (form.get("vendor_name") or "").strip():
        raise ValueError("Vendor is required for this account head.")
    payment_source = (form.get("payment_source") or "Bank").strip()
    petty_cash_request_id = form.get("petty_cash_request_id") or None
    if payment_source == "Petty Cash" and not petty_cash_request_id:
        raise ValueError("Select a petty cash request when payment source is Petty Cash.")
    tds_applicable = str(form.get("tds_applicable", "")) in ("1", "true", "True", "on")
    tds_rate = _safe_float(form.get("tds_rate"))
    tds_calc = _calc_tds(totals["grand_total"], tds_applicable, tds_rate)
    tds_section = (form.get("tds_section") or "").strip() if tds_applicable else ""
    values = (
        form.get("entry_date", "").strip(),
        form.get("project_id") or None,
        int(chart_account_id),
        (form.get("vendor_name") or "").strip(),
        (form.get("invoice_number") or "").strip(),
        (form.get("invoice_date") or "").strip(),
        (form.get("expense_type") or "Expense").strip(),
        payment_source,
        petty_cash_request_id,
        (form.get("payment_status") or "Draft").strip(),
        totals["subtotal"],
        totals["total_cgst"],
        totals["total_sgst"],
        totals["total_igst"],
        totals["grand_total"],
        (form.get("remarks") or "").strip(),
        1 if tds_applicable else 0,
        tds_section,
        tds_rate if tds_applicable else 0,
        tds_calc["tds_amount"],
        tds_calc["net_payable"],
        _safe_float(form.get("amount_paid")),
    )
    if expense_id:
        db.execute(
            "UPDATE account_expenses SET entry_date=?, project_id=?, chart_account_id=?, vendor_name=?, "
            "invoice_number=?, invoice_date=?, expense_type=?, payment_source=?, petty_cash_request_id=?, "
            "payment_status=?, subtotal=?, total_cgst=?, total_sgst=?, total_igst=?, grand_total=?, "
            "remarks=?, tds_applicable=?, tds_section=?, tds_rate=?, tds_amount=?, net_payable=?, "
            "amount_paid=?, modified_by=?, modified_at=? WHERE id=?",
            values + (username, now, expense_id),
        )
        _persist_expense_lines(db, expense_id, lines)
        _sync_tds_register_for_expense(db, expense_id)
        return expense_id
    db.execute(
        "INSERT INTO account_expenses("
        "entry_date, project_id, chart_account_id, vendor_name, invoice_number, invoice_date, "
        "expense_type, payment_source, petty_cash_request_id, payment_status, subtotal, total_cgst, "
        "total_sgst, total_igst, grand_total, remarks, tds_applicable, tds_section, tds_rate, "
        "tds_amount, net_payable, amount_paid, created_by, created_at, modified_by, modified_at, "
        "approval_status"
        ") VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,'Pending Checker')",
        values + (username, now, username, now),
    )
    new_id = db.execute("SELECT last_insert_rowid()").fetchone()[0]
    _persist_expense_lines(db, new_id, lines)
    _sync_tds_register_for_expense(db, new_id)
    return new_id


def load_account_expense(db, expense_id: int) -> dict | None:
    row = db.execute(
        "SELECT e.*, c.code AS account_code, c.name AS account_name, p.project_name "
        "FROM account_expenses e "
        "LEFT JOIN chart_of_accounts c ON e.chart_account_id = c.id "
        "LEFT JOIN projects p ON e.project_id = p.id "
        "WHERE e.id=?",
        (expense_id,),
    ).fetchone()
    if not row:
        return None
    data = dict(row)
    line_rows = db.execute(
        "SELECT * FROM account_expense_lines WHERE expense_id=? ORDER BY sort_order, id",
        (expense_id,),
    ).fetchall()
    data["lines"] = [dict(l) for l in line_rows]
    return data


def list_account_expenses(db) -> list[dict]:
    rows = db.execute(
        "SELECT e.*, c.code AS account_code, c.name AS account_name, p.project_name "
        "FROM account_expenses e "
        "LEFT JOIN chart_of_accounts c ON e.chart_account_id = c.id "
        "LEFT JOIN projects p ON e.project_id = p.id "
        "ORDER BY e.entry_date DESC, e.id DESC"
    ).fetchall()
    return [dict(r) for r in rows]


def save_payment_voucher(db, form, username: str, voucher_id: int | None = None) -> int:
    now = _now_ts()
    amount = _safe_float(form.get("amount"))
    if amount <= 0:
        raise ValueError("Amount must be greater than zero.")
    payment_mode = (form.get("payment_mode") or "Bank").strip()
    petty_cash_request_id = form.get("petty_cash_request_id") or None
    if payment_mode == "Petty Cash" and not petty_cash_request_id:
        raise ValueError("Select a petty cash request for petty cash payments.")
    tds_applicable = str(form.get("tds_applicable", "")) in ("1", "true", "True", "on")
    tds_rate = _safe_float(form.get("tds_rate"))
    tds_calc = _calc_tds(amount, tds_applicable, tds_rate)
    tds_section = (form.get("tds_section") or "").strip() if tds_applicable else ""
    values = (
        (form.get("voucher_date") or "").strip(),
        (form.get("vendor_name") or "").strip(),
        form.get("project_id") or None,
        form.get("chart_account_id") or None,
        amount,
        payment_mode,
        (form.get("bank_name") or "").strip(),
        (form.get("utr_number") or "").strip(),
        (form.get("remarks") or "").strip(),
        (form.get("payment_status") or "Draft").strip(),
        petty_cash_request_id,
        1 if tds_applicable else 0,
        tds_section,
        tds_rate if tds_applicable else 0,
        tds_calc["tds_amount"],
        tds_calc["net_payable"],
        tds_calc["gross_amount"],
    )
    if voucher_id:
        db.execute(
            "UPDATE payment_vouchers SET voucher_date=?, vendor_name=?, project_id=?, chart_account_id=?, "
            "amount=?, payment_mode=?, bank_name=?, utr_number=?, remarks=?, payment_status=?, "
            "petty_cash_request_id=?, tds_applicable=?, tds_section=?, tds_rate=?, tds_amount=?, "
            "net_payable=?, gross_amount=?, modified_by=?, modified_at=? WHERE id=?",
            values + (username, now, voucher_id),
        )
        _save_payment_allocations(db, voucher_id, form)
        _sync_tds_register_for_payment(db, voucher_id)
        return voucher_id
    voucher_number = generate_voucher_number(db, "PV")
    db.execute(
        "INSERT INTO payment_vouchers("
        "voucher_number, voucher_date, vendor_name, project_id, chart_account_id, amount, payment_mode, "
        "bank_name, utr_number, remarks, payment_status, petty_cash_request_id, tds_applicable, tds_section, "
        "tds_rate, tds_amount, net_payable, gross_amount, created_by, created_at, modified_by, modified_at, "
        "approval_status"
        ") VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,'Pending Checker')",
        (voucher_number,) + values + (username, now, username, now),
    )
    new_id = db.execute("SELECT last_insert_rowid()").fetchone()[0]
    _save_payment_allocations(db, new_id, form)
    _sync_tds_register_for_payment(db, new_id)
    return new_id


def load_payment_voucher(db, voucher_id: int) -> dict | None:
    row = db.execute(
        "SELECT v.*, c.code AS account_code, c.name AS account_name, p.project_name, "
        "pc.request_number AS petty_cash_number "
        "FROM payment_vouchers v "
        "LEFT JOIN chart_of_accounts c ON v.chart_account_id = c.id "
        "LEFT JOIN projects p ON v.project_id = p.id "
        "LEFT JOIN petty_cash_requests pc ON v.petty_cash_request_id = pc.id "
        "WHERE v.id=?",
        (voucher_id,),
    ).fetchone()
    return dict(row) if row else None


def list_payment_vouchers(db) -> list[dict]:
    rows = db.execute(
        "SELECT v.*, c.code AS account_code, c.name AS account_name, p.project_name "
        "FROM payment_vouchers v "
        "LEFT JOIN chart_of_accounts c ON v.chart_account_id = c.id "
        "LEFT JOIN projects p ON v.project_id = p.id "
        "ORDER BY v.voucher_date DESC, v.id DESC"
    ).fetchall()
    return [dict(r) for r in rows]


def save_receipt_voucher(db, form, username: str, voucher_id: int | None = None) -> int:
    now = _now_ts()
    amounts = _calc_receipt_gst_amounts(form)
    if amounts["grand_total"] <= 0:
        raise ValueError("Amount must be greater than zero.")
    payment_mode = (form.get("payment_mode") or "Bank").strip()
    values = (
        (form.get("receipt_date") or "").strip(),
        (form.get("client_name") or "").strip(),
        form.get("project_id") or None,
        (form.get("invoice_ref") or "").strip(),
        form.get("chart_account_id") or None,
        amounts["amount"],
        payment_mode,
        amounts["taxable_value"],
        amounts["gst_percent"],
        amounts["tax_type"],
        amounts["total_cgst"],
        amounts["total_sgst"],
        amounts["total_igst"],
        amounts["grand_total"],
        (form.get("bank_name") or "").strip(),
        (form.get("utr_number") or "").strip(),
        (form.get("remarks") or "").strip(),
    )
    if voucher_id:
        db.execute(
            "UPDATE receipt_vouchers SET receipt_date=?, client_name=?, project_id=?, invoice_ref=?, "
            "chart_account_id=?, amount=?, payment_mode=?, taxable_value=?, gst_percent=?, tax_type=?, "
            "total_cgst=?, total_sgst=?, total_igst=?, grand_total=?, bank_name=?, utr_number=?, remarks=?, "
            "modified_by=?, modified_at=? WHERE id=?",
            values + (username, now, voucher_id),
        )
        return voucher_id
    voucher_number = generate_voucher_number(db, "RV")
    db.execute(
        "INSERT INTO receipt_vouchers("
        "voucher_number, receipt_date, client_name, project_id, invoice_ref, chart_account_id, amount, "
        "payment_mode, taxable_value, gst_percent, tax_type, total_cgst, total_sgst, total_igst, grand_total, "
        "bank_name, utr_number, remarks, created_by, created_at, modified_by, modified_at, approval_status"
        ") VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,'Pending Checker')",
        (voucher_number,) + values + (username, now, username, now),
    )
    return db.execute("SELECT last_insert_rowid()").fetchone()[0]


def _calc_receipt_gst_amounts(form) -> dict[str, float | str]:
    taxable = _safe_float(form.get("taxable_value"))
    if taxable <= 0:
        taxable = _safe_float(form.get("amount"))
    gst_percent = _safe_float(form.get("gst_percent"))
    tax_type = (form.get("tax_type") or "CGST_SGST").strip()
    if gst_percent > 0 and taxable > 0:
        gst = calc_gst_line(1, taxable, gst_percent, tax_type)
        return {
            "taxable_value": taxable,
            "gst_percent": gst_percent,
            "tax_type": tax_type,
            "total_cgst": gst["cgst"],
            "total_sgst": gst["sgst"],
            "total_igst": gst["igst"],
            "grand_total": gst["line_total"],
            "amount": gst["line_total"],
        }
    amt = round(taxable, 2)
    return {
        "taxable_value": amt,
        "gst_percent": 0.0,
        "tax_type": tax_type,
        "total_cgst": 0.0,
        "total_sgst": 0.0,
        "total_igst": 0.0,
        "grand_total": amt,
        "amount": amt,
    }


def load_receipt_voucher(db, voucher_id: int) -> dict | None:
    row = db.execute(
        "SELECT v.*, c.code AS account_code, c.name AS account_name, p.project_name "
        "FROM receipt_vouchers v "
        "LEFT JOIN chart_of_accounts c ON v.chart_account_id = c.id "
        "LEFT JOIN projects p ON v.project_id = p.id "
        "WHERE v.id=?",
        (voucher_id,),
    ).fetchone()
    return dict(row) if row else None


def list_receipt_vouchers(db) -> list[dict]:
    rows = db.execute(
        "SELECT v.*, c.code AS account_code, c.name AS account_name, p.project_name "
        "FROM receipt_vouchers v "
        "LEFT JOIN chart_of_accounts c ON v.chart_account_id = c.id "
        "LEFT JOIN projects p ON v.project_id = p.id "
        "ORDER BY v.receipt_date DESC, v.id DESC"
    ).fetchall()
    return [dict(r) for r in rows]


def list_settled_petty_cash(db) -> list[dict]:
    if not _table_exists(db, "petty_cash_requests"):
        return []
    rows = db.execute(
        "SELECT id, request_number, request_date, staff_name, project_id, purpose, "
        "transferred_amount, expenses_total, status "
        "FROM petty_cash_requests "
        "WHERE status IN ('Settled', 'Closed', 'Amount Received', 'Settlement Pending') "
        "ORDER BY id DESC"
    ).fetchall()
    return [dict(r) for r in rows]


def get_gst_purchase_register(db) -> list[dict]:
    rows = db.execute(
        "SELECT e.entry_date, e.invoice_number, e.invoice_date, e.vendor_name, "
        "c.code AS account_code, c.name AS account_name, p.project_name, "
        "l.item_name, l.taxable_value, l.gst_percent, l.cgst, l.sgst, l.igst, l.line_total, "
        "e.approval_status, e.id AS expense_id "
        "FROM account_expense_lines l "
        "JOIN account_expenses e ON l.expense_id = e.id "
        "LEFT JOIN chart_of_accounts c ON e.chart_account_id = c.id "
        "LEFT JOIN projects p ON e.project_id = p.id "
        "WHERE e.expense_type IN ('Expense', 'Purchase') AND e.approval_status='Approved' "
        "ORDER BY e.entry_date DESC, e.id DESC, l.sort_order"
    ).fetchall()
    return [dict(r) for r in rows]


def get_gst_sales_register(db) -> list[dict]:
    rows = db.execute(
        "SELECT v.receipt_date AS entry_date, v.invoice_ref AS invoice_number, v.receipt_date AS invoice_date, "
        "v.client_name AS vendor_name, c.code AS account_code, c.name AS account_name, p.project_name, "
        "COALESCE(v.taxable_value, v.amount) AS taxable_value, COALESCE(v.gst_percent, 0) AS gst_percent, "
        "COALESCE(v.total_cgst, 0) AS cgst, COALESCE(v.total_sgst, 0) AS sgst, COALESCE(v.total_igst, 0) AS igst, "
        "COALESCE(v.grand_total, v.amount) AS line_total, "
        "v.approval_status, v.id AS expense_id "
        "FROM receipt_vouchers v "
        "LEFT JOIN chart_of_accounts c ON v.chart_account_id = c.id "
        "LEFT JOIN projects p ON v.project_id = p.id "
        "WHERE v.approval_status='Approved' "
        "ORDER BY v.receipt_date DESC, v.id DESC"
    ).fetchall()
    return [dict(r) for r in rows]


# --- Phase 2: Journal, ledgers, reports, compliance ---


def _payment_source_account_id(db, payment_source: str) -> int | None:
    mapping = {
        "Cash": COA_CASH,
        "Bank": COA_BANK,
        "Petty Cash": COA_PETTY,
        "Credit": COA_CREDITORS,
    }
    code = mapping.get(payment_source or "", COA_BANK)
    head = get_chart_account_by_code(db, code)
    return head["id"] if head else None


def _payment_mode_account_id(db, payment_mode: str) -> int | None:
    mapping = {
        "Cash": COA_CASH,
        "Bank": COA_BANK,
        "UPI": COA_BANK,
        "Cheque": COA_BANK,
        "Petty Cash": COA_PETTY,
    }
    code = mapping.get(payment_mode or "", COA_BANK)
    head = get_chart_account_by_code(db, code)
    return head["id"] if head else None


def _generate_journal_number(db) -> str:
    today = datetime.now().strftime("%Y%m%d")
    row = db.execute(
        "SELECT COUNT(*) AS c FROM journal_entries WHERE entry_number LIKE ?",
        (f"JE-{today}-%",),
    ).fetchone()
    seq = int(row["c"] or 0) + 1
    return f"JE-{today}-{seq:03d}"


def _insert_journal(
    db,
    entry_date: str,
    entry_type: str,
    reference_type: str,
    reference_id: int,
    narration: str,
    lines: list[dict],
    username: str,
) -> int:
    now = _now_ts()
    entry_number = _generate_journal_number(db)
    db.execute(
        "INSERT INTO journal_entries(entry_number, entry_date, entry_type, reference_type, "
        "reference_id, narration, created_by, created_at) VALUES(?,?,?,?,?,?,?,?)",
        (entry_number, entry_date, entry_type, reference_type, reference_id, narration, username, now),
    )
    je_id = db.execute("SELECT last_insert_rowid()").fetchone()[0]
    for idx, line in enumerate(lines):
        db.execute(
            "INSERT INTO journal_entry_lines(journal_entry_id, chart_account_id, project_id, party_name, "
            "debit, credit, line_narration, sort_order) VALUES(?,?,?,?,?,?,?,?)",
            (
                je_id,
                line["chart_account_id"],
                line.get("project_id"),
                line.get("party_name"),
                round(_safe_float(line.get("debit")), 2),
                round(_safe_float(line.get("credit")), 2),
                line.get("line_narration") or "",
                idx,
            ),
        )
    return je_id


def void_journal_for_reference(db, reference_type: str, reference_id: int) -> None:
    """Void all active journal entries for a voucher (handles duplicate-post edge cases)."""
    rows = db.execute(
        "SELECT id FROM journal_entries WHERE reference_type=? AND reference_id=? AND is_void=0",
        (reference_type, reference_id),
    ).fetchall()
    if not rows:
        return
    for row in rows:
        db.execute("UPDATE journal_entries SET is_void=1 WHERE id=?", (row["id"],))
    if reference_type == "account_expenses":
        db.execute(
            "UPDATE account_expenses SET journal_entry_id=NULL WHERE id=?",
            (reference_id,),
        )
    elif reference_type == "payment_vouchers":
        db.execute(
            "UPDATE payment_vouchers SET journal_entry_id=NULL WHERE id=?",
            (reference_id,),
        )
    elif reference_type == "receipt_vouchers":
        db.execute(
            "UPDATE receipt_vouchers SET journal_entry_id=NULL WHERE id=?",
            (reference_id,),
        )


def post_journal_on_approval(db, record_table: str, record_id: int, username: str = "system") -> int | None:
    existing = db.execute(
        "SELECT id FROM journal_entries WHERE reference_type=? AND reference_id=? AND is_void=0",
        (record_table, record_id),
    ).fetchone()
    if existing:
        return existing["id"]
    if record_table == "account_expenses":
        return _post_expense_journal(db, record_id, username)
    if record_table == "payment_vouchers":
        return _post_payment_journal(db, record_id, username)
    if record_table == "receipt_vouchers":
        return _post_receipt_journal(db, record_id, username)
    return None


def _post_expense_journal(db, expense_id: int, username: str) -> int | None:
    exp = load_account_expense(db, expense_id)
    if not exp or exp.get("approval_status") != "Approved":
        return None
    expense_acct = exp["chart_account_id"]
    grand = _safe_float(exp.get("grand_total"))
    subtotal = _safe_float(exp.get("subtotal"))
    gst_total = round(
        _safe_float(exp.get("total_cgst"))
        + _safe_float(exp.get("total_sgst"))
        + _safe_float(exp.get("total_igst")),
        2,
    )
    credit_acct = _payment_source_account_id(db, exp.get("payment_source") or "Bank")
    if exp.get("payment_status") in ("Unpaid", "Partially Paid", "Draft"):
        creditors = get_chart_account_by_code(db, COA_CREDITORS)
        credit_acct = creditors["id"] if creditors else credit_acct
    lines: list[dict] = [{
        "chart_account_id": expense_acct,
        "project_id": exp.get("project_id"),
        "party_name": exp.get("vendor_name"),
        "debit": subtotal,
        "credit": 0,
        "line_narration": f"Expense #{expense_id}",
    }]
    if gst_total > 0:
        gst_head = get_chart_account_by_code(db, COA_INPUT_GST) or get_chart_account_by_code(db, COA_GST)
        if gst_head:
            lines.append({
                "chart_account_id": gst_head["id"],
                "project_id": exp.get("project_id"),
                "party_name": exp.get("vendor_name"),
                "debit": gst_total,
                "credit": 0,
                "line_narration": "GST on purchase",
            })
    if credit_acct:
        lines.append({
            "chart_account_id": credit_acct,
            "project_id": exp.get("project_id"),
            "party_name": exp.get("vendor_name"),
            "debit": 0,
            "credit": grand,
            "line_narration": exp.get("payment_source") or "Bank",
        })
    je_id = _insert_journal(
        db,
        exp.get("entry_date") or _now_ts()[:10],
        "Expense",
        "account_expenses",
        expense_id,
        f"Expense #{expense_id} — {exp.get('vendor_name') or ''}",
        lines,
        username,
    )
    db.execute("UPDATE account_expenses SET journal_entry_id=? WHERE id=?", (je_id, expense_id))
    if exp.get("tds_applicable") and _safe_float(exp.get("tds_amount")) > 0:
        _sync_tds_register_for_expense(db, expense_id)
    return je_id


def _post_payment_journal(db, voucher_id: int, username: str) -> int | None:
    pv = load_payment_voucher(db, voucher_id)
    if not pv or pv.get("approval_status") != "Approved":
        return None
    gross = _safe_float(pv.get("gross_amount") or pv.get("amount"))
    net = _safe_float(pv.get("net_payable") or pv.get("amount"))
    tds_amt = _safe_float(pv.get("tds_amount"))
    debit_acct = pv.get("chart_account_id")
    if not debit_acct:
        creditors = get_chart_account_by_code(db, COA_CREDITORS)
        debit_acct = creditors["id"] if creditors else None
    credit_acct = _payment_mode_account_id(db, pv.get("payment_mode") or "Bank")
    if not debit_acct or not credit_acct:
        return None
    lines: list[dict] = [{
        "chart_account_id": debit_acct,
        "project_id": pv.get("project_id"),
        "party_name": pv.get("vendor_name"),
        "debit": gross,
        "credit": 0,
        "line_narration": pv.get("voucher_number") or f"PV-{voucher_id}",
    }]
    if tds_amt > 0:
        tds_head = get_chart_account_by_code(db, COA_TDS)
        if tds_head:
            lines.append({
                "chart_account_id": tds_head["id"],
                "project_id": pv.get("project_id"),
                "party_name": pv.get("vendor_name"),
                "debit": 0,
                "credit": tds_amt,
                "line_narration": "TDS deducted",
            })
    lines.append({
        "chart_account_id": credit_acct,
        "project_id": pv.get("project_id"),
        "party_name": pv.get("vendor_name"),
        "debit": 0,
        "credit": net,
        "line_narration": pv.get("payment_mode") or "Bank",
    })
    je_id = _insert_journal(
        db,
        pv.get("voucher_date") or _now_ts()[:10],
        "Payment",
        "payment_vouchers",
        voucher_id,
        f"Payment {pv.get('voucher_number') or voucher_id}",
        lines,
        username,
    )
    db.execute("UPDATE payment_vouchers SET journal_entry_id=? WHERE id=?", (je_id, voucher_id))
    _apply_payment_allocations(db, voucher_id)
    _sync_tds_register_for_payment(db, voucher_id)
    return je_id


def _post_receipt_journal(db, voucher_id: int, username: str) -> int | None:
    rv = load_receipt_voucher(db, voucher_id)
    if not rv or rv.get("approval_status") != "Approved":
        return None
    grand = _safe_float(rv.get("grand_total") or rv.get("amount"))
    taxable = _safe_float(rv.get("taxable_value") or grand)
    gst_total = round(
        _safe_float(rv.get("total_cgst"))
        + _safe_float(rv.get("total_sgst"))
        + _safe_float(rv.get("total_igst")),
        2,
    )
    debit_acct = _payment_mode_account_id(db, rv.get("payment_mode") or "Bank")
    credit_acct = rv.get("chart_account_id")
    if not debit_acct or not credit_acct:
        return None
    lines: list[dict] = [
        {
            "chart_account_id": debit_acct,
            "project_id": rv.get("project_id"),
            "party_name": rv.get("client_name"),
            "debit": grand,
            "credit": 0,
            "line_narration": rv.get("voucher_number") or f"RV-{voucher_id}",
        },
        {
            "chart_account_id": credit_acct,
            "project_id": rv.get("project_id"),
            "party_name": rv.get("client_name"),
            "debit": 0,
            "credit": taxable,
            "line_narration": "Receipt (taxable)",
        },
    ]
    if gst_total > 0:
        gst_head = get_chart_account_by_code(db, COA_GST)
        if gst_head:
            lines.append({
                "chart_account_id": gst_head["id"],
                "project_id": rv.get("project_id"),
                "party_name": rv.get("client_name"),
                "debit": 0,
                "credit": gst_total,
                "line_narration": "Output GST",
            })
    je_id = _insert_journal(
        db,
        rv.get("receipt_date") or _now_ts()[:10],
        "Receipt",
        "receipt_vouchers",
        voucher_id,
        f"Receipt {rv.get('voucher_number') or voucher_id}",
        lines,
        username,
    )
    db.execute("UPDATE receipt_vouchers SET journal_entry_id=? WHERE id=?", (je_id, voucher_id))
    return je_id


def _save_payment_allocations(db, voucher_id: int, form) -> None:
    expense_ids = form.getlist("alloc_expense_id[]")
    amounts = form.getlist("alloc_amount[]")
    if not expense_ids:
        return
    db.execute("DELETE FROM payment_allocations WHERE payment_voucher_id=?", (voucher_id,))
    for idx, exp_id in enumerate(expense_ids):
        if not exp_id:
            continue
        amt = _safe_float(amounts[idx] if idx < len(amounts) else 0)
        if amt <= 0:
            continue
        db.execute(
            "INSERT INTO payment_allocations(payment_voucher_id, expense_id, allocated_amount) VALUES(?,?,?)",
            (voucher_id, int(exp_id), amt),
        )


def _apply_payment_allocations(db, voucher_id: int) -> None:
    rows = db.execute(
        "SELECT expense_id, allocated_amount FROM payment_allocations WHERE payment_voucher_id=?",
        (voucher_id,),
    ).fetchall()
    for row in rows:
        exp_id = row["expense_id"]
        if not exp_id:
            continue
        exp = db.execute(
            "SELECT grand_total, amount_paid FROM account_expenses WHERE id=?",
            (exp_id,),
        ).fetchone()
        if not exp:
            continue
        paid = round(_safe_float(exp["amount_paid"]) + _safe_float(row["allocated_amount"]), 2)
        grand = _safe_float(exp["grand_total"])
        if paid >= grand:
            status = "Paid"
        elif paid > 0:
            status = "Partially Paid"
        else:
            status = "Unpaid"
        db.execute(
            "UPDATE account_expenses SET amount_paid=?, payment_status=? WHERE id=?",
            (paid, status, exp_id),
        )


def _tds_due_date(entry_date: str) -> str:
    try:
        dt = datetime.strptime(entry_date[:10], "%Y-%m-%d")
    except ValueError:
        dt = datetime.now()
    due = dt.replace(day=7) + timedelta(days=32)
    due = due.replace(day=7)
    return due.strftime("%Y-%m-%d")


def _sync_tds_register_for_expense(db, expense_id: int) -> None:
    exp = db.execute("SELECT * FROM account_expenses WHERE id=?", (expense_id,)).fetchone()
    if not exp:
        return
    exp = dict(exp)
    db.execute("DELETE FROM tds_register WHERE source_type='expense' AND source_id=?", (expense_id,))
    if not exp.get("tds_applicable") or _safe_float(exp.get("tds_amount")) <= 0:
        return
    db.execute(
        "INSERT INTO tds_register(source_type, source_id, entry_date, vendor_name, tds_section, "
        "tds_rate, tds_amount, gross_amount, payment_status, filing_status, due_date) "
        "VALUES('expense',?,?,?,?,?,?,?, 'Deducted', 'Pending', ?)",
        (
            expense_id,
            exp.get("entry_date"),
            exp.get("vendor_name"),
            exp.get("tds_section"),
            _safe_float(exp.get("tds_rate")),
            _safe_float(exp.get("tds_amount")),
            _safe_float(exp.get("grand_total")),
            _tds_due_date(exp.get("entry_date") or ""),
        ),
    )


def _sync_tds_register_for_payment(db, voucher_id: int) -> None:
    pv = db.execute("SELECT * FROM payment_vouchers WHERE id=?", (voucher_id,)).fetchone()
    if not pv:
        return
    pv = dict(pv)
    db.execute("DELETE FROM tds_register WHERE source_type='payment' AND source_id=?", (voucher_id,))
    if not pv.get("tds_applicable") or _safe_float(pv.get("tds_amount")) <= 0:
        return
    db.execute(
        "INSERT INTO tds_register(source_type, source_id, entry_date, vendor_name, tds_section, "
        "tds_rate, tds_amount, gross_amount, payment_status, filing_status, due_date) "
        "VALUES('payment',?,?,?,?,?,?,?, 'Deducted', 'Pending', ?)",
        (
            voucher_id,
            pv.get("voucher_date"),
            pv.get("vendor_name"),
            pv.get("tds_section"),
            _safe_float(pv.get("tds_rate")),
            _safe_float(pv.get("tds_amount")),
            _safe_float(pv.get("gross_amount") or pv.get("amount")),
            _tds_due_date(pv.get("voucher_date") or ""),
        ),
    )


def save_account_attachment(db, entity_type: str, entity_id: int, filename: str, username: str) -> None:
    db.execute(
        "INSERT INTO account_attachments(entity_type, entity_id, filename, uploaded_by, uploaded_at) "
        "VALUES(?,?,?,?,?)",
        (entity_type, entity_id, filename, username, _now_ts()),
    )


def list_account_attachments(db, entity_type: str, entity_id: int) -> list[dict]:
    rows = db.execute(
        "SELECT * FROM account_attachments WHERE entity_type=? AND entity_id=? ORDER BY id DESC",
        (entity_type, entity_id),
    ).fetchall()
    return [dict(r) for r in rows]


def _journal_lines_query(from_date: str | None = None, to_date: str | None = None) -> tuple[str, list]:
    clauses = ["je.is_void=0"]
    params: list[Any] = []
    if from_date:
        clauses.append("je.entry_date >= ?")
        params.append(from_date)
    if to_date:
        clauses.append("je.entry_date <= ?")
        params.append(to_date)
    where = " AND ".join(clauses)
    sql = (
        "SELECT jel.*, je.entry_date, je.entry_number, je.entry_type, je.narration, "
        "je.reference_type, je.reference_id, c.code AS account_code, c.name AS account_name, "
        "c.account_type, p.project_name "
        "FROM journal_entry_lines jel "
        "JOIN journal_entries je ON jel.journal_entry_id = je.id "
        "JOIN chart_of_accounts c ON jel.chart_account_id = c.id "
        "LEFT JOIN projects p ON jel.project_id = p.id "
        f"WHERE {where} "
        "ORDER BY je.entry_date, je.id, jel.sort_order"
    )
    return sql, params


def get_cash_book_v2(db, from_date: str | None = None, to_date: str | None = None) -> list[dict]:
    cash = get_chart_account_by_code(db, COA_CASH)
    petty = get_chart_account_by_code(db, COA_PETTY)
    ids = [h["id"] for h in (cash, petty) if h]
    if not ids:
        return []
    placeholders = ",".join("?" * len(ids))
    sql, params = _journal_lines_query(from_date, to_date)
    sql = sql.replace("WHERE", f"WHERE jel.chart_account_id IN ({placeholders}) AND", 1)
    rows = db.execute(sql, ids + params).fetchall()
    return [dict(r) for r in rows]


def get_bank_book_v2(db, from_date: str | None = None, to_date: str | None = None) -> list[dict]:
    bank = get_chart_account_by_code(db, COA_BANK)
    if not bank:
        return []
    sql, params = _journal_lines_query(from_date, to_date)
    sql = sql.replace("WHERE", "WHERE jel.chart_account_id=? AND", 1)
    rows = db.execute(sql, [bank["id"]] + params).fetchall()
    return [dict(r) for r in rows]


def get_day_book(db, from_date: str | None = None, to_date: str | None = None) -> list[dict]:
    """Chronological journal lines across all accounts (day book)."""
    sql, params = _journal_lines_query(from_date, to_date)
    rows = db.execute(sql, params).fetchall()
    return [dict(r) for r in rows]


def get_general_ledger(db, account_id: int | None = None, from_date: str | None = None, to_date: str | None = None) -> list[dict]:
    sql, params = _journal_lines_query(from_date, to_date)
    if account_id:
        sql = sql.replace("WHERE", "WHERE jel.chart_account_id=? AND", 1)
        params = [account_id] + params
    rows = db.execute(sql, params).fetchall()
    return [dict(r) for r in rows]


def get_vendor_ledger(db, vendor_name: str, from_date: str | None = None, to_date: str | None = None) -> list[dict]:
    vendor = (vendor_name or "").strip()
    if not vendor:
        return []
    rows: list[dict] = []
    exp_sql = (
        "SELECT entry_date AS txn_date, 'Expense' AS txn_type, id AS ref_id, invoice_number AS ref_no, "
        "grand_total AS debit, 0 AS credit, payment_status AS status, approval_status "
        "FROM account_expenses WHERE vendor_name=? AND approval_status='Approved'"
    )
    params: list[Any] = [vendor]
    if from_date:
        exp_sql += " AND entry_date>=?"
        params.append(from_date)
    if to_date:
        exp_sql += " AND entry_date<=?"
        params.append(to_date)
    for r in db.execute(exp_sql + " ORDER BY entry_date, id", params).fetchall():
        rows.append(dict(r))
    pay_sql = (
        "SELECT voucher_date AS txn_date, 'Payment' AS txn_type, id AS ref_id, voucher_number AS ref_no, "
        "0 AS debit, COALESCE(gross_amount, amount) AS credit, payment_status AS status, approval_status "
        "FROM payment_vouchers WHERE vendor_name=? AND approval_status='Approved'"
    )
    pay_params: list[Any] = [vendor]
    if from_date:
        pay_sql += " AND voucher_date>=?"
        pay_params.append(from_date)
    if to_date:
        pay_sql += " AND voucher_date<=?"
        pay_params.append(to_date)
    for r in db.execute(pay_sql + " ORDER BY voucher_date, id", pay_params).fetchall():
        rows.append(dict(r))
    rows.sort(key=lambda x: (x.get("txn_date") or "", x.get("ref_id") or 0))
    balance = 0.0
    for row in rows:
        balance += _safe_float(row.get("debit")) - _safe_float(row.get("credit"))
        row["balance"] = round(balance, 2)
    return rows


def get_client_ledger(db, client_name: str, from_date: str | None = None, to_date: str | None = None) -> list[dict]:
    client = (client_name or "").strip()
    if not client:
        return []
    sql = (
        "SELECT receipt_date AS txn_date, 'Receipt' AS txn_type, id AS ref_id, voucher_number AS ref_no, "
        "0 AS debit, COALESCE(grand_total, amount) AS credit, approval_status AS status "
        "FROM receipt_vouchers WHERE client_name=? AND approval_status='Approved'"
    )
    params: list[Any] = [client]
    if from_date:
        sql += " AND receipt_date>=?"
        params.append(from_date)
    if to_date:
        sql += " AND receipt_date<=?"
        params.append(to_date)
    rows = [dict(r) for r in db.execute(sql + " ORDER BY receipt_date, id", params).fetchall()]
    balance = 0.0
    for row in rows:
        balance -= _safe_float(row.get("credit"))
        row["balance"] = round(balance, 2)
    return rows


def _account_balances(db, from_date: str | None, to_date: str | None) -> list[dict]:
    sql = (
        "SELECT c.id, c.code, c.name, c.account_type, "
        "COALESCE(SUM(jel.debit),0) AS total_debit, COALESCE(SUM(jel.credit),0) AS total_credit "
        "FROM chart_of_accounts c "
        "LEFT JOIN journal_entry_lines jel ON jel.chart_account_id=c.id "
        "LEFT JOIN journal_entries je ON jel.journal_entry_id=je.id AND je.is_void=0 "
    )
    clauses = ["c.is_active=1"]
    params: list[Any] = []
    if from_date:
        clauses.append("(je.entry_date IS NULL OR je.entry_date>=?)")
        params.append(from_date)
    if to_date:
        clauses.append("(je.entry_date IS NULL OR je.entry_date<=?)")
        params.append(to_date)
    sql += " WHERE " + " AND ".join(clauses) + " GROUP BY c.id ORDER BY c.account_type, c.sort_order, c.code"
    rows = db.execute(sql, params).fetchall()
    result = []
    for r in rows:
        row = dict(r)
        debit = _safe_float(row["total_debit"])
        credit = _safe_float(row["total_credit"])
        acct_type = row.get("account_type") or ""
        if acct_type in ("Asset", "Expense"):
            row["balance"] = round(debit - credit, 2)
        else:
            row["balance"] = round(credit - debit, 2)
        result.append(row)
    return result


def get_trial_balance(db, from_date: str | None = None, to_date: str | None = None) -> list[dict]:
    rows = _account_balances(db, from_date, to_date)
    tb = []
    for row in rows:
        debit = _safe_float(row.get("total_debit"))
        credit = _safe_float(row.get("total_credit"))
        if debit == 0 and credit == 0:
            continue
        tb.append({
            "code": row["code"],
            "name": row["name"],
            "account_type": row["account_type"],
            "debit": debit,
            "credit": credit,
        })
    return tb


def get_profit_and_loss(db, from_date: str | None = None, to_date: str | None = None) -> dict[str, Any]:
    rows = _account_balances(db, from_date, to_date)
    income = [r for r in rows if r.get("account_type") == "Income" and r["balance"] != 0]
    expenses = [r for r in rows if r.get("account_type") == "Expense" and r["balance"] != 0]
    total_income = round(sum(r["balance"] for r in income), 2)
    total_expense = round(sum(r["balance"] for r in expenses), 2)
    return {
        "income": income,
        "expenses": expenses,
        "total_income": total_income,
        "total_expense": total_expense,
        "net_profit": round(total_income - total_expense, 2),
    }


def get_balance_sheet(db, as_of_date: str | None = None) -> dict[str, Any]:
    rows = _account_balances(db, None, as_of_date)
    assets = [r for r in rows if r.get("account_type") == "Asset" and r["balance"] != 0]
    liabilities = [r for r in rows if r.get("account_type") == "Liability" and r["balance"] != 0]
    total_assets = round(sum(r["balance"] for r in assets), 2)
    total_liabilities = round(sum(r["balance"] for r in liabilities), 2)
    pl = get_profit_and_loss(db, None, as_of_date)
    retained = pl["net_profit"]
    return {
        "assets": assets,
        "liabilities": liabilities,
        "total_assets": total_assets,
        "total_liabilities": total_liabilities,
        "retained_earnings": retained,
        "total_liabilities_equity": round(total_liabilities + retained, 2),
    }


def get_cash_flow_summary(db, from_date: str | None = None, to_date: str | None = None) -> dict[str, Any]:
    cash_rows = get_cash_book_v2(db, from_date, to_date)
    bank_rows = get_bank_book_v2(db, from_date, to_date)
    inflows = round(
        sum(_safe_float(r.get("debit")) for r in cash_rows + bank_rows), 2
    )
    outflows = round(
        sum(_safe_float(r.get("credit")) for r in cash_rows + bank_rows), 2
    )
    return {
        "cash_lines": cash_rows,
        "bank_lines": bank_rows,
        "total_inflow": inflows,
        "total_outflow": outflows,
        "net_cash_flow": round(inflows - outflows, 2),
    }


def get_project_profitability(db, project_id: int | None = None, from_date: str | None = None, to_date: str | None = None) -> list[dict]:
    clauses = ["e.approval_status='Approved'"]
    params: list[Any] = []
    if project_id:
        clauses.append("e.project_id=?")
        params.append(project_id)
    if from_date:
        clauses.append("e.entry_date>=?")
        params.append(from_date)
    if to_date:
        clauses.append("e.entry_date<=?")
        params.append(to_date)
    where = " AND ".join(clauses)
    if project_id:
        sql = (
            "SELECT p.id AS project_id, p.project_name, "
            "COALESCE(SUM(e.grand_total),0) AS account_expense_total, "
            "COALESCE(SUM(e.subtotal),0) AS expense_subtotal, COUNT(e.id) AS expense_count "
            "FROM projects p "
            f"LEFT JOIN account_expenses e ON e.project_id=p.id AND {where} "
            "WHERE p.id=? GROUP BY p.id"
        )
        params.append(project_id)
    else:
        sql = (
            "SELECT p.id AS project_id, p.project_name, "
            "COALESCE(SUM(e.grand_total),0) AS account_expense_total, "
            "COALESCE(SUM(e.subtotal),0) AS expense_subtotal, COUNT(e.id) AS expense_count "
            "FROM projects p "
            f"LEFT JOIN account_expenses e ON e.project_id=p.id AND {where} "
            "GROUP BY p.id ORDER BY p.project_name"
        )
    rows = db.execute(sql, params).fetchall()
    result = []
    for row in rows:
        item = dict(row)
        pid = item["project_id"]
        receipts = 0.0
        if _table_exists(db, "receipt_vouchers"):
            rc = db.execute(
                "SELECT COALESCE(SUM(COALESCE(grand_total, amount)),0) AS t FROM receipt_vouchers "
                "WHERE project_id=? AND approval_status='Approved'",
                (pid,),
            ).fetchone()
            receipts = _safe_float(rc["t"] if rc else 0)
        pe_total = 0.0
        if _table_exists(db, "project_expenses"):
            pe = db.execute(
                "SELECT COALESCE(SUM(amount),0) AS t FROM project_expenses WHERE project_id=?",
                (pid,),
            ).fetchone()
            pe_total = _safe_float(pe["t"] if pe else 0)
        item["receipt_total"] = receipts
        item["project_expense_legacy"] = pe_total
        item["total_cost"] = round(_safe_float(item["account_expense_total"]) + pe_total, 2)
        item["profit"] = round(receipts - item["total_cost"], 2)
        result.append(item)
    return result


def list_tds_register(db) -> list[dict]:
    rows = db.execute(
        "SELECT * FROM tds_register ORDER BY entry_date DESC, id DESC"
    ).fetchall()
    today = datetime.now().date()
    result = []
    for r in rows:
        row = dict(r)
        due = row.get("due_date")
        alert = None
        if due and row.get("filing_status") != "Filed":
            try:
                due_dt = datetime.strptime(due[:10], "%Y-%m-%d").date()
                days = (due_dt - today).days
                if days < 0:
                    alert = "Overdue"
                elif days <= 1:
                    alert = "1 day"
                elif days <= 3:
                    alert = "3 days"
                elif days <= 7:
                    alert = "7 days"
            except ValueError:
                pass
        row["due_alert"] = alert
        result.append(row)
    return result


def list_gst_filing_periods(db) -> list[dict]:
    if not _table_exists(db, "gst_filing_periods"):
        return []
    rows = db.execute(
        "SELECT * FROM gst_filing_periods ORDER BY period_year DESC, period_month DESC"
    ).fetchall()
    today = datetime.now().date()
    result = []
    for r in rows:
        row = dict(r)
        due = row.get("due_date")
        alert = None
        if due and row.get("filing_status") != "Filed":
            try:
                due_dt = datetime.strptime(due[:10], "%Y-%m-%d").date()
                days = (due_dt - today).days
                if days < 0:
                    alert = "Overdue"
                elif days <= 1:
                    alert = "1 day"
                elif days <= 3:
                    alert = "3 days"
                elif days <= 7:
                    alert = "7 days"
            except ValueError:
                pass
        row["due_alert"] = alert
        result.append(row)
    return result


def update_gst_filing_status(db, period_id: int, status: str, filed_date: str | None = None) -> None:
    db.execute(
        "UPDATE gst_filing_periods SET filing_status=?, filed_date=? WHERE id=?",
        (status, filed_date, period_id),
    )


def get_gstr_summary(db, from_date: str | None = None, to_date: str | None = None) -> dict[str, Any]:
    purchase = get_gst_purchase_register(db)
    sales = get_gst_sales_register(db)
    if from_date:
        purchase = [r for r in purchase if (r.get("entry_date") or "") >= from_date]
        sales = [r for r in sales if (r.get("entry_date") or "") >= from_date]
    if to_date:
        purchase = [r for r in purchase if (r.get("entry_date") or "") <= to_date]
        sales = [r for r in sales if (r.get("entry_date") or "") <= to_date]
    out = {
        "purchase_taxable": round(sum(_safe_float(r.get("taxable_value")) for r in purchase), 2),
        "purchase_cgst": round(sum(_safe_float(r.get("cgst")) for r in purchase), 2),
        "purchase_sgst": round(sum(_safe_float(r.get("sgst")) for r in purchase), 2),
        "purchase_igst": round(sum(_safe_float(r.get("igst")) for r in purchase), 2),
        "sales_taxable": round(sum(_safe_float(r.get("taxable_value")) for r in sales), 2),
        "sales_cgst": round(sum(_safe_float(r.get("cgst")) for r in sales), 2),
        "sales_sgst": round(sum(_safe_float(r.get("sgst")) for r in sales), 2),
        "sales_igst": round(sum(_safe_float(r.get("igst")) for r in sales), 2),
    }
    out["net_cgst"] = round(out["purchase_cgst"] - out["sales_cgst"], 2)
    out["net_sgst"] = round(out["purchase_sgst"] - out["sales_sgst"], 2)
    out["net_igst"] = round(out["purchase_igst"] - out["sales_igst"], 2)
    return out


def build_pf_esi_from_payroll(db, period_month: int, period_year: int) -> int:
    if not _table_exists(db, "staff"):
        return 0
    staff_rows = db.execute(
        "SELECT id, staff_name, salary_amount, pf_applicable, esi_applicable, pf_rate, esi_rate "
        "FROM staff WHERE status='Active' OR status IS NULL"
    ).fetchall()
    count = 0
    now = _now_ts()
    due = f"{period_year:04d}-{period_month:02d}-15"
    for s in staff_rows:
        s = dict(s)
        salary = _safe_float(s.get("salary_amount"))
        if s.get("pf_applicable"):
            emp = round(salary * _safe_float(s.get("pf_rate"), 12) / 100.0, 2)
            er = emp
            db.execute(
                "INSERT INTO pf_esi_register(register_type, staff_id, staff_name, period_month, "
                "period_year, employee_contribution, employer_contribution, total_amount, due_date, created_at) "
                "VALUES('PF',?,?,?,?,?,?,?,?,?)",
                (s["id"], s["staff_name"], period_month, period_year, emp, er, emp + er, due, now),
            )
            count += 1
        if s.get("esi_applicable"):
            emp = round(salary * _safe_float(s.get("esi_rate"), 0.75) / 100.0, 2)
            er = round(salary * 3.25 / 100.0, 2)
            db.execute(
                "INSERT INTO pf_esi_register(register_type, staff_id, staff_name, period_month, "
                "period_year, employee_contribution, employer_contribution, total_amount, due_date, created_at) "
                "VALUES('ESI',?,?,?,?,?,?,?,?,?)",
                (s["id"], s["staff_name"], period_month, period_year, emp, er, emp + er, due, now),
            )
            count += 1
    return count


def _filing_due_alert(due_date: str | None, filing_status: str | None) -> str | None:
    if not due_date or filing_status == "Filed":
        return None
    try:
        due_dt = datetime.strptime(due_date[:10], "%Y-%m-%d").date()
        days = (due_dt - datetime.now().date()).days
        if days < 0:
            return "Overdue"
        if days <= 1:
            return "1 day"
        if days <= 3:
            return "3 days"
        if days <= 7:
            return "7 days"
    except ValueError:
        pass
    return None


def list_pf_esi_register(db, register_type: str | None = None) -> list[dict]:
    if register_type:
        rows = db.execute(
            "SELECT * FROM pf_esi_register WHERE register_type=? ORDER BY period_year DESC, period_month DESC, id DESC",
            (register_type,),
        ).fetchall()
    else:
        rows = db.execute(
            "SELECT * FROM pf_esi_register ORDER BY period_year DESC, period_month DESC, id DESC"
        ).fetchall()
    result = []
    for r in rows:
        row = dict(r)
        row["due_alert"] = _filing_due_alert(row.get("due_date"), row.get("filing_status"))
        result.append(row)
    return result


def update_pf_esi_filing_status(
    db,
    register_id: int,
    status: str,
    filed_date: str | None = None,
    remarks: str | None = None,
) -> None:
    db.execute(
        "UPDATE pf_esi_register SET filing_status=?, filed_date=?, remarks=? WHERE id=?",
        (status, filed_date, (remarks or "").strip() or None, register_id),
    )


def create_expense_draft_from_petty_cash(db, request_id: int, username: str) -> int | None:
    if not _table_exists(db, "petty_cash_requests"):
        return None
    row = db.execute("SELECT * FROM petty_cash_requests WHERE id=?", (request_id,)).fetchone()
    if not row:
        return None
    row = dict(row)
    total = _safe_float(row.get("expenses_total") or row.get("transferred_amount"))
    if total <= 0:
        return None
    head = get_chart_account_by_code(db, "E006") or get_chart_account_by_code(db, "E005")
    if not head:
        heads = list_chart_of_accounts(db)
        head = heads[0] if heads else None
    if not head:
        return None
    now = _now_ts()
    db.execute(
        "INSERT INTO account_expenses("
        "entry_date, project_id, chart_account_id, vendor_name, expense_type, payment_source, "
        "petty_cash_request_id, payment_status, subtotal, grand_total, remarks, created_by, created_at, "
        "modified_by, modified_at, approval_status"
        ") VALUES(?,?,?,?, 'Expense', 'Petty Cash', ?, 'Draft', ?, ?, ?, ?, ?, ?, ?, 'Pending Checker')",
        (
            row.get("request_date") or now[:10],
            row.get("project_id"),
            head["id"],
            row.get("staff_name") or "Petty Cash",
            request_id,
            total,
            total,
            f"Auto-draft from petty cash {row.get('request_number') or request_id}",
            username,
            now,
            username,
            now,
        ),
    )
    exp_id = db.execute("SELECT last_insert_rowid()").fetchone()[0]
    db.execute(
        "INSERT INTO account_expense_lines(expense_id, item_name, quantity, rate, amount, taxable_value, "
        "line_total, sort_order) VALUES(?,?,1,?,?,?,?,0)",
        (exp_id, row.get("purpose") or "Petty cash settlement", total, total, total, total),
    )
    return exp_id


def export_report_excel(rows: list[dict], sheet_name: str = "Report") -> BytesIO:
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = sheet_name[:31]
    if not rows:
        ws.append(["No data"])
    else:
        headers = list(rows[0].keys())
        ws.append(headers)
        for row in rows:
            ws.append([row.get(h) for h in headers])
    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def export_report_csv(rows: list[dict]) -> str:
    import csv
    from io import StringIO

    si = StringIO()
    if not rows:
        si.write("No data\n")
        return si.getvalue()
    writer = csv.writer(si)
    headers = list(rows[0].keys())
    writer.writerow(headers)
    for row in rows:
        writer.writerow([row.get(h) for h in headers])
    return si.getvalue()


def export_tally_xml(db, from_date: str | None = None, to_date: str | None = None) -> str:
    sql = (
        "SELECT je.*, jel.debit, jel.credit, jel.line_narration, c.name AS account_name "
        "FROM journal_entries je "
        "JOIN journal_entry_lines jel ON jel.journal_entry_id=je.id "
        "JOIN chart_of_accounts c ON jel.chart_account_id=c.id "
        "WHERE je.is_void=0"
    )
    params: list[Any] = []
    if from_date:
        sql += " AND je.entry_date>=?"
        params.append(from_date)
    if to_date:
        sql += " AND je.entry_date<=?"
        params.append(to_date)
    sql += " ORDER BY je.entry_date, je.id, jel.sort_order"
    rows = db.execute(sql, params).fetchall()
    root = ET.Element("ENVELOPE")
    header = ET.SubElement(root, "HEADER")
    ET.SubElement(header, "TALLYREQUEST").text = "Import Data"
    body = ET.SubElement(root, "BODY")
    import_data = ET.SubElement(body, "IMPORTDATA")
    request_data = ET.SubElement(import_data, "REQUESTDATA")
    current_je = None
    voucher_el = None
    for row in rows:
        row = dict(row)
        if row["id"] != current_je:
            current_je = row["id"]
            tally_msg = ET.SubElement(request_data, "TALLYMESSAGE")
            voucher_el = ET.SubElement(tally_msg, "VOUCHER")
            voucher_el.set("VCHTYPE", row.get("entry_type") or "Journal")
            voucher_el.set("ACTION", "Create")
            ET.SubElement(voucher_el, "DATE").text = (row.get("entry_date") or "").replace("-", "")
            ET.SubElement(voucher_el, "VOUCHERNUMBER").text = row.get("entry_number") or ""
            ET.SubElement(voucher_el, "NARRATION").text = row.get("narration") or ""
        if not voucher_el:
            continue
        ledger = ET.SubElement(voucher_el, "ALLLEDGERENTRIES.LIST")
        ET.SubElement(ledger, "LEDGERNAME").text = row.get("account_name") or ""
        if _safe_float(row.get("debit")) > 0:
            ET.SubElement(ledger, "ISDEEMEDPOSITIVE").text = "Yes"
            ET.SubElement(ledger, "AMOUNT").text = f"-{_safe_float(row.get('debit')):.2f}"
        else:
            ET.SubElement(ledger, "ISDEEMEDPOSITIVE").text = "No"
            ET.SubElement(ledger, "AMOUNT").text = f"{_safe_float(row.get('credit')):.2f}"
    return ET.tostring(root, encoding="unicode")


def list_unpaid_expenses_for_vendor(db, vendor_name: str) -> list[dict]:
    vendor = (vendor_name or "").strip()
    if not vendor:
        return []
    rows = db.execute(
        "SELECT id, entry_date, invoice_number, grand_total, amount_paid, payment_status "
        "FROM account_expenses WHERE vendor_name=? AND approval_status='Approved' "
        "AND payment_status IN ('Unpaid', 'Partially Paid') ORDER BY entry_date",
        (vendor,),
    ).fetchall()
    return [dict(r) for r in rows]


def get_petty_cash_balance(db, request_id: int) -> float:
    if not _table_exists(db, "petty_cash_requests"):
        return 0.0
    row = db.execute(
        "SELECT transferred_amount, expenses_total FROM petty_cash_requests WHERE id=?",
        (request_id,),
    ).fetchone()
    if not row:
        return 0.0
    return round(
        _safe_float(row["transferred_amount"]) - _safe_float(row["expenses_total"]),
        2,
    )


def _safe_count(db, sql: str, default: int = 0) -> int:
    try:
        row = db.execute(sql).fetchone()
        return int(row["c"] if row else default)
    except sqlite3.OperationalError:
        return default


def accounts_hub_stats(db) -> dict[str, Any]:
    stats = {
        "chart_heads": 0,
        "expenses": 0,
        "payments": 0,
        "receipts": 0,
        "pending_expenses": 0,
        "journal_entries": 0,
        "tds_pending": 0,
    }
    if _table_exists(db, "chart_of_accounts"):
        stats["chart_heads"] = _safe_count(
            db,
            "SELECT COUNT(*) AS c FROM chart_of_accounts WHERE COALESCE(is_active, 1)=1",
        )
    if _table_exists(db, "account_expenses"):
        stats["expenses"] = _safe_count(db, "SELECT COUNT(*) AS c FROM account_expenses")
        stats["pending_expenses"] = _safe_count(
            db,
            "SELECT COUNT(*) AS c FROM account_expenses "
            "WHERE approval_status IN ('Pending Checker', 'Pending Approval')",
        )
    if _table_exists(db, "payment_vouchers"):
        stats["payments"] = _safe_count(db, "SELECT COUNT(*) AS c FROM payment_vouchers")
    if _table_exists(db, "receipt_vouchers"):
        stats["receipts"] = _safe_count(db, "SELECT COUNT(*) AS c FROM receipt_vouchers")
    if _table_exists(db, "journal_entries"):
        stats["journal_entries"] = _safe_count(
            db, "SELECT COUNT(*) AS c FROM journal_entries WHERE is_void=0"
        )
    if _table_exists(db, "tds_register"):
        stats["tds_pending"] = _safe_count(
            db, "SELECT COUNT(*) AS c FROM tds_register WHERE filing_status!='Filed'"
        )
    return stats


def chart_heads_payload(db, account_type: str | None = None) -> list[dict[str, Any]]:
    if account_type:
        normalized = _normalize_account_type(account_type)
        if normalized == "Expense":
            heads = list_expense_chart_heads(db)
        elif normalized == "Income":
            heads = list_income_chart_heads(db)
        else:
            heads = [
                h for h in list_chart_of_accounts(db)
                if h.get("account_type") == normalized
            ]
    else:
        heads = list_chart_of_accounts(db)
    return [
        {
            "id": h["id"],
            "code": h["code"],
            "name": h["name"],
            "account_type": h.get("account_type"),
            "requires_project": bool(h.get("requires_project")),
            "requires_vendor": bool(h.get("requires_vendor")),
            "default_gst_applicable": bool(h.get("default_gst_applicable")),
        }
        for h in heads
    ]


def chart_accounts_for_js(db) -> str:
    return json.dumps(chart_heads_payload(db))
