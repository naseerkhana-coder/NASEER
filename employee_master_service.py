"""Employee Master (MODULE-014) — company staff registry (HR/payroll; not construction workers)."""

from __future__ import annotations

import csv
import io
import json
import os
import re
from io import BytesIO
from typing import Any

from branch_master_service import ensure_branch_master_schema, get_branch_master
from company_master_service import (
    _ensure_column,
    _now_ts,
    _table_exists,
    ensure_company_master_schema,
    get_company,
    list_companies,
    validate_email,
    validate_pan_number,
    validate_phone,
)
from department_master_service import ensure_department_master_schema, get_department_master
from designation_master_service import ensure_designation_master_schema, get_designation_master

EMPLOYEE_STATUSES = ("Active", "Inactive")
EMPLOYEE_TYPES = ("Permanent", "Contract", "Probation", "Intern", "Consultant")
EMPLOYMENT_STATUSES = ("Active", "On Leave", "Suspended", "Relieved", "Absconding")
APPROVAL_STATUSES = ("Draft", "Pending", "Approved", "Rejected")
GENDER_OPTIONS = ("Male", "Female", "Other", "Prefer not to say")
BLOOD_GROUPS = ("A+", "A-", "B+", "B-", "AB+", "AB-", "O+", "O-")
MARITAL_STATUSES = ("Single", "Married", "Divorced", "Widowed")
SALARY_TYPES = ("Monthly", "Daily", "Hourly")
PHOTO_EXTENSIONS = frozenset({".jpg", ".jpeg", ".png", ".webp"})
SIGNATURE_EXTENSIONS = frozenset({".jpg", ".jpeg", ".png", ".webp", ".svg"})
AADHAAR_RE = re.compile(r"^\d{12}$")
MOBILE_RE = re.compile(r"^\+?[\d\s\-()]{7,20}$")

EMPLOYEE_SORT_COLUMNS = (
    "employee_code",
    "staff_name",
    "department_name",
    "designation_name",
    "joining_date",
    "status",
    "employment_status",
    "created_at",
    "company_id",
)
EMPLOYEE_EXPORT_COLUMNS = (
    "employee_code",
    "staff_name",
    "first_name",
    "last_name",
    "employee_type",
    "company_code",
    "branch_code",
    "department_name",
    "designation_name",
    "reporting_manager_code",
    "official_email",
    "mobile",
    "joining_date",
    "confirmation_date",
    "relieving_date",
    "employment_status",
    "pan_number",
    "aadhaar_number",
    "pf_number",
    "esi_number",
    "uan_number",
    "bank_name",
    "bank_account",
    "ifsc_code",
    "salary_type",
    "salary_amount",
    "status",
    "approval_status",
)
EMPLOYEE_AUDIT_FIELDS = (
    "employee_code",
    "staff_name",
    "first_name",
    "middle_name",
    "last_name",
    "employee_type",
    "company_id",
    "branch_id",
    "department_id",
    "designation_id",
    "reporting_manager_id",
    "official_email",
    "personal_email",
    "mobile",
    "alternative_mobile",
    "joining_date",
    "confirmation_date",
    "relieving_date",
    "employment_status",
    "gender",
    "date_of_birth",
    "blood_group",
    "nationality",
    "marital_status",
    "aadhaar_number",
    "pan_number",
    "passport_number",
    "driving_license",
    "pf_number",
    "esi_number",
    "uan_number",
    "bank_name",
    "bank_account",
    "ifsc_code",
    "branch_name",
    "salary_type",
    "salary_amount",
    "salary_structure",
    "status",
    "approval_status",
)
EMPLOYEE_REFERENCE_TABLES = (
    ("users", "staff_id"),
    ("payroll_lines", "staff_id"),
    ("salary_revisions", "staff_id"),
    ("staff_salary_components", "staff_id"),
    ("staff_salary_increments", "staff_id"),
    ("staff_bonus", "staff_id"),
    ("staff_monthly_attendance", "staff_id"),
    ("leave_requests", "staff_id"),
    ("petty_cash_requests", "staff_id"),
)


def employee_upload_dir(base_dir: str, subfolder: str = "employees") -> str:
    path = os.path.join(base_dir, "static", "uploads", subfolder)
    os.makedirs(path, exist_ok=True)
    return path


def validate_aadhaar_number(value: str) -> None:
    text = re.sub(r"\s", "", (value or "").strip())
    if text and not AADHAAR_RE.match(text):
        raise ValueError("Enter a valid 12-digit Aadhaar number.")


def validate_employee_mobile(value: str) -> None:
    text = (value or "").strip()
    if not text:
        return
    if not MOBILE_RE.match(text):
        raise ValueError("Enter a valid mobile number.")
    validate_phone(text)


def ensure_employee_master_schema(db) -> None:
    """Extend staff table and child tables for MODULE-014 (idempotent)."""
    ensure_company_master_schema(db)
    ensure_branch_master_schema(db)
    ensure_department_master_schema(db)
    ensure_designation_master_schema(db)
    db.execute(
        """
        CREATE TABLE IF NOT EXISTS staff(
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            employee_code TEXT,
            staff_name TEXT,
            mobile TEXT,
            email TEXT,
            department TEXT,
            designation TEXT,
            salary_type TEXT,
            salary_amount REAL,
            ot_applicable TEXT,
            working_hours REAL,
            joining_date TEXT,
            photo TEXT,
            status TEXT
        )
        """
    )
    staff_columns = (
        ("employee_type", "TEXT DEFAULT 'Permanent'"),
        ("company_id", "INTEGER"),
        ("branch_id", "INTEGER"),
        ("department_id", "INTEGER"),
        ("reporting_manager_id", "INTEGER"),
        ("first_name", "TEXT"),
        ("middle_name", "TEXT"),
        ("last_name", "TEXT"),
        ("official_email", "TEXT"),
        ("personal_email", "TEXT"),
        ("alternative_mobile", "TEXT"),
        ("current_address", "TEXT"),
        ("permanent_address", "TEXT"),
        ("city", "TEXT"),
        ("state", "TEXT"),
        ("pin_code", "TEXT"),
        ("country", "TEXT DEFAULT 'India'"),
        ("gender", "TEXT"),
        ("date_of_birth", "TEXT"),
        ("blood_group", "TEXT"),
        ("nationality", "TEXT DEFAULT 'Indian'"),
        ("marital_status", "TEXT"),
        ("confirmation_date", "TEXT"),
        ("relieving_date", "TEXT"),
        ("employment_status", "TEXT DEFAULT 'Active'"),
        ("passport_number", "TEXT"),
        ("driving_license", "TEXT"),
        ("pf_number", "TEXT"),
        ("esi_number", "TEXT"),
        ("uan_number", "TEXT"),
        ("salary_structure", "TEXT"),
        ("profile_photo", "TEXT"),
        ("digital_signature", "TEXT"),
        ("designation_id", "INTEGER"),
        ("reporting_manager", "TEXT"),
        ("workflow_role", "TEXT"),
        ("ot_rate_per_hour", "REAL DEFAULT 0"),
        ("holiday_pay_applicable", "TEXT DEFAULT 'No'"),
        ("aadhaar_number", "TEXT"),
        ("pan_number", "TEXT"),
        ("bank_account", "TEXT"),
        ("bank_name", "TEXT"),
        ("ifsc_code", "TEXT"),
        ("branch_name", "TEXT"),
        ("id_proof", "TEXT"),
        ("aadhaar_document", "TEXT"),
        ("pan_document", "TEXT"),
        ("company_room_provided", "TEXT DEFAULT 'No'"),
        ("company_food_provided", "TEXT DEFAULT 'No'"),
        ("approval_status", "TEXT DEFAULT 'Draft'"),
        ("approved_by", "TEXT"),
        ("approved_at", "TEXT"),
        ("created_by", "TEXT"),
        ("created_at", "TEXT"),
        ("modified_by", "TEXT"),
        ("modified_at", "TEXT"),
        ("is_deleted", "INTEGER DEFAULT 0"),
        ("deleted_by", "TEXT"),
        ("deleted_at", "TEXT"),
        ("customer_id", "INTEGER"),
    )
    for col, ctype in staff_columns:
        _ensure_column(db, "staff", col, ctype)

    db.execute(
        """
        CREATE TABLE IF NOT EXISTS employee_emergency_contacts(
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            staff_id INTEGER NOT NULL,
            contact_name TEXT,
            relationship TEXT,
            phone TEXT,
            email TEXT,
            address TEXT,
            is_primary INTEGER DEFAULT 0,
            created_by TEXT,
            created_at TEXT,
            modified_by TEXT,
            modified_at TEXT,
            is_deleted INTEGER DEFAULT 0,
            FOREIGN KEY(staff_id) REFERENCES staff(id)
        )
        """
    )
    db.execute(
        """
        CREATE TABLE IF NOT EXISTS employee_qualifications(
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            staff_id INTEGER NOT NULL,
            qualification TEXT,
            institution TEXT,
            board_university TEXT,
            year_completed TEXT,
            grade TEXT,
            document_path TEXT,
            created_by TEXT,
            created_at TEXT,
            modified_by TEXT,
            modified_at TEXT,
            is_deleted INTEGER DEFAULT 0,
            FOREIGN KEY(staff_id) REFERENCES staff(id)
        )
        """
    )
    db.execute(
        """
        CREATE TABLE IF NOT EXISTS employee_experience(
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            staff_id INTEGER NOT NULL,
            company_name TEXT,
            designation TEXT,
            from_date TEXT,
            to_date TEXT,
            responsibilities TEXT,
            created_by TEXT,
            created_at TEXT,
            modified_by TEXT,
            modified_at TEXT,
            is_deleted INTEGER DEFAULT 0,
            FOREIGN KEY(staff_id) REFERENCES staff(id)
        )
        """
    )
    db.execute(
        """
        CREATE TABLE IF NOT EXISTS employee_skills(
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            staff_id INTEGER NOT NULL,
            skill_name TEXT,
            proficiency TEXT,
            years_experience REAL,
            certification TEXT,
            created_by TEXT,
            created_at TEXT,
            modified_by TEXT,
            modified_at TEXT,
            is_deleted INTEGER DEFAULT 0,
            FOREIGN KEY(staff_id) REFERENCES staff(id)
        )
        """
    )
    db.execute(
        """
        CREATE TABLE IF NOT EXISTS employee_employment_history(
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            staff_id INTEGER NOT NULL,
            effective_date TEXT,
            change_type TEXT,
            department_id INTEGER,
            designation_id INTEGER,
            branch_id INTEGER,
            reporting_manager_id INTEGER,
            remarks TEXT,
            created_by TEXT,
            created_at TEXT,
            FOREIGN KEY(staff_id) REFERENCES staff(id)
        )
        """
    )
    db.execute("UPDATE staff SET status='Active' WHERE status IS NULL OR TRIM(status)=''")
    db.execute(
        "UPDATE staff SET employment_status='Active' WHERE employment_status IS NULL OR TRIM(employment_status)=''"
    )
    _migrate_legacy_staff(db)
    try:
        from audit_trail_service import ensure_audit_schema

        ensure_audit_schema(db)
    except Exception:
        pass


def _migrate_legacy_staff(db) -> None:
    if not _table_exists(db, "staff"):
        return
    first_company = db.execute(
        "SELECT id FROM companies WHERE COALESCE(is_deleted,0)=0 ORDER BY id LIMIT 1"
    ).fetchone()
    company_id = int(first_company[0]) if first_company else None
    rows = db.execute(
        """
        SELECT id, staff_name, first_name, last_name, email, official_email, photo, profile_photo,
               department, department_id, designation, designation_id, company_id
        FROM staff WHERE COALESCE(is_deleted,0)=0
        """
    ).fetchall()
    for row in rows:
        sid = int(row[0])
        staff_name = (row[1] or "").strip()
        first = (row[2] or "").strip() if len(row) > 2 else ""
        last = (row[3] or "").strip() if len(row) > 3 else ""
        email = (row[4] or "").strip() if len(row) > 4 else ""
        official = (row[5] or "").strip() if len(row) > 5 else ""
        photo = (row[6] or "").strip() if len(row) > 6 else ""
        profile = (row[7] or "").strip() if len(row) > 7 else ""
        updates: list[str] = []
        params: list[Any] = []
        if company_id and not (row[12] if len(row) > 12 else None):
            updates.append("company_id=?")
            params.append(company_id)
        if not first and staff_name:
            parts = staff_name.split(None, 1)
            updates.extend(["first_name=?", "last_name=?"])
            params.extend([parts[0], parts[1] if len(parts) > 1 else ""])
        if not official and email:
            updates.append("official_email=?")
            params.append(email)
        if not profile and photo:
            updates.append("profile_photo=?")
            params.append(photo)
        dept_id = row[9] if len(row) > 9 else None
        dept_text = (row[8] or "").strip() if len(row) > 8 else ""
        if not dept_id and dept_text and company_id and _table_exists(db, "departments"):
            hit = db.execute(
                """
                SELECT id FROM departments
                WHERE department_name=? AND company_id=? AND COALESCE(is_deleted,0)=0 LIMIT 1
                """,
                (dept_text, company_id),
            ).fetchone()
            if hit:
                updates.append("department_id=?")
                params.append(int(hit[0]))
        des_id = row[11] if len(row) > 11 else None
        des_text = (row[10] or "").strip() if len(row) > 10 else ""
        if not des_id and des_text and _table_exists(db, "designations"):
            hit = db.execute(
                """
                SELECT id FROM designations
                WHERE designation_name=? AND COALESCE(is_deleted,0)=0 LIMIT 1
                """,
                (des_text,),
            ).fetchone()
            if hit:
                updates.append("designation_id=?")
                params.append(int(hit[0]))
        if updates:
            params.append(sid)
            db.execute(f"UPDATE staff SET {', '.join(updates)} WHERE id=?", params)


_EMP_CODE_RE = re.compile(r"^EMP\s*(\d+)$", re.IGNORECASE)


def _employee_code_number(code: str) -> int | None:
    """Parse EMP101 / EMP 114 style codes into their numeric suffix."""
    match = _EMP_CODE_RE.match(str(code or "").strip())
    if not match:
        return None
    try:
        return int(match.group(1))
    except ValueError:
        return None


def generate_employee_code(db) -> str:
    """Next EMP code from the highest existing staff employee_code (EMP101+).

    Includes soft-deleted rows so relieved/deleted employees still advance the
    sequence and codes are not reused.
    """
    rows = db.execute(
        """
        SELECT employee_code FROM staff
        WHERE employee_code IS NOT NULL AND TRIM(employee_code) != ''
        """
    ).fetchall()
    max_num = 100
    for row in rows:
        code = row[0] if not hasattr(row, "keys") else row["employee_code"]
        num = _employee_code_number(code)
        if num is not None:
            max_num = max(max_num, num)
    return f"EMP{max_num + 1}"


def validate_employee_uniqueness(
    db,
    *,
    employee_code: str,
    official_email: str = "",
    staff_id: int | None = None,
) -> None:
    code = (employee_code or "").strip().upper()
    if not code:
        raise ValueError("Employee code is required.")
    row = db.execute(
        """
        SELECT id FROM staff
        WHERE UPPER(employee_code)=? AND COALESCE(is_deleted,0)=0
        """,
        (code,),
    ).fetchone()
    if row and (not staff_id or int(row[0]) != int(staff_id)):
        raise ValueError(f"Employee code '{code}' already exists.")
    email = (official_email or "").strip().lower()
    if email:
        row = db.execute(
            """
            SELECT id FROM staff
            WHERE LOWER(COALESCE(official_email, email))=? AND COALESCE(is_deleted,0)=0
            """,
            (email,),
        ).fetchone()
        if row and (not staff_id or int(row[0]) != int(staff_id)):
            raise ValueError(f"Official email '{official_email}' is already registered.")


def validate_employee_form_data(data: dict[str, Any], *, staff_id: int | None = None) -> None:
    if not (data.get("first_name") or data.get("staff_name") or "").strip():
        raise ValueError("Employee name is required.")
    if not data.get("department_id"):
        raise ValueError("Department is required.")
    if not data.get("designation_id"):
        raise ValueError("Designation is required.")
    if not (data.get("joining_date") or "").strip():
        raise ValueError("Joining date is required.")
    if data.get("official_email"):
        validate_email(data["official_email"])
    if data.get("personal_email"):
        validate_email(data["personal_email"])
    if data.get("mobile"):
        validate_employee_mobile(data["mobile"])
    if data.get("alternative_mobile"):
        validate_employee_mobile(data["alternative_mobile"])
    if data.get("pan_number"):
        validate_pan_number(data["pan_number"])
    if data.get("aadhaar_number"):
        validate_aadhaar_number(data["aadhaar_number"])
    status = (data.get("status") or "Active").strip()
    if status not in EMPLOYEE_STATUSES:
        raise ValueError("Select a valid status.")
    emp_type = (data.get("employee_type") or "Permanent").strip()
    if emp_type not in EMPLOYEE_TYPES:
        raise ValueError("Select a valid employee type.")
    emp_status = (data.get("employment_status") or "Active").strip()
    if emp_status not in EMPLOYMENT_STATUSES:
        raise ValueError("Select a valid employment status.")


def employee_has_references(db, staff_id: int) -> list[str]:
    blockers: list[str] = []
    for table, col in EMPLOYEE_REFERENCE_TABLES:
        if not _table_exists(db, table):
            continue
        cols = {r[1] for r in db.execute(f"PRAGMA table_info({table})").fetchall()}
        if col not in cols:
            continue
        row = db.execute(f"SELECT COUNT(*) FROM {table} WHERE {col}=?", (staff_id,)).fetchone()
        if int(row[0] if row else 0) > 0:
            blockers.append(table)
    return blockers


def log_employee_audit(
    db,
    staff_id: int,
    action: str,
    username: str,
    *,
    field_name: str | None = None,
    old_value: str | None = None,
    new_value: str | None = None,
    remarks: str | None = None,
) -> None:
    try:
        from audit_trail_service import log_audit_event

        log_audit_event(
            db,
            record_table="staff",
            record_id=staff_id,
            action=action,
            changed_by=username,
            field_name=field_name,
            old_value=old_value,
            new_value=new_value,
            remarks=remarks,
        )
    except Exception:
        pass


def log_employee_field_changes(
    db,
    before: dict[str, Any] | None,
    after: dict[str, Any] | None,
    username: str,
) -> None:
    if not before or not after:
        return
    staff_id = int(after.get("id") or before.get("id") or 0)
    if not staff_id:
        return
    for field in EMPLOYEE_AUDIT_FIELDS:
        old_val = before.get(field)
        new_val = after.get(field)
        if str(old_val or "") != str(new_val or ""):
            log_employee_audit(
                db,
                staff_id,
                "update",
                username,
                field_name=field,
                old_value=str(old_val or ""),
                new_value=str(new_val or ""),
            )


def list_employee_audit_trail(db, staff_id: int, limit: int = 100) -> list[dict[str, Any]]:
    try:
        from audit_trail_service import list_audit_trail

        return list_audit_trail(db, "staff", staff_id, limit=limit)
    except Exception:
        return []


def _form_getlist(form, key: str) -> list[str]:
    if hasattr(form, "getlist"):
        return [str(v or "").strip() for v in form.getlist(key)]
    val = form.get(key) if hasattr(form, "get") else None
    if isinstance(val, list):
        return [str(v or "").strip() for v in val]
    return [str(val).strip()] if val not in (None, "") else []


def _form_value(form, key: str, default: str = "") -> str:
    if hasattr(form, "get"):
        return str(form.get(key) or default).strip()
    return str((form or {}).get(key) or default).strip()


def _employee_row_to_dict(row) -> dict[str, Any]:
    if not row:
        return {}
    data = dict(row)
    data["id"] = int(data.get("id") or 0)
    photo = data.get("profile_photo") or data.get("photo") or ""
    data["profile_photo"] = photo
    data["photo"] = photo
    data["email"] = data.get("official_email") or data.get("email") or ""
    return data


def _employee_base_sql() -> str:
    return (
        "SELECT s.*, c.company_code, c.company_name, b.branch_code, b.branch_name, "
        "d.department_code, d.department_name AS dept_master_name, "
        "des.designation_code, des.designation_name AS desig_master_name, "
        "rm.staff_name AS reporting_manager_name, rm.employee_code AS reporting_manager_code "
        "FROM staff s "
        "LEFT JOIN companies c ON s.company_id = c.id "
        "LEFT JOIN company_branches b ON s.branch_id = b.id "
        "LEFT JOIN departments d ON s.department_id = d.id "
        "LEFT JOIN designations des ON s.designation_id = des.id "
        "LEFT JOIN staff rm ON s.reporting_manager_id = rm.id "
        "WHERE 1=1"
    )


def list_employees_master(
    db,
    *,
    search: str = "",
    company_id: int | None = None,
    branch_id: int | None = None,
    department_id: int | None = None,
    designation_id: int | None = None,
    status: str = "",
    employment_status: str = "",
    employee_type: str = "",
    include_deleted: bool = False,
    page: int = 1,
    per_page: int = 25,
    sort_by: str = "staff_name",
    sort_dir: str = "asc",
) -> dict[str, Any]:
    sql = _employee_base_sql()
    count_sql = (
        "SELECT COUNT(*) FROM staff s "
        "LEFT JOIN departments d ON s.department_id = d.id "
        "LEFT JOIN designations des ON s.designation_id = des.id "
        "WHERE 1=1"
    )
    params: list[Any] = []
    if not include_deleted:
        sql += " AND COALESCE(s.is_deleted,0)=0"
        count_sql += " AND COALESCE(s.is_deleted,0)=0"
    if company_id:
        sql += " AND s.company_id=?"
        count_sql += " AND s.company_id=?"
        params.append(company_id)
    if branch_id:
        sql += " AND s.branch_id=?"
        count_sql += " AND s.branch_id=?"
        params.append(branch_id)
    if department_id:
        sql += " AND s.department_id=?"
        count_sql += " AND s.department_id=?"
        params.append(department_id)
    if designation_id:
        sql += " AND s.designation_id=?"
        count_sql += " AND s.designation_id=?"
        params.append(designation_id)
    if status:
        sql += " AND s.status=?"
        count_sql += " AND s.status=?"
        params.append(status)
    if employment_status:
        sql += " AND s.employment_status=?"
        count_sql += " AND s.employment_status=?"
        params.append(employment_status)
    if employee_type:
        sql += " AND s.employee_type=?"
        count_sql += " AND s.employee_type=?"
        params.append(employee_type)
    if search:
        clause = (
            " AND (s.staff_name LIKE ? OR s.employee_code LIKE ? OR s.first_name LIKE ? "
            "OR s.last_name LIKE ? OR s.official_email LIKE ? OR s.mobile LIKE ? "
            "OR s.department LIKE ? OR s.designation LIKE ? OR d.department_name LIKE ? "
            "OR des.designation_name LIKE ?)"
        )
        sql += clause
        count_sql += clause
        like = f"%{search}%"
        params.extend([like] * 10)
    sort_map = {
        "department_name": "d.department_name",
        "designation_name": "des.designation_name",
    }
    sort_col = sort_map.get(sort_by, f"s.{sort_by}" if sort_by in EMPLOYEE_SORT_COLUMNS else "s.staff_name")
    direction = "DESC" if str(sort_dir).lower() == "desc" else "ASC"
    sql += f" ORDER BY {sort_col} {direction}, s.id DESC"
    per_page = max(1, min(int(per_page or 25), 10000))
    page = max(1, int(page or 1))
    offset = (page - 1) * per_page
    sql += " LIMIT ? OFFSET ?"
    total = int(db.execute(count_sql, params).fetchone()[0])
    rows = db.execute(sql, [*params, per_page, offset]).fetchall()
    items = [_employee_row_to_dict(r) for r in rows]
    pages = (total + per_page - 1) // per_page if total else 0
    return {"items": items, "total": total, "page": page, "per_page": per_page, "pages": pages}


def get_employee_master(db, staff_id: int, *, include_deleted: bool = False) -> dict[str, Any] | None:
    if not staff_id or not _table_exists(db, "staff"):
        return None
    sql = _employee_base_sql() + " AND s.id=?"
    if not include_deleted:
        sql += " AND COALESCE(s.is_deleted,0)=0"
    row = db.execute(sql, (staff_id,)).fetchone()
    if not row:
        return None
    data = _employee_row_to_dict(row)
    data["emergency_contacts"] = list_employee_emergency_contacts(db, staff_id)
    data["qualifications"] = list_employee_qualifications(db, staff_id)
    data["experience"] = list_employee_experience(db, staff_id)
    data["skills"] = list_employee_skills(db, staff_id)
    data["employment_history"] = list_employee_employment_history(db, staff_id)
    return data


def list_employee_emergency_contacts(db, staff_id: int) -> list[dict[str, Any]]:
    rows = db.execute(
        """
        SELECT * FROM employee_emergency_contacts
        WHERE staff_id=? AND COALESCE(is_deleted,0)=0 ORDER BY is_primary DESC, id
        """,
        (staff_id,),
    ).fetchall()
    return [dict(r) for r in rows]


def list_employee_qualifications(db, staff_id: int) -> list[dict[str, Any]]:
    rows = db.execute(
        """
        SELECT * FROM employee_qualifications
        WHERE staff_id=? AND COALESCE(is_deleted,0)=0 ORDER BY year_completed DESC, id
        """,
        (staff_id,),
    ).fetchall()
    return [dict(r) for r in rows]


def list_employee_experience(db, staff_id: int) -> list[dict[str, Any]]:
    rows = db.execute(
        """
        SELECT * FROM employee_experience
        WHERE staff_id=? AND COALESCE(is_deleted,0)=0 ORDER BY from_date DESC, id
        """,
        (staff_id,),
    ).fetchall()
    return [dict(r) for r in rows]


def list_employee_skills(db, staff_id: int) -> list[dict[str, Any]]:
    rows = db.execute(
        """
        SELECT * FROM employee_skills
        WHERE staff_id=? AND COALESCE(is_deleted,0)=0 ORDER BY skill_name, id
        """,
        (staff_id,),
    ).fetchall()
    return [dict(r) for r in rows]


def list_employee_employment_history(db, staff_id: int) -> list[dict[str, Any]]:
    rows = db.execute(
        """
        SELECT h.*, d.department_name, des.designation_name, b.branch_name,
               rm.staff_name AS reporting_manager_name
        FROM employee_employment_history h
        LEFT JOIN departments d ON h.department_id = d.id
        LEFT JOIN designations des ON h.designation_id = des.id
        LEFT JOIN company_branches b ON h.branch_id = b.id
        LEFT JOIN staff rm ON h.reporting_manager_id = rm.id
        WHERE h.staff_id=? ORDER BY h.effective_date DESC, h.id DESC
        """,
        (staff_id,),
    ).fetchall()
    return [dict(r) for r in rows]


def list_active_employees_for_dropdown(db, company_id: int | None = None) -> list[dict[str, Any]]:
    listing = list_employees_master(
        db, company_id=company_id, status="Active", per_page=5000, sort_by="staff_name"
    )
    return [
        {
            "id": r["id"],
            "employee_code": r.get("employee_code"),
            "staff_name": r.get("staff_name"),
            "department_name": r.get("dept_master_name") or r.get("department"),
        }
        for r in listing["items"]
    ]


def _parse_employee_form(form) -> dict[str, Any]:
    def _int(key: str) -> int | None:
        raw = _form_value(form, key)
        return int(raw) if raw not in ("", "0") else None

    first = _form_value(form, "first_name")
    middle = _form_value(form, "middle_name")
    last = _form_value(form, "last_name")
    staff_name = _form_value(form, "staff_name") or " ".join(p for p in (first, middle, last) if p).strip()
    salary_raw = _form_value(form, "salary_amount")
    salary_amount = float(salary_raw) if salary_raw else None
    return {
        "employee_code": (_form_value(form, "employee_code") or "").upper(),
        "staff_name": staff_name,
        "first_name": first,
        "middle_name": middle,
        "last_name": last,
        "employee_type": _form_value(form, "employee_type") or "Permanent",
        "company_id": _int("company_id"),
        "branch_id": _int("branch_id"),
        "department_id": _int("department_id"),
        "designation_id": _int("designation_id"),
        "reporting_manager_id": _int("reporting_manager_id"),
        "official_email": _form_value(form, "official_email"),
        "personal_email": _form_value(form, "personal_email"),
        "mobile": _form_value(form, "mobile"),
        "alternative_mobile": _form_value(form, "alternative_mobile"),
        "current_address": _form_value(form, "current_address"),
        "permanent_address": _form_value(form, "permanent_address"),
        "city": _form_value(form, "city"),
        "state": _form_value(form, "state"),
        "pin_code": _form_value(form, "pin_code"),
        "country": _form_value(form, "country") or "India",
        "gender": _form_value(form, "gender"),
        "date_of_birth": _form_value(form, "date_of_birth"),
        "blood_group": _form_value(form, "blood_group"),
        "nationality": _form_value(form, "nationality") or "Indian",
        "marital_status": _form_value(form, "marital_status"),
        "joining_date": _form_value(form, "joining_date"),
        "confirmation_date": _form_value(form, "confirmation_date"),
        "relieving_date": _form_value(form, "relieving_date"),
        "employment_status": _form_value(form, "employment_status") or "Active",
        "aadhaar_number": re.sub(r"\s", "", _form_value(form, "aadhaar_number")),
        "pan_number": _form_value(form, "pan_number").upper(),
        "passport_number": _form_value(form, "passport_number"),
        "driving_license": _form_value(form, "driving_license"),
        "pf_number": _form_value(form, "pf_number"),
        "esi_number": _form_value(form, "esi_number"),
        "uan_number": _form_value(form, "uan_number"),
        "bank_name": _form_value(form, "bank_name"),
        "bank_account": _form_value(form, "bank_account"),
        "ifsc_code": _form_value(form, "ifsc_code").upper(),
        "branch_name": _form_value(form, "bank_branch_name") or _form_value(form, "branch_name"),
        "salary_type": _form_value(form, "salary_type"),
        "salary_amount": salary_amount,
        "salary_structure": _form_value(form, "salary_structure"),
        "ot_applicable": _form_value(form, "ot_applicable") or "No",
        "working_hours": _form_value(form, "working_hours"),
        "ot_rate_per_hour": _form_value(form, "ot_rate_per_hour"),
        "holiday_pay_applicable": _form_value(form, "holiday_pay_applicable") or "No",
        "workflow_role": _form_value(form, "workflow_role"),
        "company_room_provided": _form_value(form, "company_room_provided") or "No",
        "company_food_provided": _form_value(form, "company_food_provided") or "No",
        "profile_photo": _form_value(form, "profile_photo"),
        "digital_signature": _form_value(form, "digital_signature"),
        "id_proof": _form_value(form, "id_proof"),
        "aadhaar_document": _form_value(form, "aadhaar_document"),
        "pan_document": _form_value(form, "pan_document"),
        "status": _form_value(form, "status") or "Active",
        "record_employment_history": _form_value(form, "record_employment_history") == "1",
        "history_remarks": _form_value(form, "history_remarks"),
    }


def _resolve_org_labels(db, data: dict[str, Any]) -> tuple[str, str, str]:
    dept_name = ""
    desig_name = ""
    manager_name = ""
    if data.get("department_id"):
        dept = get_department_master(db, int(data["department_id"]))
        dept_name = (dept or {}).get("department_name") or ""
    if data.get("designation_id"):
        des = get_designation_master(db, int(data["designation_id"]))
        desig_name = (des or {}).get("designation_name") or ""
    if data.get("reporting_manager_id"):
        mgr = get_employee_master(db, int(data["reporting_manager_id"]))
        manager_name = (mgr or {}).get("staff_name") or ""
    return dept_name, desig_name, manager_name


def _validate_org_assignment(db, data: dict[str, Any], staff_id: int | None = None) -> None:
    company_id = data.get("company_id")
    if not company_id:
        raise ValueError("Company is required.")
    if not get_company(db, int(company_id)):
        raise ValueError("Selected company was not found.")
    if data.get("branch_id"):
        branch = get_branch_master(db, int(data["branch_id"]))
        if not branch:
            raise ValueError("Selected branch was not found.")
        if int(branch.get("company_id") or 0) != int(company_id):
            raise ValueError("Branch must belong to the selected company.")
    if not data.get("department_id"):
        raise ValueError("Department is required.")
    dept = get_department_master(db, int(data["department_id"]))
    if not dept:
        raise ValueError("Selected department was not found.")
    if int(dept.get("company_id") or 0) != int(company_id):
        raise ValueError("Department must belong to the selected company.")
    if not data.get("designation_id"):
        raise ValueError("Designation is required.")
    des = get_designation_master(db, int(data["designation_id"]))
    if not des:
        raise ValueError("Selected designation was not found.")
    mgr_id = data.get("reporting_manager_id")
    if mgr_id:
        if staff_id and int(mgr_id) == int(staff_id):
            raise ValueError("Reporting manager cannot be the same employee.")
        mgr = get_employee_master(db, int(mgr_id))
        if not mgr:
            raise ValueError("Reporting manager was not found.")
        if int(mgr.get("company_id") or 0) != int(company_id):
            raise ValueError("Reporting manager must belong to the same company.")


def _save_child_rows(
    db,
    staff_id: int,
    form,
    username: str,
    *,
    table: str,
    id_key: str,
    fields: tuple[str, ...],
    required_field: str | None = None,
) -> None:
    now = _now_ts()
    ids = _form_getlist(form, f"{id_key}[]")
    seen: set[int] = set()
    lists = {f: _form_getlist(form, f"{f}[]") for f in fields}
    max_len = max([len(ids)] + [len(v) for v in lists.values()] + [0])
    for idx in range(max_len):
        values = {f: (lists[f][idx] if idx < len(lists[f]) else "") for f in fields}
        if required_field and not values.get(required_field):
            continue
        if not any(values.values()):
            continue
        rid_raw = ids[idx] if idx < len(ids) else ""
        if rid_raw.isdigit():
            rid = int(rid_raw)
            seen.add(rid)
            sets = ", ".join(f"{f}=?" for f in fields)
            db.execute(
                f"UPDATE {table} SET {sets}, modified_by=?, modified_at=? WHERE id=? AND staff_id=?",
                (*[values[f] for f in fields], username, now, rid, staff_id),
            )
        else:
            cols = "staff_id, " + ", ".join(fields) + ", created_by, created_at, modified_by, modified_at"
            placeholders = ",".join(["?"] * (len(fields) + 5))
            cur = db.execute(
                f"INSERT INTO {table}({cols}) VALUES({placeholders})",
                (staff_id, *[values[f] for f in fields], username, now, username, now),
            )
            seen.add(int(cur.lastrowid))
    existing = db.execute(f"SELECT id FROM {table} WHERE staff_id=? AND COALESCE(is_deleted,0)=0", (staff_id,)).fetchall()
    for row in existing:
        rid = int(row[0])
        if max_len == 0:
            continue
        if rid not in seen:
            db.execute(
                f"UPDATE {table} SET is_deleted=1, modified_by=?, modified_at=? WHERE id=?",
                (username, now, rid),
            )


def _maybe_record_employment_history(
    db,
    staff_id: int,
    before: dict[str, Any] | None,
    after: dict[str, Any],
    username: str,
    *,
    force: bool = False,
    remarks: str = "",
) -> None:
    if not before and not force:
        return
    changed = force or any(
        str(before.get(k) or "") != str(after.get(k) or "")
        for k in ("department_id", "designation_id", "branch_id", "reporting_manager_id")
    )
    if not changed:
        return
    now = _now_ts()
    db.execute(
        """
        INSERT INTO employee_employment_history(
            staff_id, effective_date, change_type, department_id, designation_id,
            branch_id, reporting_manager_id, remarks, created_by, created_at
        ) VALUES(?,?,?,?,?,?,?,?,?,?)
        """,
        (
            staff_id,
            _now_ts()[:10],
            "Transfer/Promotion" if before else "Initial",
            after.get("department_id"),
            after.get("designation_id"),
            after.get("branch_id"),
            after.get("reporting_manager_id"),
            remarks or "Org assignment updated",
            username,
            now,
        ),
    )


def save_employee_master(
    db,
    form,
    username: str,
    staff_id: int | None = None,
    *,
    customer_id: int | None = None,
) -> int:
    data = _parse_employee_form(form)
    validate_employee_form_data(data, staff_id=staff_id)
    _validate_org_assignment(db, data, staff_id=staff_id)
    if not data["employee_code"]:
        data["employee_code"] = generate_employee_code(db)
    validate_employee_uniqueness(
        db,
        employee_code=data["employee_code"],
        official_email=data["official_email"],
        staff_id=staff_id,
    )
    dept_name, desig_name, manager_name = _resolve_org_labels(db, data)
    data["department"] = dept_name
    data["designation"] = desig_name
    data["reporting_manager"] = manager_name
    data["email"] = data["official_email"] or data.get("personal_email") or ""
    working_hours = float(data["working_hours"]) if data.get("working_hours") else None
    ot_rate = float(data["ot_rate_per_hour"]) if data.get("ot_rate_per_hour") else 0.0
    photo_path = data.get("profile_photo") or data.get("photo") or ""
    signature_path = data.get("digital_signature") or ""
    id_proof = data.get("id_proof") or ""
    aadhaar_doc = data.get("aadhaar_document") or ""
    pan_doc = data.get("pan_document") or ""
    now = _now_ts()
    core = (
        data["employee_code"],
        data["staff_name"],
        data["first_name"],
        data["middle_name"],
        data["last_name"],
        data["employee_type"],
        data["company_id"],
        data["branch_id"],
        data["department_id"],
        data["designation_id"],
        data["reporting_manager_id"],
        data["official_email"],
        data["personal_email"],
        data["mobile"],
        data["alternative_mobile"],
        data["current_address"],
        data["permanent_address"],
        data["city"],
        data["state"],
        data["pin_code"],
        data["country"],
        data["gender"],
        data["date_of_birth"],
        data["blood_group"],
        data["nationality"],
        data["marital_status"],
        data["joining_date"],
        data["confirmation_date"],
        data["relieving_date"],
        data["employment_status"],
        data["aadhaar_number"],
        data["pan_number"],
        data["passport_number"],
        data["driving_license"],
        data["pf_number"],
        data["esi_number"],
        data["uan_number"],
        data["bank_name"],
        data["bank_account"],
        data["ifsc_code"],
        data["branch_name"],
        data["salary_type"],
        data["salary_amount"],
        data["salary_structure"],
        data["ot_applicable"],
        working_hours,
        ot_rate,
        data["holiday_pay_applicable"],
        data["workflow_role"],
        data["company_room_provided"],
        data["company_food_provided"],
        dept_name,
        desig_name,
        manager_name,
        data["email"],
        photo_path,
        photo_path,
        signature_path,
        id_proof,
        aadhaar_doc,
        pan_doc,
        data["status"],
    )
    if staff_id:
        existing = get_employee_master(db, staff_id, include_deleted=True)
        if not existing:
            raise ValueError("Employee not found.")
        if not photo_path and existing.get("profile_photo"):
            core_list = list(core)
            core_list[-7] = existing["profile_photo"]
            core_list[-6] = existing["profile_photo"]
            core = tuple(core_list)
        if not signature_path and existing.get("digital_signature"):
            core_list = list(core)
            core_list[-5] = existing["digital_signature"]
            core = tuple(core_list)
        if not id_proof and existing.get("id_proof"):
            core_list = list(core)
            core_list[-4] = existing["id_proof"]
            core = tuple(core_list)
        if not aadhaar_doc and existing.get("aadhaar_document"):
            core_list = list(core)
            core_list[-3] = existing["aadhaar_document"]
            core = tuple(core_list)
        if not pan_doc and existing.get("pan_document"):
            core_list = list(core)
            core_list[-2] = existing["pan_document"]
            core = tuple(core_list)
        db.execute(
            """
            UPDATE staff SET employee_code=?, staff_name=?, first_name=?, middle_name=?, last_name=?,
            employee_type=?, company_id=?, branch_id=?, department_id=?, designation_id=?,
            reporting_manager_id=?, official_email=?, personal_email=?, mobile=?, alternative_mobile=?,
            current_address=?, permanent_address=?, city=?, state=?, pin_code=?, country=?,
            gender=?, date_of_birth=?, blood_group=?, nationality=?, marital_status=?,
            joining_date=?, confirmation_date=?, relieving_date=?, employment_status=?,
            aadhaar_number=?, pan_number=?, passport_number=?, driving_license=?,
            pf_number=?, esi_number=?, uan_number=?, bank_name=?, bank_account=?, ifsc_code=?,
            branch_name=?, salary_type=?, salary_amount=?, salary_structure=?,
            ot_applicable=?, working_hours=?, ot_rate_per_hour=?, holiday_pay_applicable=?,
            workflow_role=?, company_room_provided=?, company_food_provided=?,
            department=?, designation=?, reporting_manager=?, email=?,
            photo=?, profile_photo=?, digital_signature=?, id_proof=?, aadhaar_document=?, pan_document=?,
            status=?, modified_by=?, modified_at=? WHERE id=?
            """,
            (*core, username, now, staff_id),
        )
        if customer_id is not None:
            db.execute("UPDATE staff SET customer_id=? WHERE id=?", (customer_id, staff_id))
        _save_child_rows(
            db,
            staff_id,
            form,
            username,
            table="employee_emergency_contacts",
            id_key="ec_id",
            fields=("contact_name", "relationship", "phone", "email", "address"),
            required_field="contact_name",
        )
        _save_child_rows(
            db,
            staff_id,
            form,
            username,
            table="employee_qualifications",
            id_key="qual_id",
            fields=("qualification", "institution", "board_university", "year_completed", "grade"),
            required_field="qualification",
        )
        _save_child_rows(
            db,
            staff_id,
            form,
            username,
            table="employee_experience",
            id_key="exp_id",
            fields=("company_name", "designation", "from_date", "to_date", "responsibilities"),
            required_field="company_name",
        )
        _save_child_rows(
            db,
            staff_id,
            form,
            username,
            table="employee_skills",
            id_key="skill_id",
            fields=("skill_name", "proficiency", "years_experience", "certification"),
            required_field="skill_name",
        )
        after = get_employee_master(db, staff_id, include_deleted=True)
        _maybe_record_employment_history(
            db,
            staff_id,
            existing,
            after or {},
            username,
            force=data.get("record_employment_history"),
            remarks=data.get("history_remarks") or "",
        )
        log_employee_field_changes(db, existing, after, username)
        return staff_id

    approval_status = _form_value(form, "approval_status") or "Draft"
    if approval_status not in APPROVAL_STATUSES:
        approval_status = "Draft"
    insert_cols = (
        "employee_code, staff_name, first_name, middle_name, last_name, employee_type, company_id, branch_id, "
        "department_id, designation_id, reporting_manager_id, official_email, personal_email, mobile, "
        "alternative_mobile, current_address, permanent_address, city, state, pin_code, country, gender, "
        "date_of_birth, blood_group, nationality, marital_status, joining_date, confirmation_date, relieving_date, "
        "employment_status, aadhaar_number, pan_number, passport_number, driving_license, pf_number, esi_number, "
        "uan_number, bank_name, bank_account, ifsc_code, branch_name, salary_type, salary_amount, salary_structure, "
        "ot_applicable, working_hours, ot_rate_per_hour, holiday_pay_applicable, workflow_role, company_room_provided, "
        "company_food_provided, department, designation, reporting_manager, email, photo, profile_photo, "
        "digital_signature, id_proof, aadhaar_document, pan_document, status, approval_status, created_by, created_at, "
        "modified_by, modified_at"
    )
    placeholders = ",".join(["?"] * 67)
    vals = (*core, approval_status, username, now, username, now)
    if customer_id is not None:
        cur = db.execute(
            f"INSERT INTO staff({insert_cols}, customer_id) VALUES({placeholders},?)",
            (*vals, customer_id),
        )
    else:
        cur = db.execute(f"INSERT INTO staff({insert_cols}) VALUES({placeholders})", vals)
    new_id = int(cur.lastrowid)
    _save_child_rows(
        db,
        new_id,
        form,
        username,
        table="employee_emergency_contacts",
        id_key="ec_id",
        fields=("contact_name", "relationship", "phone", "email", "address"),
        required_field="contact_name",
    )
    _save_child_rows(
        db,
        new_id,
        form,
        username,
        table="employee_qualifications",
        id_key="qual_id",
        fields=("qualification", "institution", "board_university", "year_completed", "grade"),
        required_field="qualification",
    )
    _save_child_rows(
        db,
        new_id,
        form,
        username,
        table="employee_experience",
        id_key="exp_id",
        fields=("company_name", "designation", "from_date", "to_date", "responsibilities"),
        required_field="company_name",
    )
    _save_child_rows(
        db,
        new_id,
        form,
        username,
        table="employee_skills",
        id_key="skill_id",
        fields=("skill_name", "proficiency", "years_experience", "certification"),
        required_field="skill_name",
    )
    log_employee_audit(
        db,
        new_id,
        "create",
        username,
        remarks=f"Created employee {data['employee_code']}",
    )
    return new_id


def apply_employee_upload_paths(db, staff_id: int, *, profile_photo: str = "", digital_signature: str = "", username: str = "") -> None:
    if not staff_id:
        return
    now = _now_ts()
    sets: list[str] = []
    params: list[Any] = []
    if profile_photo:
        sets.extend(["profile_photo=?", "photo=?"])
        params.extend([profile_photo, profile_photo])
    if digital_signature:
        sets.append("digital_signature=?")
        params.append(digital_signature)
    if not sets:
        return
    sets.extend(["modified_by=?", "modified_at=?"])
    params.extend([username, now, staff_id])
    db.execute(f"UPDATE staff SET {', '.join(sets)} WHERE id=?", params)


def soft_delete_employee_master(db, staff_id: int, username: str) -> None:
    if not staff_id:
        raise ValueError("Invalid employee.")
    row = get_employee_master(db, staff_id, include_deleted=True)
    if not row:
        raise ValueError("Employee not found.")
    if row.get("is_deleted"):
        return
    blockers = employee_has_references(db, staff_id)
    if blockers:
        raise ValueError(
            "Employee cannot be deleted because linked records exist in: "
            + ", ".join(blockers)
            + ". Deactivate instead."
        )
    now = _now_ts()
    db.execute(
        """
        UPDATE staff SET is_deleted=1, deleted_by=?, deleted_at=?,
        modified_by=?, modified_at=? WHERE id=?
        """,
        (username, now, username, now, staff_id),
    )
    log_employee_audit(
        db,
        staff_id,
        "soft_delete",
        username,
        remarks=f"Soft-deleted employee {row.get('employee_code')}",
    )


def activate_employee_master(db, staff_id: int, username: str) -> None:
    if not get_employee_master(db, staff_id, include_deleted=True):
        raise ValueError("Employee not found.")
    now = _now_ts()
    db.execute(
        "UPDATE staff SET status='Active', employment_status='Active', modified_by=?, modified_at=? WHERE id=?",
        (username, now, staff_id),
    )
    log_employee_audit(db, staff_id, "activate", username)


def deactivate_employee_master(db, staff_id: int, username: str) -> None:
    if not get_employee_master(db, staff_id, include_deleted=True):
        raise ValueError("Employee not found.")
    now = _now_ts()
    db.execute(
        "UPDATE staff SET status='Inactive', modified_by=?, modified_at=? WHERE id=?",
        (username, now, staff_id),
    )
    log_employee_audit(db, staff_id, "deactivate", username)


def approve_employee_master(db, staff_id: int, username: str) -> None:
    if not get_employee_master(db, staff_id):
        raise ValueError("Employee not found.")
    now = _now_ts()
    db.execute(
        "UPDATE staff SET approval_status='Approved', approved_by=?, approved_at=?, modified_by=?, modified_at=? WHERE id=?",
        (username, now, username, now, staff_id),
    )
    log_employee_audit(db, staff_id, "approve", username)


def reject_employee_master(db, staff_id: int, username: str, remarks: str = "") -> None:
    if not get_employee_master(db, staff_id):
        raise ValueError("Employee not found.")
    now = _now_ts()
    db.execute(
        "UPDATE staff SET approval_status='Rejected', modified_by=?, modified_at=? WHERE id=?",
        (username, now, staff_id),
    )
    log_employee_audit(db, staff_id, "reject", username, remarks=remarks or "Employee rejected")


def user_can_employee_master(
    db,
    user_id: int | None,
    action: str,
    *,
    is_admin: bool = False,
) -> bool:
    if is_admin:
        return True
    if not user_id:
        return False
    action_map = {"deactivate": "edit", "activate": "edit"}
    check = action_map.get(action, action)
    try:
        from user_permission_service import empty_permission_actions, normalize_permission_actions

        row = db.execute(
            """
            SELECT granted, action_flags FROM user_tab_permissions
            WHERE user_id=? AND granted=1 AND endpoint='employee_master'
            LIMIT 1
            """,
            (user_id,),
        ).fetchone()
        if not row:
            return False
        raw_flags = row["action_flags"] if hasattr(row, "keys") else row[1]
        actions = normalize_permission_actions(
            json.loads(raw_flags) if raw_flags else empty_permission_actions()
        )
        if check == "import":
            return bool(actions.get("import") or actions.get("create"))
        if check == "delete":
            return bool(actions.get("delete") or actions.get("edit"))
        return bool(actions.get(check))
    except Exception:
        return False


def employees_for_export(db, *, include_deleted: bool = False, **filters) -> list[dict[str, Any]]:
    listing = list_employees_master(db, include_deleted=include_deleted, per_page=10000, **filters)
    rows: list[dict[str, Any]] = []
    for item in listing["items"]:
        row = {col: item.get(col, "") for col in EMPLOYEE_EXPORT_COLUMNS}
        row["department_name"] = item.get("dept_master_name") or item.get("department") or ""
        row["designation_name"] = item.get("desig_master_name") or item.get("designation") or ""
        row["company_code"] = item.get("company_code") or ""
        row["branch_code"] = item.get("branch_code") or ""
        row["reporting_manager_code"] = item.get("reporting_manager_code") or ""
        rows.append(row)
    return rows


def export_employees_excel(db, **filters) -> BytesIO:
    from openpyxl import Workbook

    rows = employees_for_export(db, **filters)
    wb = Workbook()
    ws = wb.active
    ws.title = "Employees"
    headers = list(EMPLOYEE_EXPORT_COLUMNS)
    ws.append(headers)
    for row in rows:
        ws.append([row.get(h, "") for h in headers])
    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def export_employees_csv(db, **filters) -> str:
    rows = employees_for_export(db, **filters)
    si = io.StringIO()
    writer = csv.writer(si)
    headers = list(EMPLOYEE_EXPORT_COLUMNS)
    writer.writerow(headers)
    for row in rows:
        writer.writerow([row.get(h, "") for h in headers])
    return si.getvalue()


def export_employees_pdf(db, *, report_title: str = "Employee Master Report", **filters) -> BytesIO:
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.pdfgen import canvas

    rows = employees_for_export(db, **filters)
    buf = BytesIO()
    page_size = landscape(A4)
    c = canvas.Canvas(buf, pagesize=page_size)
    _, height = page_size
    y = height - 40
    c.setFont("Helvetica-Bold", 14)
    c.drawString(40, y, f"MAXEK ERP — {report_title}")
    y -= 24
    c.setFont("Helvetica", 9)
    for row in rows[:250]:
        line = (
            f"{row.get('employee_code')} | {row.get('staff_name')} | "
            f"{row.get('department_name') or '—'} | {row.get('designation_name') or '—'} | "
            f"{row.get('status')}"
        )
        if y < 40:
            c.showPage()
            y = height - 40
            c.setFont("Helvetica", 9)
        c.drawString(40, y, line[:130])
        y -= 14
    c.save()
    buf.seek(0)
    return buf


def employee_report(db, report_key: str, **filters) -> list[dict[str, Any]]:
    key = (report_key or "directory").lower().strip()
    if key == "active":
        filters["status"] = "Active"
    elif key == "inactive":
        filters["status"] = "Inactive"
    elif key == "joining":
        listing = list_employees_master(db, per_page=5000, **filters)
        return sorted(listing["items"], key=lambda r: r.get("joining_date") or "")
    elif key == "relieving":
        listing = list_employees_master(db, per_page=5000, employment_status="Relieved", **filters)
        return sorted(listing["items"], key=lambda r: r.get("relieving_date") or "", reverse=True)
    elif key == "department_wise":
        listing = list_employees_master(db, per_page=5000, **filters)
        return sorted(
            listing["items"],
            key=lambda r: (r.get("dept_master_name") or r.get("department") or "", r.get("staff_name") or ""),
        )
    elif key == "designation_wise":
        listing = list_employees_master(db, per_page=5000, **filters)
        return sorted(
            listing["items"],
            key=lambda r: (r.get("desig_master_name") or r.get("designation") or "", r.get("staff_name") or ""),
        )
    listing = list_employees_master(db, per_page=5000, **filters)
    return listing["items"]


def employee_import_template() -> BytesIO:
    from bulk_import_service import build_xlsx_template

    return build_xlsx_template(
        [
            "Employee Code",
            "First Name",
            "Last Name",
            "Employee Type",
            "Company Code",
            "Branch Code",
            "Department Code",
            "Designation Code",
            "Reporting Manager Code",
            "Official Email",
            "Mobile",
            "Joining Date",
            "PAN",
            "Aadhaar",
            "PF Number",
            "ESI Number",
            "UAN",
            "Bank Name",
            "Account Number",
            "IFSC",
            "Salary Type",
            "Salary Amount",
            "Status",
        ],
        [
            "EMP101",
            "Rajesh",
            "Kumar",
            "Permanent",
            "CO001",
            "BR001",
            "DEPT-HO",
            "DES-ENG",
            "",
            "rajesh@company.in",
            "9876543210",
            "2026-01-01",
            "ABCDE1234F",
            "123456789012",
            "",
            "",
            "",
            "HDFC Bank",
            "1234567890",
            "HDFC0001234",
            "Monthly",
            "50000",
            "Active",
        ],
    )


def ai_validate_employee(
    db,
    staff_id: int | None = None,
    form: dict | None = None,
) -> dict[str, Any]:
    data = dict(form or {})
    if staff_id and not form:
        row = get_employee_master(db, staff_id)
        if row:
            data = dict(row)
    issues: list[str] = []
    suggestions: list[str] = []
    duplicates: list[str] = []
    missing: list[str] = []
    code = (data.get("employee_code") or "").strip()
    email = (data.get("official_email") or data.get("email") or "").strip()
    pan = (data.get("pan_number") or "").strip()
    aadhaar = (data.get("aadhaar_number") or "").strip()
    if not (data.get("first_name") or data.get("staff_name")):
        issues.append("Employee name is missing.")
        missing.append("first_name")
    if not data.get("department_id") and not data.get("department"):
        issues.append("Department is not assigned.")
        missing.append("department_id")
    if not data.get("designation_id") and not data.get("designation"):
        issues.append("Designation is not assigned.")
        missing.append("designation_id")
    if not data.get("joining_date"):
        issues.append("Joining date is missing.")
        missing.append("joining_date")
    if not data.get("profile_photo") and not data.get("photo"):
        missing.append("profile_photo")
        suggestions.append("Upload a profile photo for HR records.")
    if not data.get("aadhaar_document"):
        missing.append("aadhaar_document")
    if not data.get("pan_document") and pan:
        missing.append("pan_document")
    if code:
        hit = db.execute(
            "SELECT id, staff_name FROM staff WHERE UPPER(employee_code)=? AND COALESCE(is_deleted,0)=0",
            (code.upper(),),
        ).fetchone()
        if hit and (not staff_id or int(hit[0]) != int(staff_id)):
            duplicates.append(f"Employee code {code} already used by {hit[1]}.")
    if email:
        hit = db.execute(
            """
            SELECT id, staff_name FROM staff
            WHERE LOWER(COALESCE(official_email, email))=? AND COALESCE(is_deleted,0)=0
            """,
            (email.lower(),),
        ).fetchone()
        if hit and (not staff_id or int(hit[0]) != int(staff_id)):
            duplicates.append(f"Email already registered to {hit[1]}.")
    if pan:
        hit = db.execute(
            "SELECT id, staff_name FROM staff WHERE UPPER(pan_number)=? AND COALESCE(is_deleted,0)=0",
            (pan.upper(),),
        ).fetchone()
        if hit and (not staff_id or int(hit[0]) != int(staff_id)):
            duplicates.append(f"PAN already registered to {hit[1]}.")
    try:
        if pan:
            validate_pan_number(pan)
        if aadhaar:
            validate_aadhaar_number(aadhaar)
        if email:
            validate_email(email)
    except ValueError as exc:
        issues.append(str(exc))
    expiring: list[str] = []
    passport = (data.get("passport_number") or "").strip()
    dl = (data.get("driving_license") or "").strip()
    if passport and not data.get("passport_expiry"):
        suggestions.append("Passport expiry date not tracked — add when available.")
    if dl:
        suggestions.append("Verify driving licence validity periodically.")
    return {
        "ok": not issues and not duplicates,
        "issues": issues,
        "suggestions": suggestions,
        "duplicates": duplicates,
        "missing_documents": missing,
        "expiring_ids": expiring,
        "ai_available": False,
        "message": "Rule-based validation (AI service not configured).",
    }
