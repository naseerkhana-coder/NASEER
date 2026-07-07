"""Client Master (MODULE-011) — clients linked to projects, DMS, and billing."""

from __future__ import annotations

import csv
import io
import json
import re
from io import BytesIO
from typing import Any

from company_master_service import (
    COMPANY_COUNTRIES,
    _ensure_column,
    _now_ts,
    _table_exists,
    ensure_company_master_schema,
    get_company,
    list_companies,
    validate_email,
    validate_gst_number,
    validate_pan_number,
    validate_phone,
)

CLIENT_STATUSES = ("Active", "Inactive")
CLIENT_TYPES = ("Corporate", "Government", "PSU", "Individual", "NGO", "Other")
ADDRESS_TYPES = ("Billing", "Site", "Other")
APPROVAL_STATUSES = ("Draft", "Pending", "Approved", "Rejected")
CLIENT_SORT_COLUMNS = (
    "client_code",
    "client_name",
    "company_name",
    "legal_name",
    "client_type",
    "industry",
    "city",
    "status",
    "created_at",
    "company_id",
)
CLIENT_EXPORT_COLUMNS = (
    "client_code",
    "client_name",
    "legal_name",
    "company_name",
    "client_type",
    "industry",
    "gst_number",
    "pan_number",
    "registration_number",
    "email",
    "phone",
    "mobile",
    "website",
    "billing_address",
    "site_address",
    "country",
    "state",
    "district",
    "city",
    "pin_code",
    "primary_contact_name",
    "primary_contact_email",
    "primary_contact_mobile",
    "payment_terms",
    "credit_limit",
    "bank_name",
    "account_number",
    "ifsc_swift",
    "status",
    "approval_status",
    "company_code",
)
CLIENT_AUDIT_FIELDS = (
    "client_code",
    "client_name",
    "legal_name",
    "company_name",
    "client_type",
    "industry",
    "gst_number",
    "pan_number",
    "registration_number",
    "email",
    "phone",
    "mobile",
    "website",
    "billing_address",
    "site_address",
    "address",
    "country",
    "state",
    "district",
    "city",
    "pin_code",
    "contact_person",
    "payment_terms",
    "credit_limit",
    "bank_name",
    "account_number",
    "ifsc_swift",
    "status",
    "approval_status",
    "company_id",
)
TERMINAL_PROJECT_STATUSES = frozenset(
    {"COMPLETED", "CANCELLED", "CLOSED", "INACTIVE", "ARCHIVED"}
)
CLIENT_INDUSTRIES = (
    "Construction",
    "Infrastructure",
    "Real Estate",
    "Government",
    "Manufacturing",
    "Oil & Gas",
    "Power",
    "Consultancy",
    "Other",
)


def ensure_client_master_schema(db) -> None:
    """Extend clients table and child tables for MODULE-011 (idempotent)."""
    ensure_company_master_schema(db)
    db.execute(
        """
        CREATE TABLE IF NOT EXISTS clients(
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            client_name TEXT,
            company_name TEXT,
            mobile TEXT,
            email TEXT,
            address TEXT,
            gst_number TEXT,
            status TEXT
        )
        """
    )
    for col, ctype in (
        ("client_code", "TEXT"),
        ("legal_name", "TEXT"),
        ("client_type", "TEXT"),
        ("industry", "TEXT"),
        ("pan_number", "TEXT"),
        ("registration_number", "TEXT"),
        ("phone", "TEXT"),
        ("website", "TEXT"),
        ("billing_address", "TEXT"),
        ("site_address", "TEXT"),
        ("country", "TEXT"),
        ("state", "TEXT"),
        ("district", "TEXT"),
        ("city", "TEXT"),
        ("pin_code", "TEXT"),
        ("contact_person", "TEXT"),
        ("payment_terms", "TEXT"),
        ("credit_limit", "REAL"),
        ("bank_name", "TEXT"),
        ("account_number", "TEXT"),
        ("ifsc_swift", "TEXT"),
        ("company_id", "INTEGER"),
        ("approval_status", "TEXT DEFAULT 'Draft'"),
        ("approved_by", "TEXT"),
        ("approved_at", "TEXT"),
        ("created_by", "TEXT"),
        ("created_at", "TEXT"),
        ("modified_by", "TEXT"),
        ("modified_at", "TEXT"),
        ("is_deleted", "INTEGER DEFAULT 0"),
        ("deleted_by", "TEXT"),
        ("deleted_at", "TEXT"),
        ("customer_id", "INTEGER"),
    ):
        _ensure_column(db, "clients", col, ctype)
    db.execute(
        """
        CREATE TABLE IF NOT EXISTS client_contacts(
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            client_id INTEGER NOT NULL,
            contact_name TEXT,
            designation TEXT,
            email TEXT,
            mobile TEXT,
            is_primary INTEGER DEFAULT 0,
            created_by TEXT,
            created_at TEXT,
            modified_by TEXT,
            modified_at TEXT,
            is_deleted INTEGER DEFAULT 0,
            FOREIGN KEY(client_id) REFERENCES clients(id)
        )
        """
    )
    db.execute(
        """
        CREATE TABLE IF NOT EXISTS client_addresses(
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            client_id INTEGER NOT NULL,
            address_type TEXT,
            address_line1 TEXT,
            address_line2 TEXT,
            city TEXT,
            district TEXT,
            state TEXT,
            country TEXT,
            pin_code TEXT,
            is_primary INTEGER DEFAULT 0,
            created_by TEXT,
            created_at TEXT,
            modified_by TEXT,
            modified_at TEXT,
            is_deleted INTEGER DEFAULT 0,
            FOREIGN KEY(client_id) REFERENCES clients(id)
        )
        """
    )
    db.execute(
        """
        CREATE TABLE IF NOT EXISTS client_bank_details(
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            client_id INTEGER NOT NULL,
            bank_name TEXT,
            account_number TEXT,
            ifsc_swift TEXT,
            branch_name TEXT,
            account_holder TEXT,
            is_primary INTEGER DEFAULT 0,
            created_by TEXT,
            created_at TEXT,
            modified_by TEXT,
            modified_at TEXT,
            is_deleted INTEGER DEFAULT 0,
            FOREIGN KEY(client_id) REFERENCES clients(id)
        )
        """
    )
    db.execute("UPDATE clients SET status='Active' WHERE status IS NULL OR TRIM(status)=''")
    _migrate_legacy_clients(db)
    try:
        from audit_trail_service import ensure_audit_schema

        ensure_audit_schema(db)
    except Exception:
        pass


def _migrate_legacy_clients(db) -> None:
    if not _table_exists(db, "clients"):
        return
    rows = db.execute(
        "SELECT id, client_code, client_name, company_name, contact_person, mobile, email, address "
        "FROM clients WHERE COALESCE(is_deleted,0)=0"
    ).fetchall()
    for row in rows:
        cid = int(row[0])
        code = (row[1] or "").strip() if len(row) > 1 else ""
        if not code:
            code = generate_client_code(db)
            db.execute("UPDATE clients SET client_code=? WHERE id=?", (code, cid))
        legal = (row[3] or row[2] or "").strip() if len(row) > 3 else ""
        if legal:
            db.execute(
                "UPDATE clients SET legal_name=COALESCE(NULLIF(TRIM(legal_name),''), ?) WHERE id=?",
                (legal, cid),
            )
        contact_count = db.execute(
            "SELECT COUNT(*) FROM client_contacts WHERE client_id=? AND COALESCE(is_deleted,0)=0",
            (cid,),
        ).fetchone()[0]
        if int(contact_count or 0) == 0:
            cp = (row[4] or "").strip() if len(row) > 4 else ""
            mob = (row[5] or "").strip() if len(row) > 5 else ""
            em = (row[6] or "").strip() if len(row) > 6 else ""
            if cp or mob or em:
                now = _now_ts()
                db.execute(
                    """
                    INSERT INTO client_contacts(client_id, contact_name, email, mobile, is_primary,
                    created_by, created_at, modified_by, modified_at)
                    VALUES(?,?,?,?,1,'migration',?,?,?)
                    """,
                    (cid, cp or "Primary Contact", em, mob, now, "migration", now),
                )
        addr_count = db.execute(
            "SELECT COUNT(*) FROM client_addresses WHERE client_id=? AND COALESCE(is_deleted,0)=0",
            (cid,),
        ).fetchone()[0]
        addr = (row[7] or "").strip() if len(row) > 7 else ""
        if int(addr_count or 0) == 0 and addr:
            now = _now_ts()
            db.execute(
                """
                INSERT INTO client_addresses(client_id, address_type, address_line1, is_primary,
                created_by, created_at, modified_by, modified_at)
                VALUES(?, 'Billing', ?, 1, 'migration', ?, 'migration', ?)
                """,
                (cid, addr, now, now),
            )


def generate_client_code(db) -> str:
    rows = db.execute("SELECT client_code FROM clients WHERE client_code LIKE 'CLT%'").fetchall()
    max_code = 100
    for row in rows:
        code = str(row[0] if not hasattr(row, "keys") else row["client_code"] or "").strip().upper()
        number = code[3:]
        if number.isdigit():
            max_code = max(max_code, int(number))
    return f"CLT{max_code + 1}"


def validate_client_uniqueness(
    db,
    *,
    client_code: str,
    gst_number: str = "",
    client_id: int | None = None,
) -> None:
    code = (client_code or "").strip().upper()
    if not code:
        raise ValueError("Client code is required.")
    row = db.execute(
        """
        SELECT id FROM clients
        WHERE UPPER(client_code)=? AND COALESCE(is_deleted,0)=0
        """,
        (code,),
    ).fetchone()
    if row and (not client_id or int(row[0]) != int(client_id)):
        raise ValueError(f"Client code '{code}' already exists.")
    gst = (gst_number or "").strip().upper()
    if gst:
        row = db.execute(
            """
            SELECT id FROM clients
            WHERE UPPER(gst_number)=? AND COALESCE(is_deleted,0)=0
            """,
            (gst,),
        ).fetchone()
        if row and (not client_id or int(row[0]) != int(client_id)):
            raise ValueError(f"GST number '{gst}' is already registered to another client.")


def validate_client_form_data(data: dict[str, Any], *, client_id: int | None = None) -> None:
    name = (data.get("client_name") or "").strip()
    if not name:
        raise ValueError("Client name is required.")
    if data.get("gst_number"):
        validate_gst_number(data["gst_number"])
    if data.get("pan_number"):
        validate_pan_number(data["pan_number"])
    if data.get("email"):
        validate_email(data["email"])
    if data.get("phone"):
        validate_phone(data["phone"])
    if data.get("mobile"):
        validate_phone(data["mobile"])
    status = (data.get("status") or "Active").strip()
    if status not in CLIENT_STATUSES:
        raise ValueError("Select a valid status.")
    client_type = (data.get("client_type") or "").strip()
    if client_type and client_type not in CLIENT_TYPES:
        raise ValueError("Select a valid client type.")


def client_has_active_projects(db, client_id: int) -> bool:
    if not _table_exists(db, "projects"):
        return False
    rows = db.execute(
        "SELECT status FROM projects WHERE client_id=?",
        (client_id,),
    ).fetchall()
    for row in rows:
        status = str(row[0] if not hasattr(row, "keys") else row["status"] or "Active").strip().upper()
        if status not in TERMINAL_PROJECT_STATUSES:
            return True
    return False


def count_client_projects(db, client_id: int) -> int:
    if not _table_exists(db, "projects"):
        return 0
    row = db.execute("SELECT COUNT(*) FROM projects WHERE client_id=?", (client_id,)).fetchone()
    return int(row[0] if row else 0)


def log_client_audit(
    db,
    client_id: int,
    action: str,
    username: str,
    *,
    field_name: str | None = None,
    old_value: str | None = None,
    new_value: str | None = None,
    remarks: str | None = None,
) -> None:
    try:
        from audit_trail_service import log_audit_event

        log_audit_event(
            db,
            record_table="clients",
            record_id=client_id,
            action=action,
            changed_by=username,
            field_name=field_name,
            old_value=old_value,
            new_value=new_value,
            remarks=remarks,
        )
    except Exception:
        pass


def log_client_field_changes(
    db,
    before: dict[str, Any] | None,
    after: dict[str, Any] | None,
    username: str,
) -> None:
    if not before or not after:
        return
    client_id = int(after.get("id") or before.get("id") or 0)
    if not client_id:
        return
    for field in CLIENT_AUDIT_FIELDS:
        old_val = before.get(field)
        new_val = after.get(field)
        if str(old_val or "") != str(new_val or ""):
            log_client_audit(
                db,
                client_id,
                "update",
                username,
                field_name=field,
                old_value=str(old_val or ""),
                new_value=str(new_val or ""),
            )


def list_client_audit_trail(db, client_id: int, limit: int = 100) -> list[dict[str, Any]]:
    try:
        from audit_trail_service import list_audit_trail

        return list_audit_trail(db, "clients", client_id, limit=limit)
    except Exception:
        return []


def _form_getlist(form, key: str) -> list[str]:
    if hasattr(form, "getlist"):
        return [str(v or "").strip() for v in form.getlist(key)]
    val = form.get(key) if hasattr(form, "get") else None
    if isinstance(val, list):
        return [str(v or "").strip() for v in val]
    return [str(val).strip()] if val not in (None, "") else []


def _form_value(form, key: str, default: str = "") -> str:
    if hasattr(form, "get"):
        return str(form.get(key) or default).strip()
    return str((form or {}).get(key) or default).strip()


def _parse_client_form(form) -> dict[str, Any]:
    company_raw = _form_value(form, "company_id")
    company_id = int(company_raw) if company_raw not in ("", "0") else None
    credit_raw = _form_value(form, "credit_limit")
    credit_limit = float(credit_raw) if credit_raw else None
    client_name = _form_value(form, "client_name")
    legal_name = _form_value(form, "legal_name") or _form_value(form, "company_name") or client_name
    company_name = _form_value(form, "company_name") or legal_name
    return {
        "company_id": company_id,
        "client_code": (_form_value(form, "client_code") or "").upper(),
        "client_name": client_name,
        "legal_name": legal_name,
        "company_name": company_name,
        "client_type": _form_value(form, "client_type"),
        "industry": _form_value(form, "industry"),
        "gst_number": _form_value(form, "gst_number").upper(),
        "pan_number": _form_value(form, "pan_number").upper(),
        "registration_number": _form_value(form, "registration_number"),
        "email": _form_value(form, "email"),
        "phone": _form_value(form, "phone"),
        "mobile": _form_value(form, "mobile"),
        "website": _form_value(form, "website"),
        "billing_address": _form_value(form, "billing_address"),
        "site_address": _form_value(form, "site_address"),
        "address": _form_value(form, "address") or _form_value(form, "billing_address"),
        "country": _form_value(form, "country") or "India",
        "state": _form_value(form, "state"),
        "district": _form_value(form, "district"),
        "city": _form_value(form, "city"),
        "pin_code": _form_value(form, "pin_code"),
        "contact_person": _form_value(form, "contact_person"),
        "payment_terms": _form_value(form, "payment_terms"),
        "credit_limit": credit_limit,
        "bank_name": _form_value(form, "bank_name"),
        "account_number": _form_value(form, "account_number"),
        "ifsc_swift": _form_value(form, "ifsc_swift").upper(),
        "status": _form_value(form, "status") or "Active",
    }


def list_client_contacts(db, client_id: int) -> list[dict[str, Any]]:
    rows = db.execute(
        """
        SELECT * FROM client_contacts
        WHERE client_id=? AND COALESCE(is_deleted,0)=0
        ORDER BY is_primary DESC, id
        """,
        (client_id,),
    ).fetchall()
    return [dict(r) for r in rows]


def list_client_addresses(db, client_id: int) -> list[dict[str, Any]]:
    rows = db.execute(
        """
        SELECT * FROM client_addresses
        WHERE client_id=? AND COALESCE(is_deleted,0)=0
        ORDER BY is_primary DESC, id
        """,
        (client_id,),
    ).fetchall()
    return [dict(r) for r in rows]


def list_client_bank_details(db, client_id: int) -> list[dict[str, Any]]:
    rows = db.execute(
        """
        SELECT * FROM client_bank_details
        WHERE client_id=? AND COALESCE(is_deleted,0)=0
        ORDER BY is_primary DESC, id
        """,
        (client_id,),
    ).fetchall()
    return [dict(r) for r in rows]


def _save_client_contacts(db, client_id: int, form, username: str) -> None:
    now = _now_ts()
    ids = _form_getlist(form, "contact_id[]")
    names = _form_getlist(form, "contact_name[]")
    designations = _form_getlist(form, "contact_designation[]")
    emails = _form_getlist(form, "contact_email[]")
    mobiles = _form_getlist(form, "contact_mobile[]")
    primary_raw = _form_value(form, "contact_primary_index")
    primary_idx = int(primary_raw) if primary_raw.isdigit() else 0
    seen_ids: set[int] = set()
    primary_name = ""
    primary_email = ""
    primary_mobile = ""
    max_len = max(len(names), len(designations), len(emails), len(mobiles), len(ids))
    for idx in range(max_len):
        name = names[idx] if idx < len(names) else ""
        if not name:
            continue
        cid_raw = ids[idx] if idx < len(ids) else ""
        designation = designations[idx] if idx < len(designations) else ""
        email = emails[idx] if idx < len(emails) else ""
        mobile = mobiles[idx] if idx < len(mobiles) else ""
        is_primary = 1 if idx == primary_idx else 0
        if is_primary:
            primary_name = name
            primary_email = email
            primary_mobile = mobile
        if cid_raw.isdigit():
            contact_id = int(cid_raw)
            seen_ids.add(contact_id)
            db.execute(
                """
                UPDATE client_contacts SET contact_name=?, designation=?, email=?, mobile=?,
                is_primary=?, modified_by=?, modified_at=? WHERE id=? AND client_id=?
                """,
                (name, designation, email, mobile, is_primary, username, now, contact_id, client_id),
            )
        else:
            cur = db.execute(
                """
                INSERT INTO client_contacts(client_id, contact_name, designation, email, mobile,
                is_primary, created_by, created_at, modified_by, modified_at)
                VALUES(?,?,?,?,?,?,?,?,?,?)
                """,
                (client_id, name, designation, email, mobile, is_primary, username, now, username, now),
            )
            seen_ids.add(int(cur.lastrowid))
    existing = db.execute(
        "SELECT id FROM client_contacts WHERE client_id=? AND COALESCE(is_deleted,0)=0",
        (client_id,),
    ).fetchall()
    for row in existing:
        rid = int(row[0])
        if max_len == 0:
            continue
        if rid not in seen_ids:
            db.execute(
                "UPDATE client_contacts SET is_deleted=1, modified_by=?, modified_at=? WHERE id=?",
                (username, now, rid),
            )
    if primary_name:
        db.execute(
            "UPDATE clients SET contact_person=?, email=COALESCE(NULLIF(email,''), ?), mobile=COALESCE(NULLIF(mobile,''), ?) WHERE id=?",
            (primary_name, primary_email, primary_mobile, client_id),
        )


def _save_client_addresses(db, client_id: int, form, username: str) -> None:
    now = _now_ts()
    ids = _form_getlist(form, "address_id[]")
    types = _form_getlist(form, "address_type[]")
    line1s = _form_getlist(form, "address_line1[]")
    line2s = _form_getlist(form, "address_line2[]")
    cities = _form_getlist(form, "address_city[]")
    districts = _form_getlist(form, "address_district[]")
    states = _form_getlist(form, "address_state[]")
    countries = _form_getlist(form, "address_country[]")
    pins = _form_getlist(form, "address_pin_code[]")
    primary_raw = _form_value(form, "address_primary_index")
    primary_idx = int(primary_raw) if primary_raw.isdigit() else 0
    seen_ids: set[int] = set()
    max_len = max(len(line1s), len(types), len(ids))
    billing_line = ""
    site_line = ""
    for idx in range(max_len):
        line1 = line1s[idx] if idx < len(line1s) else ""
        if not line1:
            continue
        addr_type = (types[idx] if idx < len(types) else "Billing") or "Billing"
        if addr_type not in ADDRESS_TYPES:
            addr_type = "Other"
        line2 = line2s[idx] if idx < len(line2s) else ""
        city = cities[idx] if idx < len(cities) else ""
        district = districts[idx] if idx < len(districts) else ""
        state = states[idx] if idx < len(states) else ""
        country = (countries[idx] if idx < len(countries) else "") or "India"
        pin = pins[idx] if idx < len(pins) else ""
        is_primary = 1 if idx == primary_idx else 0
        full = ", ".join(p for p in (line1, line2, city, state, pin) if p)
        if addr_type == "Billing" and not billing_line:
            billing_line = full
        if addr_type == "Site" and not site_line:
            site_line = full
        aid_raw = ids[idx] if idx < len(ids) else ""
        if aid_raw.isdigit():
            addr_id = int(aid_raw)
            seen_ids.add(addr_id)
            db.execute(
                """
                UPDATE client_addresses SET address_type=?, address_line1=?, address_line2=?,
                city=?, district=?, state=?, country=?, pin_code=?, is_primary=?,
                modified_by=?, modified_at=? WHERE id=? AND client_id=?
                """,
                (
                    addr_type,
                    line1,
                    line2,
                    city,
                    district,
                    state,
                    country,
                    pin,
                    is_primary,
                    username,
                    now,
                    addr_id,
                    client_id,
                ),
            )
        else:
            cur = db.execute(
                """
                INSERT INTO client_addresses(client_id, address_type, address_line1, address_line2,
                city, district, state, country, pin_code, is_primary,
                created_by, created_at, modified_by, modified_at)
                VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?)
                """,
                (
                    client_id,
                    addr_type,
                    line1,
                    line2,
                    city,
                    district,
                    state,
                    country,
                    pin,
                    is_primary,
                    username,
                    now,
                    username,
                    now,
                ),
            )
            seen_ids.add(int(cur.lastrowid))
    existing = db.execute(
        "SELECT id FROM client_addresses WHERE client_id=? AND COALESCE(is_deleted,0)=0",
        (client_id,),
    ).fetchall()
    for row in existing:
        rid = int(row[0])
        if rid not in seen_ids and max_len > 0:
            db.execute(
                "UPDATE client_addresses SET is_deleted=1, modified_by=?, modified_at=? WHERE id=?",
                (username, now, rid),
            )
    updates: list[str] = []
    params: list[Any] = []
    if billing_line:
        updates.append("billing_address=?")
        params.append(billing_line)
        updates.append("address=?")
        params.append(billing_line)
    if site_line:
        updates.append("site_address=?")
        params.append(site_line)
    if updates:
        params.append(client_id)
        db.execute(f"UPDATE clients SET {', '.join(updates)} WHERE id=?", params)


def _save_client_bank_details(db, client_id: int, form, username: str) -> None:
    now = _now_ts()
    ids = _form_getlist(form, "bank_detail_id[]")
    names = _form_getlist(form, "bank_detail_name[]")
    accounts = _form_getlist(form, "bank_detail_account[]")
    ifscs = _form_getlist(form, "bank_detail_ifsc[]")
    branches = _form_getlist(form, "bank_detail_branch[]")
    holders = _form_getlist(form, "bank_detail_holder[]")
    seen_ids: set[int] = set()
    max_len = max(len(names), len(accounts), len(ids))
    for idx in range(max_len):
        bank_name = names[idx] if idx < len(names) else ""
        account = accounts[idx] if idx < len(accounts) else ""
        if not bank_name and not account:
            continue
        ifsc = (ifscs[idx] if idx < len(ifscs) else "").upper()
        branch = branches[idx] if idx < len(branches) else ""
        holder = holders[idx] if idx < len(holders) else ""
        bid_raw = ids[idx] if idx < len(ids) else ""
        if bid_raw.isdigit():
            bank_id = int(bid_raw)
            seen_ids.add(bank_id)
            db.execute(
                """
                UPDATE client_bank_details SET bank_name=?, account_number=?, ifsc_swift=?,
                branch_name=?, account_holder=?, modified_by=?, modified_at=?
                WHERE id=? AND client_id=?
                """,
                (bank_name, account, ifsc, branch, holder, username, now, bank_id, client_id),
            )
        else:
            cur = db.execute(
                """
                INSERT INTO client_bank_details(client_id, bank_name, account_number, ifsc_swift,
                branch_name, account_holder, is_primary, created_by, created_at, modified_by, modified_at)
                VALUES(?,?,?,?,?,?,0,?,?,?,?)
                """,
                (client_id, bank_name, account, ifsc, branch, holder, username, now, username, now),
            )
            seen_ids.add(int(cur.lastrowid))
    existing = db.execute(
        "SELECT id FROM client_bank_details WHERE client_id=? AND COALESCE(is_deleted,0)=0",
        (client_id,),
    ).fetchall()
    for row in existing:
        rid = int(row[0])
        if rid not in seen_ids and max_len > 0:
            db.execute(
                "UPDATE client_bank_details SET is_deleted=1, modified_by=?, modified_at=? WHERE id=?",
                (username, now, rid),
            )


def list_clients_master(
    db,
    *,
    search: str = "",
    company_id: int | None = None,
    client_type: str = "",
    status: str = "",
    industry: str = "",
    include_deleted: bool = False,
    page: int = 1,
    per_page: int = 25,
    sort_by: str = "client_name",
    sort_dir: str = "asc",
) -> dict[str, Any]:
    if not _table_exists(db, "clients"):
        return {"items": [], "total": 0, "page": 1, "per_page": per_page, "pages": 0}
    sql = (
        "SELECT c.*, co.company_code, co.company_name AS linked_company_name "
        "FROM clients c LEFT JOIN companies co ON c.company_id = co.id WHERE 1=1"
    )
    count_sql = "SELECT COUNT(*) FROM clients c WHERE 1=1"
    params: list[Any] = []
    if not include_deleted:
        sql += " AND COALESCE(c.is_deleted,0)=0"
        count_sql += " AND COALESCE(c.is_deleted,0)=0"
    if company_id:
        sql += " AND c.company_id=?"
        count_sql += " AND c.company_id=?"
        params.append(company_id)
    if client_type:
        sql += " AND c.client_type=?"
        count_sql += " AND c.client_type=?"
        params.append(client_type)
    if status:
        sql += " AND c.status=?"
        count_sql += " AND c.status=?"
        params.append(status)
    if industry:
        sql += " AND c.industry=?"
        count_sql += " AND c.industry=?"
        params.append(industry)
    if search:
        clause = (
            " AND (c.client_name LIKE ? OR c.client_code LIKE ? OR c.company_name LIKE ? "
            "OR c.legal_name LIKE ? OR c.gst_number LIKE ? OR c.pan_number LIKE ? "
            "OR c.email LIKE ? OR c.mobile LIKE ? OR c.contact_person LIKE ? OR c.city LIKE ?)"
        )
        sql += clause
        count_sql += clause
        like = f"%{search}%"
        params.extend([like] * 10)
    sort_col = sort_by if sort_by in CLIENT_SORT_COLUMNS else "client_name"
    sort_col = f"c.{sort_col}"
    direction = "DESC" if str(sort_dir).lower() == "desc" else "ASC"
    sql += f" ORDER BY {sort_col} {direction}, c.id DESC"
    per_page = max(1, min(int(per_page or 25), 10000))
    page = max(1, int(page or 1))
    offset = (page - 1) * per_page
    sql += " LIMIT ? OFFSET ?"
    total = int(db.execute(count_sql, params).fetchone()[0])
    rows = db.execute(sql, [*params, per_page, offset]).fetchall()
    items = [dict(r) for r in rows]
    pages = (total + per_page - 1) // per_page if total else 0
    return {"items": items, "total": total, "page": page, "per_page": per_page, "pages": pages}


def get_client_master(db, client_id: int, *, include_deleted: bool = False) -> dict[str, Any] | None:
    if not client_id or not _table_exists(db, "clients"):
        return None
    sql = (
        "SELECT c.*, co.company_code, co.company_name AS linked_company_name "
        "FROM clients c LEFT JOIN companies co ON c.company_id = co.id WHERE c.id=?"
    )
    if not include_deleted:
        sql += " AND COALESCE(c.is_deleted,0)=0"
    row = db.execute(sql, (client_id,)).fetchone()
    if not row:
        return None
    data = dict(row)
    data["contacts"] = list_client_contacts(db, client_id)
    data["addresses"] = list_client_addresses(db, client_id)
    data["bank_details"] = list_client_bank_details(db, client_id)
    data["project_count"] = count_client_projects(db, client_id)
    primary_contact = next((c for c in data["contacts"] if c.get("is_primary")), None)
    if primary_contact:
        data["primary_contact_name"] = primary_contact.get("contact_name")
        data["primary_contact_email"] = primary_contact.get("email")
        data["primary_contact_mobile"] = primary_contact.get("mobile")
    else:
        data["primary_contact_name"] = data.get("contact_person")
        data["primary_contact_email"] = data.get("email")
        data["primary_contact_mobile"] = data.get("mobile")
    return data


def list_companies_for_client_form(db) -> list[dict[str, Any]]:
    listing = list_companies(db, status="Active", per_page=500)
    return listing.get("items", [])


def save_client_master(
    db,
    form,
    username: str,
    client_id: int | None = None,
    *,
    customer_id: int | None = None,
) -> int:
    data = _parse_client_form(form)
    if not data["client_name"]:
        raise ValueError("Client name is required.")
    if data["company_id"] and not get_company(db, data["company_id"]):
        raise ValueError("Selected company was not found.")
    validate_client_form_data(data, client_id=client_id)
    if not data["client_code"]:
        data["client_code"] = generate_client_code(db)
    validate_client_uniqueness(
        db,
        client_code=data["client_code"],
        gst_number=data["gst_number"],
        client_id=client_id,
    )
    now = _now_ts()
    core = (
        data["client_code"],
        data["client_name"],
        data["legal_name"],
        data["company_name"],
        data["client_type"],
        data["industry"],
        data["gst_number"],
        data["pan_number"],
        data["registration_number"],
        data["email"],
        data["phone"],
        data["mobile"],
        data["website"],
        data["billing_address"],
        data["site_address"],
        data["address"],
        data["country"],
        data["state"],
        data["district"],
        data["city"],
        data["pin_code"],
        data["contact_person"],
        data["payment_terms"],
        data["credit_limit"],
        data["bank_name"],
        data["account_number"],
        data["ifsc_swift"],
        data["status"],
        data["company_id"],
    )
    if client_id:
        existing = get_client_master(db, client_id, include_deleted=True)
        if not existing:
            raise ValueError("Client not found.")
        db.execute(
            """
            UPDATE clients SET client_code=?, client_name=?, legal_name=?, company_name=?,
            client_type=?, industry=?, gst_number=?, pan_number=?, registration_number=?,
            email=?, phone=?, mobile=?, website=?, billing_address=?, site_address=?, address=?,
            country=?, state=?, district=?, city=?, pin_code=?, contact_person=?,
            payment_terms=?, credit_limit=?, bank_name=?, account_number=?, ifsc_swift=?,
            status=?, company_id=?, modified_by=?, modified_at=? WHERE id=?
            """,
            (*core, username, now, client_id),
        )
        if customer_id is not None:
            db.execute("UPDATE clients SET customer_id=? WHERE id=?", (customer_id, client_id))
        _save_client_contacts(db, client_id, form, username)
        _save_client_addresses(db, client_id, form, username)
        _save_client_bank_details(db, client_id, form, username)
        log_client_field_changes(
            db,
            existing,
            get_client_master(db, client_id, include_deleted=True),
            username,
        )
        return client_id
    approval_status = _form_value(form, "approval_status") or "Draft"
    if approval_status not in APPROVAL_STATUSES:
        approval_status = "Draft"
    insert_cols = (
        "client_code, client_name, legal_name, company_name, client_type, industry, gst_number, "
        "pan_number, registration_number, email, phone, mobile, website, billing_address, "
        "site_address, address, country, state, district, city, pin_code, contact_person, "
        "payment_terms, credit_limit, bank_name, account_number, ifsc_swift, status, company_id, "
        "approval_status, created_by, created_at, modified_by, modified_at"
    )
    placeholders = ",".join(["?"] * 34)
    vals = (*core, approval_status, username, now, username, now)
    if customer_id is not None:
        cur = db.execute(
            f"INSERT INTO clients({insert_cols}, customer_id) VALUES({placeholders},?)",
            (*vals, customer_id),
        )
    else:
        cur = db.execute(f"INSERT INTO clients({insert_cols}) VALUES({placeholders})", vals)
    new_id = int(cur.lastrowid)
    _save_client_contacts(db, new_id, form, username)
    _save_client_addresses(db, new_id, form, username)
    _save_client_bank_details(db, new_id, form, username)
    log_client_audit(
        db,
        new_id,
        "create",
        username,
        remarks=f"Created client {data['client_code']}",
    )
    return new_id


def soft_delete_client_master(db, client_id: int, username: str) -> None:
    if not client_id:
        raise ValueError("Invalid client.")
    row = get_client_master(db, client_id, include_deleted=True)
    if not row:
        raise ValueError("Client not found.")
    if row.get("is_deleted"):
        return
    if client_has_active_projects(db, client_id):
        raise ValueError("Client cannot be deleted while active projects reference it. Deactivate instead.")
    now = _now_ts()
    db.execute(
        """
        UPDATE clients SET is_deleted=1, deleted_by=?, deleted_at=?,
        modified_by=?, modified_at=? WHERE id=?
        """,
        (username, now, username, now, client_id),
    )
    log_client_audit(
        db,
        client_id,
        "soft_delete",
        username,
        remarks=f"Soft-deleted client {row.get('client_code')}",
    )


def activate_client_master(db, client_id: int, username: str) -> None:
    if not get_client_master(db, client_id):
        raise ValueError("Client not found.")
    now = _now_ts()
    db.execute(
        "UPDATE clients SET status='Active', modified_by=?, modified_at=? WHERE id=?",
        (username, now, client_id),
    )
    log_client_audit(
        db,
        client_id,
        "activate",
        username,
        field_name="status",
        old_value="Inactive",
        new_value="Active",
    )


def deactivate_client_master(db, client_id: int, username: str) -> None:
    if not get_client_master(db, client_id):
        raise ValueError("Client not found.")
    now = _now_ts()
    db.execute(
        "UPDATE clients SET status='Inactive', modified_by=?, modified_at=? WHERE id=?",
        (username, now, client_id),
    )
    log_client_audit(
        db,
        client_id,
        "deactivate",
        username,
        field_name="status",
        old_value="Active",
        new_value="Inactive",
    )


def approve_client_master(db, client_id: int, username: str) -> None:
    if not get_client_master(db, client_id):
        raise ValueError("Client not found.")
    now = _now_ts()
    db.execute(
        """
        UPDATE clients SET approval_status='Approved', approved_by=?, approved_at=?,
        modified_by=?, modified_at=? WHERE id=?
        """,
        (username, now, username, now, client_id),
    )
    log_client_audit(db, client_id, "approve", username, remarks="Client approved")


def reject_client_master(db, client_id: int, username: str, remarks: str = "") -> None:
    if not get_client_master(db, client_id):
        raise ValueError("Client not found.")
    now = _now_ts()
    db.execute(
        "UPDATE clients SET approval_status='Rejected', modified_by=?, modified_at=? WHERE id=?",
        (username, now, client_id),
    )
    log_client_audit(db, client_id, "reject", username, remarks=remarks or "Client rejected")


def user_can_client_master(
    db,
    user_id: int | None,
    action: str,
    *,
    is_admin: bool = False,
) -> bool:
    if is_admin:
        return True
    if not user_id:
        return False
    action_map = {"deactivate": "edit", "activate": "edit"}
    check = action_map.get(action, action)
    try:
        from user_permission_service import empty_permission_actions, normalize_permission_actions

        row = db.execute(
            """
            SELECT granted, action_flags FROM user_tab_permissions
            WHERE user_id=? AND granted=1 AND endpoint='client_master'
            LIMIT 1
            """,
            (user_id,),
        ).fetchone()
        if not row:
            return False
        raw_flags = row["action_flags"] if hasattr(row, "keys") else row[1]
        actions = normalize_permission_actions(
            json.loads(raw_flags) if raw_flags else empty_permission_actions()
        )
        if check == "import":
            return bool(actions.get("import") or actions.get("create"))
        if check == "delete":
            return bool(actions.get("delete") or actions.get("edit"))
        return bool(actions.get(check))
    except Exception:
        return False


def clients_for_export(db, *, include_deleted: bool = False, **filters) -> list[dict[str, Any]]:
    listing = list_clients_master(db, include_deleted=include_deleted, per_page=10000, **filters)
    rows: list[dict[str, Any]] = []
    for item in listing["items"]:
        row = {col: item.get(col, "") for col in CLIENT_EXPORT_COLUMNS}
        row["company_code"] = item.get("company_code") or ""
        rows.append(row)
    return rows


def export_clients_excel(db, **filters) -> BytesIO:
    from openpyxl import Workbook

    rows = clients_for_export(db, **filters)
    wb = Workbook()
    ws = wb.active
    ws.title = "Clients"
    headers = list(CLIENT_EXPORT_COLUMNS)
    ws.append(headers)
    for row in rows:
        ws.append([row.get(h, "") for h in headers])
    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def export_clients_csv(db, **filters) -> str:
    rows = clients_for_export(db, **filters)
    si = io.StringIO()
    writer = csv.writer(si)
    headers = list(CLIENT_EXPORT_COLUMNS)
    writer.writerow(headers)
    for row in rows:
        writer.writerow([row.get(h, "") for h in headers])
    return si.getvalue()


def export_clients_pdf(db, *, report_title: str = "Client Master Report", **filters) -> BytesIO:
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.pdfgen import canvas

    rows = clients_for_export(db, **filters)
    buf = BytesIO()
    page_size = landscape(A4)
    c = canvas.Canvas(buf, pagesize=page_size)
    _, height = page_size
    y = height - 40
    c.setFont("Helvetica-Bold", 14)
    c.drawString(40, y, f"MAXEK ERP — {report_title}")
    y -= 24
    c.setFont("Helvetica", 9)
    for row in rows[:250]:
        line = (
            f"{row.get('client_code')} | {row.get('client_name')} | {row.get('gst_number') or '—'} | "
            f"{row.get('city') or '—'} | {row.get('status')}"
        )
        if y < 40:
            c.showPage()
            y = height - 40
            c.setFont("Helvetica", 9)
        c.drawString(40, y, line[:130])
        y -= 14
    c.save()
    buf.seek(0)
    return buf


def client_report(db, report_key: str, **filters) -> list[dict[str, Any]]:
    key = (report_key or "directory").lower().strip()
    if key == "active":
        filters["status"] = "Active"
    elif key == "inactive":
        filters["status"] = "Inactive"
    elif key in ("directory", "list"):
        filters.setdefault("status", "")
    if key == "contacts":
        clients = list_clients_master(db, per_page=5000, status=filters.get("status") or "", **{
            k: v for k, v in filters.items() if k in ("company_id", "client_type", "industry", "search")
        })
        out: list[dict[str, Any]] = []
        for client in clients["items"]:
            cid = int(client["id"])
            for contact in list_client_contacts(db, cid):
                out.append(
                    {
                        "client_code": client.get("client_code"),
                        "client_name": client.get("client_name"),
                        "contact_name": contact.get("contact_name"),
                        "designation": contact.get("designation"),
                        "email": contact.get("email"),
                        "mobile": contact.get("mobile"),
                        "is_primary": contact.get("is_primary"),
                        "status": client.get("status"),
                    }
                )
        return out
    listing = list_clients_master(db, per_page=5000, **filters)
    return listing["items"]


def client_import_template() -> BytesIO:
    from bulk_import_service import build_xlsx_template

    return build_xlsx_template(
        [
            "Client Code",
            "Client Name",
            "Legal Name",
            "Client Type",
            "Industry",
            "GST Number",
            "PAN Number",
            "Email",
            "Phone",
            "Mobile",
            "Billing Address",
            "City",
            "State",
            "Pin Code",
            "Primary Contact Name",
            "Primary Contact Email",
            "Primary Contact Mobile",
            "Payment Terms",
            "Credit Limit",
            "Bank Name",
            "Account Number",
            "IFSC/SWIFT",
            "Status",
        ],
        [
            "CLT101",
            "Metro Infra Ltd",
            "Metro Infrastructure Private Limited",
            "Corporate",
            "Infrastructure",
            "29ABCDE1234F1Z5",
            "ABCDE1234F",
            "billing@metroinfra.in",
            "08012345678",
            "9876543210",
            "12 MG Road",
            "Bengaluru",
            "Karnataka",
            "560001",
            "Rajesh Kumar",
            "rajesh@metroinfra.in",
            "9876543210",
            "Net 30",
            "5000000",
            "HDFC Bank",
            "12345678901234",
            "HDFC0001234",
            "Active",
        ],
    )


def ai_validate_client(
    db,
    client_id: int | None = None,
    form: dict | None = None,
) -> dict[str, Any]:
    data = dict(form or {})
    if client_id and not form:
        row = get_client_master(db, client_id)
        if row:
            data = dict(row)
    issues: list[str] = []
    suggestions: list[str] = []
    duplicates: list[str] = []
    missing: list[str] = []
    code = (data.get("client_code") or "").strip()
    name = (data.get("client_name") or "").strip()
    gst = (data.get("gst_number") or "").strip()
    if not name:
        issues.append("Client name is missing.")
        missing.append("client_name")
    if not code:
        suggestions.append("Client code will be auto-generated on save.")
    else:
        try:
            validate_client_uniqueness(db, client_code=code, gst_number=gst, client_id=client_id)
        except ValueError as exc:
            duplicates.append(str(exc))
    if gst:
        try:
            validate_gst_number(gst)
        except ValueError as exc:
            issues.append(str(exc))
    else:
        suggestions.append("Add GST number for tax-compliant billing.")
    if not (data.get("email") or "").strip():
        missing.append("email")
        suggestions.append("Add billing email for invoices and correspondence.")
    if not (data.get("mobile") or data.get("phone") or "").strip():
        missing.append("mobile")
    if not (data.get("billing_address") or data.get("address") or "").strip():
        missing.append("billing_address")
        suggestions.append("Add billing address for invoices.")
    if not (data.get("pan_number") or "").strip() and gst:
        suggestions.append("PAN is recommended when GST is provided.")
    if name:
        like_name = f"%{name[:20]}%"
        hits = db.execute(
            """
            SELECT client_code, client_name FROM clients
            WHERE client_name LIKE ? AND COALESCE(is_deleted,0)=0
            """,
            (like_name,),
        ).fetchall()
        for hit in hits:
            hit_id = None
            if client_id:
                existing = get_client_master(db, client_id, include_deleted=True)
                hit_id = existing.get("client_code") if existing else None
            hit_code = hit[0] if not hasattr(hit, "keys") else hit["client_code"]
            if hit_code != hit_id and hit_code != code:
                hit_name = hit[1] if not hasattr(hit, "keys") else hit["client_name"]
                duplicates.append(f"Similar client: {hit_code} — {hit_name}")
    try:
        from ai_service import ai_assist_validation

        ai_extra = ai_assist_validation("client_master", data)
        if isinstance(ai_extra, dict):
            issues.extend(ai_extra.get("issues") or [])
            suggestions.extend(ai_extra.get("suggestions") or [])
    except Exception:
        pass
    return {
        "ok": not issues and not duplicates,
        "issues": issues,
        "suggestions": suggestions,
        "duplicates": duplicates,
        "missing": missing,
    }
