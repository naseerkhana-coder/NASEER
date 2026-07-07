"""Department Master (MODULE-003) — departments linked to Company and optional Branch."""

from __future__ import annotations

import csv
import io
import json
import re
from io import BytesIO
from typing import Any

from branch_master_service import ensure_branch_master_schema, get_branch_master
from company_master_service import (
    _ensure_column,
    _now_ts,
    _table_exists,
    ensure_company_master_schema,
    get_company,
    list_companies,
)

DEPARTMENT_STATUSES = ("Active", "Inactive")
APPROVAL_STATUSES = ("Draft", "Pending", "Approved", "Rejected")
DEPARTMENT_SORT_COLUMNS = (
    "department_code",
    "department_name",
    "department_short_name",
    "status",
    "created_at",
    "company_id",
    "branch_id",
)
DEPARTMENT_EXPORT_COLUMNS = (
    "company_code",
    "company_name",
    "branch_code",
    "branch_name",
    "department_code",
    "department_name",
    "department_short_name",
    "department_head",
    "description",
    "status",
    "approval_status",
)
DEPARTMENT_AUDIT_FIELDS = (
    "department_code",
    "department_name",
    "department_short_name",
    "company_id",
    "branch_id",
    "department_head",
    "description",
    "status",
    "approval_status",
)
EMPLOYEE_DEPARTMENT_TABLES = (
    "staff",
    "users",
    "workers",
    "payroll_lines",
    "payroll_records",
    "user_maker_assignments",
    "petty_cash_requests",
    "head_office_expenses",
)
DEFAULT_DEPARTMENT_NAMES = (
    "Head Office",
    "Accounts",
    "Site Operations",
    "Store",
    "Purchase",
    "HR & Payroll",
    "Projects",
    "Management",
)


def ensure_department_master_schema(db) -> None:
    """Extend departments table for MODULE-003 (idempotent)."""
    ensure_company_master_schema(db)
    ensure_branch_master_schema(db)
    db.execute(
        """
        CREATE TABLE IF NOT EXISTS departments(
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            department_code TEXT,
            department_name TEXT,
            department_short_name TEXT,
            company_id INTEGER,
            branch_id INTEGER,
            department_head TEXT,
            description TEXT,
            status TEXT DEFAULT 'Active',
            approval_status TEXT DEFAULT 'Draft',
            approved_by TEXT,
            approved_at TEXT,
            created_by TEXT,
            created_at TEXT,
            modified_by TEXT,
            modified_at TEXT,
            is_deleted INTEGER DEFAULT 0,
            deleted_by TEXT,
            deleted_at TEXT,
            customer_id INTEGER
        )
        """
    )
    for col, ctype in (
        ("department_code", "TEXT"),
        ("department_short_name", "TEXT"),
        ("company_id", "INTEGER"),
        ("branch_id", "INTEGER"),
        ("department_head", "TEXT"),
        ("approval_status", "TEXT DEFAULT 'Draft'"),
        ("approved_by", "TEXT"),
        ("approved_at", "TEXT"),
        ("created_by", "TEXT"),
        ("created_at", "TEXT"),
        ("modified_by", "TEXT"),
        ("modified_at", "TEXT"),
        ("is_deleted", "INTEGER DEFAULT 0"),
        ("deleted_by", "TEXT"),
        ("deleted_at", "TEXT"),
        ("customer_id", "INTEGER"),
    ):
        _ensure_column(db, "departments", col, ctype)
    _migrate_legacy_departments(db)
    try:
        from audit_trail_service import ensure_audit_schema

        ensure_audit_schema(db)
    except Exception:
        pass


def _migrate_legacy_departments(db) -> None:
    if not _table_exists(db, "departments"):
        return
    first_company = db.execute(
        "SELECT id FROM companies WHERE COALESCE(is_deleted,0)=0 ORDER BY id LIMIT 1"
    ).fetchone()
    company_id = int(first_company[0]) if first_company else None
    null_company = db.execute(
        "SELECT id, department_name FROM departments WHERE company_id IS NULL AND COALESCE(is_deleted,0)=0"
    ).fetchall()
    for row in null_company:
        dept_id = int(row[0])
        name = (row[1] or "").strip()
        updates: list[Any] = []
        sets: list[str] = []
        if company_id:
            sets.append("company_id=?")
            updates.append(company_id)
        code_row = db.execute(
            "SELECT department_code FROM departments WHERE id=?",
            (dept_id,),
        ).fetchone()
        code = (code_row[0] if code_row else "") or ""
        if not code and name and company_id:
            code = _next_department_code(db, company_id)
            sets.append("department_code=?")
            updates.append(code)
        if sets:
            db.execute(
                f"UPDATE departments SET {', '.join(sets)} WHERE id=?",
                (*updates, dept_id),
            )
    count = int(db.execute("SELECT COUNT(*) FROM departments").fetchone()[0])
    if count == 0 and company_id:
        now = _now_ts()
        for idx, name in enumerate(DEFAULT_DEPARTMENT_NAMES, start=1):
            code = f"DEPT-{company_id}-{idx:03d}"
            db.execute(
                """
                INSERT INTO departments(
                    department_code, department_name, company_id, status,
                    approval_status, created_by, created_at, modified_by, modified_at
                ) VALUES (?, ?, ?, 'Active', 'Approved', 'system', ?, 'system', ?)
                """,
                (code, name, company_id, now, now),
            )


def _next_department_code(db, company_id: int) -> str:
    prefix = f"DEPT-{company_id}-"
    row = db.execute(
        "SELECT department_code FROM departments WHERE company_id=? AND department_code LIKE ? "
        "ORDER BY id DESC LIMIT 1",
        (company_id, f"{prefix}%"),
    ).fetchone()
    seq = 1
    if row and row[0]:
        match = re.search(r"-(\d+)$", str(row[0]))
        if match:
            seq = int(match.group(1)) + 1
    return f"{prefix}{seq:03d}"


def validate_department_uniqueness(
    db,
    *,
    company_id: int,
    department_code: str,
    department_name: str,
    department_id: int | None = None,
) -> None:
    code = (department_code or "").strip().upper()
    name = (department_name or "").strip()
    if not code:
        raise ValueError("Department code is required.")
    if not name:
        raise ValueError("Department name is required.")
    row = db.execute(
        """
        SELECT id FROM departments
        WHERE company_id=? AND UPPER(department_code)=? AND COALESCE(is_deleted,0)=0
        """,
        (company_id, code),
    ).fetchone()
    if row and (not department_id or int(row[0]) != int(department_id)):
        raise ValueError(f"Department code '{code}' already exists for this company.")
    row = db.execute(
        """
        SELECT id FROM departments
        WHERE company_id=? AND LOWER(TRIM(department_name))=LOWER(?) AND COALESCE(is_deleted,0)=0
        """,
        (company_id, name),
    ).fetchone()
    if row and (not department_id or int(row[0]) != int(department_id)):
        raise ValueError(f"Department name '{name}' already exists for this company.")


def department_has_employees(db, department_id: int) -> bool:
    row = get_department_master(db, department_id, include_deleted=True)
    if not row:
        return False
    dept_name = (row.get("department_name") or "").strip()
    if not dept_name:
        return False
    for table in EMPLOYEE_DEPARTMENT_TABLES:
        if not _table_exists(db, table):
            continue
        cols = {r[1] for r in db.execute(f"PRAGMA table_info({table})").fetchall()}
        if "department_id" in cols:
            hit = db.execute(
                f"SELECT 1 FROM {table} WHERE department_id=? LIMIT 1",
                (department_id,),
            ).fetchone()
            if hit:
                return True
        if "department" in cols:
            hit = db.execute(
                f"SELECT 1 FROM {table} WHERE TRIM(department)=? LIMIT 1",
                (dept_name,),
            ).fetchone()
            if hit:
                return True
    return False


def log_department_audit(
    db,
    department_id: int,
    action: str,
    username: str,
    *,
    field_name: str | None = None,
    old_value: str | None = None,
    new_value: str | None = None,
    remarks: str | None = None,
) -> None:
    try:
        from audit_trail_service import log_audit_event

        log_audit_event(
            db,
            record_table="departments",
            record_id=department_id,
            action=action,
            changed_by=username,
            field_name=field_name,
            old_value=old_value,
            new_value=new_value,
            remarks=remarks,
        )
    except Exception:
        pass


def log_department_field_changes(
    db,
    before: dict[str, Any] | None,
    after: dict[str, Any] | None,
    username: str,
) -> None:
    if not before or not after:
        return
    dept_id = int(after.get("id") or before.get("id") or 0)
    if not dept_id:
        return
    for field in DEPARTMENT_AUDIT_FIELDS:
        old_val = before.get(field)
        new_val = after.get(field)
        if str(old_val or "") != str(new_val or ""):
            log_department_audit(
                db,
                dept_id,
                "update",
                username,
                field_name=field,
                old_value=str(old_val or ""),
                new_value=str(new_val or ""),
            )


def list_department_audit_trail(db, department_id: int, limit: int = 100) -> list[dict[str, Any]]:
    try:
        from audit_trail_service import list_audit_trail

        return list_audit_trail(db, "departments", department_id, limit=limit)
    except Exception:
        return []


def _department_row_to_dict(row) -> dict[str, Any]:
    return dict(row)


def list_departments_master(
    db,
    *,
    search: str = "",
    company_id: int | None = None,
    branch_id: int | None = None,
    status: str = "",
    include_deleted: bool = False,
    page: int = 1,
    per_page: int = 25,
    sort_by: str = "department_name",
    sort_dir: str = "asc",
) -> dict[str, Any]:
    if not _table_exists(db, "departments"):
        return {"items": [], "total": 0, "page": 1, "per_page": per_page, "pages": 0}
    sql = (
        "SELECT d.*, c.company_code, c.company_name, c.legal_name AS company_legal_name, "
        "b.branch_code, b.branch_name "
        "FROM departments d "
        "JOIN companies c ON d.company_id = c.id "
        "LEFT JOIN company_branches b ON d.branch_id = b.id "
        "WHERE COALESCE(c.is_deleted,0)=0"
    )
    count_sql = (
        "SELECT COUNT(*) FROM departments d "
        "JOIN companies c ON d.company_id = c.id "
        "LEFT JOIN company_branches b ON d.branch_id = b.id "
        "WHERE COALESCE(c.is_deleted,0)=0"
    )
    params: list[Any] = []
    if not include_deleted:
        sql += " AND COALESCE(d.is_deleted,0)=0"
        count_sql += " AND COALESCE(d.is_deleted,0)=0"
    if company_id:
        sql += " AND d.company_id=?"
        count_sql += " AND d.company_id=?"
        params.append(company_id)
    if branch_id:
        sql += " AND d.branch_id=?"
        count_sql += " AND d.branch_id=?"
        params.append(branch_id)
    if status:
        sql += " AND d.status=?"
        count_sql += " AND d.status=?"
        params.append(status)
    if search:
        clause = (
            " AND (d.department_name LIKE ? OR d.department_code LIKE ? "
            "OR d.department_short_name LIKE ? OR d.department_head LIKE ? "
            "OR c.company_name LIKE ? OR c.company_code LIKE ? OR b.branch_name LIKE ?)"
        )
        sql += clause
        count_sql += clause
        like = f"%{search}%"
        params.extend([like, like, like, like, like, like, like])
    sort_col = sort_by if sort_by in DEPARTMENT_SORT_COLUMNS else "department_name"
    if sort_col in ("company_id", "branch_id"):
        sort_col = f"d.{sort_col}"
    else:
        sort_col = f"d.{sort_col}"
    direction = "DESC" if str(sort_dir).lower() == "desc" else "ASC"
    sql += f" ORDER BY {sort_col} {direction}, d.id DESC"
    per_page = max(1, min(int(per_page or 25), 10000))
    page = max(1, int(page or 1))
    offset = (page - 1) * per_page
    sql += " LIMIT ? OFFSET ?"
    total = int(db.execute(count_sql, params).fetchone()[0])
    rows = db.execute(sql, [*params, per_page, offset]).fetchall()
    items = [_department_row_to_dict(r) for r in rows]
    pages = (total + per_page - 1) // per_page if total else 0
    return {"items": items, "total": total, "page": page, "per_page": per_page, "pages": pages}


def get_department_master(db, department_id: int, *, include_deleted: bool = False) -> dict[str, Any] | None:
    if not department_id or not _table_exists(db, "departments"):
        return None
    sql = (
        "SELECT d.*, c.company_code, c.company_name, c.legal_name AS company_legal_name, "
        "b.branch_code, b.branch_name "
        "FROM departments d "
        "JOIN companies c ON d.company_id = c.id "
        "LEFT JOIN company_branches b ON d.branch_id = b.id "
        "WHERE d.id=?"
    )
    if not include_deleted:
        sql += " AND COALESCE(d.is_deleted,0)=0"
    row = db.execute(sql, (department_id,)).fetchone()
    return _department_row_to_dict(row) if row else None


def list_active_department_names(db, company_id: int | None = None) -> list[str]:
    listing = list_departments_master(db, company_id=company_id, status="Active", per_page=5000)
    return sorted({str(r.get("department_name") or "").strip() for r in listing["items"] if r.get("department_name")})


def _parse_department_form(form) -> dict[str, Any]:
    branch_raw = form.get("branch_id")
    branch_id = int(branch_raw) if branch_raw not in (None, "", "0") else None
    return {
        "company_id": int(form.get("company_id") or 0),
        "branch_id": branch_id,
        "department_code": (form.get("department_code") or "").strip().upper(),
        "department_name": (form.get("department_name") or "").strip(),
        "department_short_name": (form.get("department_short_name") or "").strip(),
        "department_head": (form.get("department_head") or "").strip(),
        "description": (form.get("description") or "").strip(),
        "status": (form.get("status") or "Active").strip(),
    }


def save_department_master(
    db,
    form,
    username: str,
    department_id: int | None = None,
    *,
    customer_id: int | None = None,
) -> int:
    data = _parse_department_form(form)
    company_id = data["company_id"]
    if not company_id:
        raise ValueError("Company is required.")
    if not get_company(db, company_id):
        raise ValueError("Selected company was not found.")
    if data["branch_id"]:
        branch = get_branch_master(db, data["branch_id"])
        if not branch:
            raise ValueError("Selected branch was not found.")
        if int(branch.get("company_id") or 0) != company_id:
            raise ValueError("Branch must belong to the selected company.")
    if not data["department_code"]:
        raise ValueError("Department code is required.")
    if not data["department_name"]:
        raise ValueError("Department name is required.")
    if data["status"] not in DEPARTMENT_STATUSES:
        raise ValueError("Select a valid status.")
    validate_department_uniqueness(
        db,
        company_id=company_id,
        department_code=data["department_code"],
        department_name=data["department_name"],
        department_id=department_id,
    )
    now = _now_ts()
    core = (
        company_id,
        data["branch_id"],
        data["department_code"],
        data["department_name"],
        data["department_short_name"],
        data["department_head"],
        data["description"],
        data["status"],
    )
    if department_id:
        existing = get_department_master(db, department_id, include_deleted=True)
        if not existing:
            raise ValueError("Department not found.")
        db.execute(
            """
            UPDATE departments SET company_id=?, branch_id=?, department_code=?, department_name=?,
            department_short_name=?, department_head=?, description=?, status=?,
            modified_by=?, modified_at=? WHERE id=?
            """,
            (*core, username, now, department_id),
        )
        if customer_id is not None:
            db.execute("UPDATE departments SET customer_id=? WHERE id=?", (customer_id, department_id))
        log_department_field_changes(
            db,
            existing,
            get_department_master(db, department_id, include_deleted=True),
            username,
        )
        return department_id
    approval_status = (form.get("approval_status") or "Draft").strip()
    if approval_status not in APPROVAL_STATUSES:
        approval_status = "Draft"
    insert_cols = (
        "company_id, branch_id, department_code, department_name, department_short_name, "
        "department_head, description, status, approval_status, created_by, created_at, "
        "modified_by, modified_at"
    )
    placeholders = ",".join(["?"] * 13)
    vals = (*core, approval_status, username, now, username, now)
    if customer_id is not None:
        cur = db.execute(
            f"INSERT INTO departments({insert_cols}, customer_id) VALUES({placeholders},?)",
            (*vals, customer_id),
        )
    else:
        cur = db.execute(f"INSERT INTO departments({insert_cols}) VALUES({placeholders})", vals)
    new_id = int(cur.lastrowid)
    log_department_audit(
        db,
        new_id,
        "create",
        username,
        remarks=f"Created department {data['department_code']}",
    )
    return new_id


def soft_delete_department_master(db, department_id: int, username: str) -> None:
    if not department_id:
        raise ValueError("Invalid department.")
    row = get_department_master(db, department_id, include_deleted=True)
    if not row:
        raise ValueError("Department not found.")
    if row.get("is_deleted"):
        return
    if department_has_employees(db, department_id):
        raise ValueError("Department cannot be deleted because employees exist. Deactivate instead.")
    now = _now_ts()
    db.execute(
        """
        UPDATE departments SET is_deleted=1, deleted_by=?, deleted_at=?,
        modified_by=?, modified_at=? WHERE id=?
        """,
        (username, now, username, now, department_id),
    )
    log_department_audit(
        db,
        department_id,
        "soft_delete",
        username,
        remarks=f"Soft-deleted department {row.get('department_code')}",
    )


def activate_department_master(db, department_id: int, username: str) -> None:
    if not get_department_master(db, department_id):
        raise ValueError("Department not found.")
    now = _now_ts()
    db.execute(
        "UPDATE departments SET status='Active', modified_by=?, modified_at=? WHERE id=?",
        (username, now, department_id),
    )
    log_department_audit(
        db,
        department_id,
        "activate",
        username,
        field_name="status",
        old_value="Inactive",
        new_value="Active",
    )


def deactivate_department_master(db, department_id: int, username: str) -> None:
    if not get_department_master(db, department_id):
        raise ValueError("Department not found.")
    now = _now_ts()
    db.execute(
        "UPDATE departments SET status='Inactive', modified_by=?, modified_at=? WHERE id=?",
        (username, now, department_id),
    )
    log_department_audit(
        db,
        department_id,
        "deactivate",
        username,
        field_name="status",
        old_value="Active",
        new_value="Inactive",
    )


def approve_department_master(db, department_id: int, username: str) -> None:
    if not get_department_master(db, department_id):
        raise ValueError("Department not found.")
    now = _now_ts()
    db.execute(
        """
        UPDATE departments SET approval_status='Approved', approved_by=?, approved_at=?,
        modified_by=?, modified_at=? WHERE id=?
        """,
        (username, now, username, now, department_id),
    )
    log_department_audit(db, department_id, "approve", username, remarks="Department approved")


def reject_department_master(db, department_id: int, username: str, remarks: str = "") -> None:
    if not get_department_master(db, department_id):
        raise ValueError("Department not found.")
    now = _now_ts()
    db.execute(
        "UPDATE departments SET approval_status='Rejected', modified_by=?, modified_at=? WHERE id=?",
        (username, now, department_id),
    )
    log_department_audit(db, department_id, "reject", username, remarks=remarks or "Department rejected")


def user_can_department_master(
    db,
    user_id: int | None,
    action: str,
    *,
    is_admin: bool = False,
) -> bool:
    if is_admin:
        return True
    if not user_id:
        return False
    action_map = {"deactivate": "edit", "activate": "edit"}
    check = action_map.get(action, action)
    try:
        from user_permission_service import (
            empty_permission_actions,
            ensure_user_tab_permissions_schema,
            normalize_permission_actions,
        )

        ensure_user_tab_permissions_schema(db)
        row = db.execute(
            """
            SELECT granted, action_flags FROM user_tab_permissions
            WHERE user_id=? AND granted=1 AND endpoint='department_master'
            LIMIT 1
            """,
            (user_id,),
        ).fetchone()
        if not row:
            return False
        raw_flags = row["action_flags"] if hasattr(row, "keys") else row[1]
        actions = normalize_permission_actions(
            json.loads(raw_flags) if raw_flags else empty_permission_actions()
        )
        if check == "import":
            return bool(actions.get("import") or actions.get("create"))
        if check == "delete":
            return bool(actions.get("delete") or actions.get("edit"))
        return bool(actions.get(check))
    except Exception:
        return False


def departments_for_export(db, *, include_deleted: bool = False, **filters) -> list[dict[str, Any]]:
    listing = list_departments_master(db, include_deleted=include_deleted, per_page=10000, **filters)
    return [{col: item.get(col, "") for col in DEPARTMENT_EXPORT_COLUMNS} for item in listing["items"]]


def export_departments_excel(db, **filters) -> BytesIO:
    from openpyxl import Workbook

    rows = departments_for_export(db, **filters)
    wb = Workbook()
    ws = wb.active
    ws.title = "Departments"
    headers = list(DEPARTMENT_EXPORT_COLUMNS)
    ws.append(headers)
    for row in rows:
        ws.append([row.get(h, "") for h in headers])
    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def export_departments_csv(db, **filters) -> str:
    rows = departments_for_export(db, **filters)
    si = io.StringIO()
    writer = csv.writer(si)
    headers = list(DEPARTMENT_EXPORT_COLUMNS)
    writer.writerow(headers)
    for row in rows:
        writer.writerow([row.get(h, "") for h in headers])
    return si.getvalue()


def export_departments_pdf(db, *, report_title: str = "Department Master Report", **filters) -> BytesIO:
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.pdfgen import canvas

    rows = departments_for_export(db, **filters)
    buf = BytesIO()
    page_size = landscape(A4)
    c = canvas.Canvas(buf, pagesize=page_size)
    _, height = page_size
    y = height - 40
    c.setFont("Helvetica-Bold", 14)
    c.drawString(40, y, f"MAXEK ERP — {report_title}")
    y -= 24
    c.setFont("Helvetica", 9)
    for row in rows[:250]:
        line = (
            f"{row.get('company_code')} | {row.get('department_code')} | {row.get('department_name')} | "
            f"{row.get('branch_name') or '—'} | {row.get('status')}"
        )
        if y < 40:
            c.showPage()
            y = height - 40
            c.setFont("Helvetica", 9)
        c.drawString(40, y, line[:130])
        y -= 14
    c.save()
    buf.seek(0)
    return buf


def department_report(db, report_key: str, **filters) -> list[dict[str, Any]]:
    key = (report_key or "list").lower().strip()
    if key == "active":
        filters["status"] = "Active"
    elif key == "inactive":
        filters["status"] = "Inactive"
    elif key == "directory":
        filters["status"] = "Active"
    listing = list_departments_master(db, per_page=5000, **filters)
    return listing["items"]


def department_import_template() -> BytesIO:
    from bulk_import_service import build_xlsx_template

    return build_xlsx_template(
        [
            "Company Code",
            "Branch Code",
            "Department Code",
            "Department Name",
            "Department Short Name",
            "Department Head",
            "Description",
            "Status",
        ],
        [
            "CO-2026-0001",
            "",
            "DEPT-HR-01",
            "HR & Payroll",
            "HR",
            "Manager Name",
            "Human resources and payroll",
            "Active",
        ],
    )


def ai_validate_department(
    db,
    department_id: int | None = None,
    form: dict | None = None,
) -> dict[str, Any]:
    data = dict(form or {})
    if department_id and not form:
        row = get_department_master(db, department_id)
        if row:
            data = dict(row)
    issues: list[str] = []
    suggestions: list[str] = []
    duplicates: list[str] = []
    missing: list[str] = []
    company_id = int(data.get("company_id") or 0)
    code = (data.get("department_code") or "").strip()
    name = (data.get("department_name") or "").strip()
    if not company_id:
        issues.append("Company is not selected.")
        missing.append("company_id")
    if not code:
        issues.append("Department code is missing.")
        missing.append("department_code")
    elif company_id:
        try:
            validate_department_uniqueness(
                db,
                company_id=company_id,
                department_code=code,
                department_name=name or code,
                department_id=department_id,
            )
        except ValueError as exc:
            duplicates.append(str(exc))
    if not name:
        issues.append("Department name is missing.")
        missing.append("department_name")
    if not (data.get("department_short_name") or "").strip():
        suggestions.append("Add a short name for reports and mobile views.")
    if not (data.get("department_head") or "").strip():
        suggestions.append("Assign a department head for workflow routing.")
    if not (data.get("description") or "").strip():
        suggestions.append("Add a description for HR and project teams.")
    branch_id = data.get("branch_id")
    if branch_id and company_id:
        branch = get_branch_master(db, int(branch_id))
        if branch and int(branch.get("company_id") or 0) != company_id:
            issues.append("Branch does not belong to the selected company.")
    result = {
        "ok": not issues and not duplicates,
        "issues": issues,
        "duplicates": duplicates,
        "suggestions": suggestions,
        "missing": missing,
    }
    try:
        from ai_service import chat_completion_json

        prompt = json.dumps({"department": data, "rule_findings": result}, ensure_ascii=False)
        ai = chat_completion_json(
            "You validate ERP department master records. Return JSON with keys: issues (array), suggestions (array).",
            prompt,
            max_tokens=500,
        )
        for key in ("issues", "suggestions"):
            extra = ai.get(key) or []
            if isinstance(extra, list):
                result[key].extend(str(x) for x in extra if x)
        result["ok"] = not result["issues"] and not result["duplicates"]
        result["ai_enriched"] = True
    except Exception:
        result["ai_enriched"] = False
    return result


def list_companies_for_department_form(db) -> list[dict[str, Any]]:
    listing = list_companies(db, per_page=1000)
    return [
        {
            "id": c["id"],
            "company_code": c.get("company_code"),
            "company_name": c.get("company_name") or c.get("legal_name"),
        }
        for c in listing.get("items", [])
    ]


def list_branches_for_department_form(db, company_id: int | None = None) -> list[dict[str, Any]]:
    from branch_master_service import list_branches_master

    listing = list_branches_master(db, company_id=company_id, status="Active", per_page=1000)
    return [
        {
            "id": b["id"],
            "branch_code": b.get("branch_code"),
            "branch_name": b.get("branch_name"),
            "company_id": b.get("company_id"),
        }
        for b in listing.get("items", [])
    ]
