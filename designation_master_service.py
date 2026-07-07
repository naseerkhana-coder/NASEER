"""Designation Master (MODULE-005) — job titles linked to Company and optional Department."""

from __future__ import annotations

import csv
import io
import json
import re
from io import BytesIO
from typing import Any

from company_master_service import (
    _ensure_column,
    _now_ts,
    _table_exists,
    ensure_company_master_schema,
    get_company,
    list_companies,
)
from department_master_service import ensure_department_master_schema, get_department_master

DESIGNATION_STATUSES = ("Active", "Inactive")
APPROVAL_STATUSES = ("Draft", "Pending", "Approved", "Rejected")
WORKFLOW_ROLE_DEFAULTS = ("Maker", "Checker", "Approver", "None")
DESIGNATION_SORT_COLUMNS = (
    "designation_code",
    "designation_name",
    "designation_short_name",
    "grade_level",
    "status",
    "created_at",
    "company_id",
    "department_id",
)
DESIGNATION_EXPORT_COLUMNS = (
    "company_code",
    "company_name",
    "department_code",
    "department_name",
    "designation_code",
    "designation_name",
    "designation_short_name",
    "grade_level",
    "workflow_role_default",
    "description",
    "status",
    "approval_status",
)
DESIGNATION_AUDIT_FIELDS = (
    "designation_code",
    "designation_name",
    "designation_short_name",
    "company_id",
    "department_id",
    "grade_level",
    "workflow_role_default",
    "description",
    "status",
    "approval_status",
)
DESIGNATION_REFERENCE_TABLES = (
    ("users", "designation_id"),
    ("staff", "designation_id"),
    ("workflow_master", "maker_designation_id"),
    ("workflow_master", "checker_designation_id"),
    ("workflow_master", "approver_designation_id"),
)
DEFAULT_DESIGNATION_NAMES = (
    "Site Engineer",
    "Supervisor",
    "Store Keeper",
    "Project Staff",
    "Project Engineer",
    "Purchase Manager",
    "HR Staff",
    "Accounts Staff",
    "Department Head",
    "Accounts Manager",
    "Store Manager",
    "Project Manager",
    "Managing Director",
)


def ensure_designation_master_schema(db) -> None:
    """Extend designations table for MODULE-005 (idempotent)."""
    ensure_company_master_schema(db)
    ensure_department_master_schema(db)
    db.execute(
        """
        CREATE TABLE IF NOT EXISTS designations(
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            designation_name TEXT UNIQUE,
            description TEXT,
            status TEXT DEFAULT 'Active'
        )
        """
    )
    for col, ctype in (
        ("designation_code", "TEXT"),
        ("designation_short_name", "TEXT"),
        ("company_id", "INTEGER"),
        ("department_id", "INTEGER"),
        ("grade_level", "TEXT"),
        ("workflow_role_default", "TEXT DEFAULT 'None'"),
        ("approval_status", "TEXT DEFAULT 'Draft'"),
        ("approved_by", "TEXT"),
        ("approved_at", "TEXT"),
        ("created_by", "TEXT"),
        ("created_at", "TEXT"),
        ("modified_by", "TEXT"),
        ("modified_at", "TEXT"),
        ("is_deleted", "INTEGER DEFAULT 0"),
        ("deleted_by", "TEXT"),
        ("deleted_at", "TEXT"),
        ("customer_id", "INTEGER"),
    ):
        _ensure_column(db, "designations", col, ctype)
    db.execute(
        "UPDATE designations SET status='Active' WHERE status IS NULL OR TRIM(status)=''"
    )
    _migrate_legacy_designations(db)
    try:
        from audit_trail_service import ensure_audit_schema

        ensure_audit_schema(db)
    except Exception:
        pass


def _migrate_legacy_designations(db) -> None:
    if not _table_exists(db, "designations"):
        return
    first_company = db.execute(
        "SELECT id FROM companies WHERE COALESCE(is_deleted,0)=0 ORDER BY id LIMIT 1"
    ).fetchone()
    company_id = int(first_company[0]) if first_company else None
    rows = db.execute(
        """
        SELECT id, designation_name, designation_code FROM designations
        WHERE COALESCE(is_deleted,0)=0
        """
    ).fetchall()
    for row in rows:
        did = int(row[0])
        name = (row[1] or "").strip()
        code = (row[2] or "").strip() if len(row) > 2 else ""
        updates: list[str] = []
        params: list[Any] = []
        if company_id and not db.execute(
            "SELECT company_id FROM designations WHERE id=?", (did,)
        ).fetchone()[0]:
            updates.append("company_id=?")
            params.append(company_id)
        if not code and name:
            slug = re.sub(r"[^A-Za-z0-9]+", "-", name).strip("-").upper()[:20]
            generated = f"DES-{slug or did}"
            updates.append("designation_code=?")
            params.append(generated)
        if not db.execute(
            "SELECT approval_status FROM designations WHERE id=?", (did,)
        ).fetchone()[0]:
            updates.append("approval_status=?")
            params.append("Approved")
        if updates:
            params.append(did)
            db.execute(f"UPDATE designations SET {', '.join(updates)} WHERE id=?", params)


def _next_designation_code(db, company_id: int) -> str:
    prefix = f"DG-{company_id}-"
    row = db.execute(
        """
        SELECT designation_code FROM designations
        WHERE company_id=? AND designation_code LIKE ? AND COALESCE(is_deleted,0)=0
        ORDER BY id DESC LIMIT 1
        """,
        (company_id, f"{prefix}%"),
    ).fetchone()
    seq = 1
    if row and row[0]:
        m = re.search(r"-(\d+)$", str(row[0]))
        if m:
            seq = int(m.group(1)) + 1
    return f"{prefix}{seq:03d}"


def validate_designation_uniqueness(
    db,
    *,
    company_id: int,
    designation_code: str,
    designation_name: str,
    designation_id: int | None = None,
) -> None:
    code = (designation_code or "").strip().upper()
    name = (designation_name or "").strip()
    if not code:
        raise ValueError("Designation code is required.")
    if not name:
        raise ValueError("Designation name is required.")
    row = db.execute(
        """
        SELECT id FROM designations
        WHERE company_id=? AND UPPER(designation_code)=? AND COALESCE(is_deleted,0)=0
        """,
        (company_id, code),
    ).fetchone()
    if row and (not designation_id or int(row[0]) != int(designation_id)):
        raise ValueError(f"Designation code '{designation_code}' already exists for this company.")
    row = db.execute(
        """
        SELECT id FROM designations
        WHERE LOWER(designation_name)=LOWER(?) AND COALESCE(is_deleted,0)=0
        """,
        (name,),
    ).fetchone()
    if row and (not designation_id or int(row[0]) != int(designation_id)):
        raise ValueError(f"Designation name '{designation_name}' is already in use.")


def log_designation_audit(
    db,
    designation_id: int,
    action: str,
    username: str,
    *,
    field_name: str | None = None,
    old_value: str | None = None,
    new_value: str | None = None,
    remarks: str | None = None,
) -> None:
    try:
        from audit_trail_service import log_audit_event

        log_audit_event(
            db,
            record_table="designations",
            record_id=designation_id,
            action=action,
            changed_by=username,
            field_name=field_name,
            old_value=old_value,
            new_value=new_value,
            remarks=remarks,
        )
    except Exception:
        pass


def log_designation_field_changes(
    db,
    before: dict[str, Any] | None,
    after: dict[str, Any] | None,
    actor: str,
) -> None:
    if not before or not after:
        return
    did = int(after.get("id") or before.get("id") or 0)
    if not did:
        return
    for field in DESIGNATION_AUDIT_FIELDS:
        old_val = before.get(field)
        new_val = after.get(field)
        if str(old_val or "") != str(new_val or ""):
            log_designation_audit(
                db,
                did,
                "update",
                actor,
                field_name=field,
                old_value=str(old_val or ""),
                new_value=str(new_val or ""),
            )


def list_designation_audit_trail(db, designation_id: int, limit: int = 100) -> list[dict[str, Any]]:
    try:
        from audit_trail_service import list_audit_trail

        return list_audit_trail(db, "designations", designation_id, limit=limit)
    except Exception:
        return []


def list_designations_master(
    db,
    *,
    search: str = "",
    company_id: int | None = None,
    department_id: int | None = None,
    status: str = "",
    workflow_role: str = "",
    customer_id: int | None = None,
    include_deleted: bool = False,
    page: int = 1,
    per_page: int = 25,
    sort_by: str = "designation_name",
    sort_dir: str = "asc",
) -> dict[str, Any]:
    if not _table_exists(db, "designations"):
        return {"items": [], "total": 0, "page": 1, "per_page": per_page, "pages": 0}
    sql = (
        "SELECT d.*, c.company_code, c.company_name, "
        "dept.department_code, dept.department_name "
        "FROM designations d "
        "LEFT JOIN companies c ON d.company_id = c.id "
        "LEFT JOIN departments dept ON d.department_id = dept.id "
        "WHERE 1=1"
    )
    count_sql = "SELECT COUNT(*) FROM designations d WHERE 1=1"
    params: list[Any] = []
    if not include_deleted:
        sql += " AND COALESCE(d.is_deleted,0)=0"
        count_sql += " AND COALESCE(d.is_deleted,0)=0"
    if customer_id is not None:
        sql += " AND (d.customer_id IS ? OR (d.customer_id IS NULL AND ? IS NULL))"
        count_sql += " AND (d.customer_id IS ? OR (d.customer_id IS NULL AND ? IS NULL))"
        params.extend([customer_id, customer_id])
    if company_id:
        sql += " AND d.company_id=?"
        count_sql += " AND d.company_id=?"
        params.append(company_id)
    if department_id:
        sql += " AND d.department_id=?"
        count_sql += " AND d.department_id=?"
        params.append(department_id)
    if status:
        sql += " AND d.status=?"
        count_sql += " AND d.status=?"
        params.append(status)
    if workflow_role:
        sql += " AND d.workflow_role_default=?"
        count_sql += " AND d.workflow_role_default=?"
        params.append(workflow_role)
    if search:
        clause = (
            " AND (d.designation_code LIKE ? OR d.designation_name LIKE ? "
            "OR d.designation_short_name LIKE ? OR d.description LIKE ?)"
        )
        sql += clause
        count_sql += clause
        like = f"%{search}%"
        params.extend([like, like, like, like])
    sort_col = sort_by if sort_by in DESIGNATION_SORT_COLUMNS else "designation_name"
    sort_col = f"d.{sort_col}"
    direction = "DESC" if str(sort_dir).lower() == "desc" else "ASC"
    sql += f" ORDER BY {sort_col} {direction}, d.id DESC"
    per_page = max(1, min(int(per_page or 25), 10000))
    page = max(1, int(page or 1))
    offset = (page - 1) * per_page
    sql += " LIMIT ? OFFSET ?"
    total = int(db.execute(count_sql, params).fetchone()[0])
    rows = db.execute(sql, [*params, per_page, offset]).fetchall()
    items = [dict(r) for r in rows]
    pages = (total + per_page - 1) // per_page if total else 0
    return {"items": items, "total": total, "page": page, "per_page": per_page, "pages": pages}


def get_designation_master(db, designation_id: int, *, include_deleted: bool = False) -> dict[str, Any] | None:
    if not designation_id:
        return None
    row = db.execute(
        """
        SELECT d.*, c.company_code, c.company_name,
        dept.department_code, dept.department_name
        FROM designations d
        LEFT JOIN companies c ON d.company_id = c.id
        LEFT JOIN departments dept ON d.department_id = dept.id
        WHERE d.id=?
        """
        + ("" if include_deleted else " AND COALESCE(d.is_deleted,0)=0"),
        (designation_id,),
    ).fetchone()
    return dict(row) if row else None


def _parse_designation_form(form) -> dict[str, Any]:
    dept_raw = form.get("department_id")
    return {
        "company_id": int(form.get("company_id") or 0),
        "department_id": int(dept_raw) if dept_raw not in (None, "", "0") else None,
        "designation_code": (form.get("designation_code") or "").strip(),
        "designation_name": (form.get("designation_name") or "").strip(),
        "designation_short_name": (form.get("designation_short_name") or "").strip(),
        "grade_level": (form.get("grade_level") or "").strip(),
        "workflow_role_default": (form.get("workflow_role_default") or "None").strip(),
        "description": (form.get("description") or "").strip(),
        "status": (form.get("status") or "Active").strip(),
    }


def designation_has_references(db, designation_id: int) -> bool:
    for table, column in DESIGNATION_REFERENCE_TABLES:
        if not _table_exists(db, table):
            continue
        if not _column_exists(db, table, column):
            continue
        row = db.execute(
            f"SELECT 1 FROM {table} WHERE {column}=? LIMIT 1",
            (designation_id,),
        ).fetchone()
        if row:
            return True
    return False


def _column_exists(db, table: str, column: str) -> bool:
    cols = {r[1] for r in db.execute(f"PRAGMA table_info({table})").fetchall()}
    return column in cols


def save_designation_master(
    db,
    form,
    username: str,
    designation_id: int | None = None,
    *,
    customer_id: int | None = None,
) -> int:
    data = _parse_designation_form(form)
    company_id = data["company_id"]
    if not company_id:
        raise ValueError("Company is required.")
    if not get_company(db, company_id):
        raise ValueError("Selected company was not found.")
    if data["department_id"]:
        dept = get_department_master(db, data["department_id"])
        if not dept:
            raise ValueError("Selected department was not found.")
        if int(dept.get("company_id") or 0) != company_id:
            raise ValueError("Department must belong to the selected company.")
    if not data["designation_code"]:
        data["designation_code"] = _next_designation_code(db, company_id)
    if data["status"] not in DESIGNATION_STATUSES:
        raise ValueError("Select a valid status.")
    if data["workflow_role_default"] not in WORKFLOW_ROLE_DEFAULTS:
        raise ValueError("Select a valid workflow role default.")
    validate_designation_uniqueness(
        db,
        company_id=company_id,
        designation_code=data["designation_code"],
        designation_name=data["designation_name"],
        designation_id=designation_id,
    )
    now = _now_ts()
    core = (
        data["company_id"],
        data["department_id"],
        data["designation_code"],
        data["designation_name"],
        data["designation_short_name"],
        data["grade_level"],
        data["workflow_role_default"],
        data["description"],
        data["status"],
    )
    if designation_id:
        existing = get_designation_master(db, designation_id, include_deleted=True)
        if not existing:
            raise ValueError("Designation not found.")
        db.execute(
            """
            UPDATE designations SET company_id=?, department_id=?, designation_code=?,
            designation_name=?, designation_short_name=?, grade_level=?, workflow_role_default=?,
            description=?, status=?, modified_by=?, modified_at=? WHERE id=?
            """,
            (*core, username, now, designation_id),
        )
        if customer_id is not None:
            db.execute("UPDATE designations SET customer_id=? WHERE id=?", (customer_id, designation_id))
        log_designation_field_changes(
            db, existing, get_designation_master(db, designation_id, include_deleted=True), username
        )
        return designation_id
    insert_cols = (
        "company_id, department_id, designation_code, designation_name, designation_short_name, "
        "grade_level, workflow_role_default, description, status, approval_status, "
        "created_by, created_at, modified_by, modified_at"
    )
    placeholders = ",".join(["?"] * 14)
    vals = (*core, "Draft", username, now, username, now)
    if customer_id is not None:
        cur = db.execute(
            f"INSERT INTO designations({insert_cols}, customer_id) VALUES({placeholders},?)",
            (*vals, customer_id),
        )
    else:
        cur = db.execute(f"INSERT INTO designations({insert_cols}) VALUES({placeholders})", vals)
    new_id = int(cur.lastrowid)
    log_designation_audit(
        db, new_id, "create", username, remarks=f"Created designation {data['designation_code']}"
    )
    return new_id


def soft_delete_designation_master(db, designation_id: int, username: str) -> None:
    row = get_designation_master(db, designation_id, include_deleted=True)
    if not row:
        raise ValueError("Designation not found.")
    if row.get("is_deleted"):
        return
    if designation_has_references(db, designation_id):
        raise ValueError("Cannot delete designation — it is assigned to users, staff, or workflow.")
    now = _now_ts()
    db.execute(
        """
        UPDATE designations SET is_deleted=1, deleted_by=?, deleted_at=?, status='Inactive',
        modified_by=?, modified_at=? WHERE id=?
        """,
        (username, now, username, now, designation_id),
    )
    log_designation_audit(
        db, designation_id, "soft_delete", username, remarks=f"Soft-deleted {row.get('designation_code')}"
    )


def activate_designation_master(db, designation_id: int, username: str) -> None:
    if not get_designation_master(db, designation_id):
        raise ValueError("Designation not found.")
    now = _now_ts()
    db.execute(
        "UPDATE designations SET status='Active', modified_by=?, modified_at=? WHERE id=?",
        (username, now, designation_id),
    )
    log_designation_audit(
        db, designation_id, "activate", username, field_name="status", old_value="Inactive", new_value="Active"
    )


def deactivate_designation_master(db, designation_id: int, username: str) -> None:
    if not get_designation_master(db, designation_id):
        raise ValueError("Designation not found.")
    now = _now_ts()
    db.execute(
        "UPDATE designations SET status='Inactive', modified_by=?, modified_at=? WHERE id=?",
        (username, now, designation_id),
    )
    log_designation_audit(
        db, designation_id, "deactivate", username, field_name="status", old_value="Active", new_value="Inactive"
    )


def approve_designation_master(db, designation_id: int, username: str) -> None:
    if not get_designation_master(db, designation_id):
        raise ValueError("Designation not found.")
    now = _now_ts()
    db.execute(
        """
        UPDATE designations SET approval_status='Approved', approved_by=?, approved_at=?,
        modified_by=?, modified_at=? WHERE id=?
        """,
        (username, now, username, now, designation_id),
    )
    log_designation_audit(db, designation_id, "approve", username, remarks="Designation approved")


def reject_designation_master(db, designation_id: int, username: str, remarks: str = "") -> None:
    if not get_designation_master(db, designation_id):
        raise ValueError("Designation not found.")
    now = _now_ts()
    db.execute(
        "UPDATE designations SET approval_status='Rejected', modified_by=?, modified_at=? WHERE id=?",
        (username, now, designation_id),
    )
    log_designation_audit(db, designation_id, "reject", username, remarks=remarks or "Designation rejected")


def user_can_designation_master(
    db,
    user_id: int | None,
    action: str,
    *,
    is_admin: bool = False,
) -> bool:
    if is_admin:
        return True
    if not user_id:
        return False
    action_map = {"deactivate": "edit", "activate": "edit"}
    check = action_map.get(action, action)
    try:
        from user_permission_service import (
            empty_permission_actions,
            ensure_user_tab_permissions_schema,
            normalize_permission_actions,
        )

        ensure_user_tab_permissions_schema(db)
        row = db.execute(
            """
            SELECT granted, action_flags FROM user_tab_permissions
            WHERE user_id=? AND granted=1 AND endpoint='designation_master'
            LIMIT 1
            """,
            (user_id,),
        ).fetchone()
        if not row:
            return False
        raw_flags = row["action_flags"] if hasattr(row, "keys") else row[1]
        actions = normalize_permission_actions(
            json.loads(raw_flags) if raw_flags else empty_permission_actions()
        )
        if check == "import":
            return bool(actions.get("import") or actions.get("create"))
        if check == "delete":
            return bool(actions.get("delete") or actions.get("edit"))
        return bool(actions.get(check))
    except Exception:
        return False


def designations_for_export(db, *, include_deleted: bool = False, **filters) -> list[dict[str, Any]]:
    listing = list_designations_master(db, include_deleted=include_deleted, per_page=10000, **filters)
    return [{col: item.get(col, "") for col in DESIGNATION_EXPORT_COLUMNS} for item in listing["items"]]


def export_designations_excel(db, **filters) -> BytesIO:
    from openpyxl import Workbook

    rows = designations_for_export(db, **filters)
    wb = Workbook()
    ws = wb.active
    ws.title = "Designations"
    headers = list(DESIGNATION_EXPORT_COLUMNS)
    ws.append(headers)
    for row in rows:
        ws.append([row.get(h, "") for h in headers])
    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def export_designations_csv(db, **filters) -> str:
    rows = designations_for_export(db, **filters)
    si = io.StringIO()
    writer = csv.writer(si)
    headers = list(DESIGNATION_EXPORT_COLUMNS)
    writer.writerow(headers)
    for row in rows:
        writer.writerow([row.get(h, "") for h in headers])
    return si.getvalue()


def export_designations_pdf(db, *, report_title: str = "Designation Master Report", **filters) -> BytesIO:
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.pdfgen import canvas

    rows = designations_for_export(db, **filters)
    buf = BytesIO()
    page_size = landscape(A4)
    c = canvas.Canvas(buf, pagesize=page_size)
    _, height = page_size
    y = height - 40
    c.setFont("Helvetica-Bold", 14)
    c.drawString(40, y, f"MAXEK ERP — {report_title}")
    y -= 24
    c.setFont("Helvetica", 9)
    for row in rows[:250]:
        line = (
            f"{row.get('company_code')} | {row.get('designation_code')} | {row.get('designation_name')} | "
            f"{row.get('workflow_role_default')} | {row.get('status')}"
        )
        if y < 40:
            c.showPage()
            y = height - 40
            c.setFont("Helvetica", 9)
        c.drawString(40, y, line[:130])
        y -= 14
    c.save()
    buf.seek(0)
    return buf


def designation_report(db, report_key: str, **filters) -> list[dict[str, Any]]:
    key = (report_key or "list").lower().strip()
    if key == "active":
        filters["status"] = "Active"
    elif key == "inactive":
        filters["status"] = "Inactive"
    elif key == "workflow":
        listing = list_designations_master(db, per_page=5000, **filters)
        return [r for r in listing["items"] if r.get("workflow_role_default") not in (None, "", "None")]
    listing = list_designations_master(db, per_page=5000, **filters)
    return listing["items"]


def designation_import_template() -> BytesIO:
    from bulk_import_service import build_xlsx_template

    return build_xlsx_template(
        [
            "Company Code",
            "Department Code",
            "Designation Code",
            "Designation Name",
            "Short Name",
            "Grade Level",
            "Workflow Role",
            "Description",
            "Status",
        ],
        [
            "CO-2026-0001",
            "DEPT-HR-01",
            "DG-SE-01",
            "Site Engineer",
            "SE",
            "L3",
            "Maker",
            "Site engineering role",
            "Active",
        ],
    )


def ai_validate_designation(
    db, designation_id: int | None = None, form: dict | None = None
) -> dict[str, Any]:
    data = dict(form or {})
    if designation_id and not form:
        row = get_designation_master(db, designation_id)
        if row:
            data = dict(row)
    issues: list[str] = []
    suggestions: list[str] = []
    duplicates: list[str] = []
    missing: list[str] = []
    if not int(data.get("company_id") or 0):
        issues.append("Company is required.")
        missing.append("company_id")
    name = (data.get("designation_name") or "").strip()
    if not name:
        issues.append("Designation name is required.")
        missing.append("designation_name")
    else:
        try:
            validate_designation_uniqueness(
                db,
                company_id=int(data.get("company_id") or 0) or 1,
                designation_code=(data.get("designation_code") or name).strip(),
                designation_name=name,
                designation_id=designation_id,
            )
        except ValueError as exc:
            duplicates.append(str(exc))
    if not (data.get("grade_level") or "").strip():
        suggestions.append("Add a grade level for payroll and hierarchy reporting.")
    result = {
        "ok": not issues and not duplicates,
        "issues": issues,
        "duplicates": duplicates,
        "suggestions": suggestions,
        "missing": missing,
    }
    try:
        from ai_service import chat_completion_json

        ai = chat_completion_json(
            "Validate ERP designation records. Return JSON with keys: issues, suggestions.",
            json.dumps({"designation": data, "rule_findings": result}, ensure_ascii=False),
            max_tokens=400,
        )
        for key in ("issues", "suggestions"):
            extra = ai.get(key) or []
            if isinstance(extra, list):
                result[key].extend(str(x) for x in extra if x)
        result["ok"] = not result["issues"] and not result["duplicates"]
        result["ai_enriched"] = True
    except Exception:
        result["ai_enriched"] = False
    return result


def list_companies_for_designation_form(db) -> list[dict[str, Any]]:
    listing = list_companies(db, per_page=1000)
    return [
        {"id": c["id"], "company_code": c.get("company_code"), "company_name": c.get("company_name")}
        for c in listing.get("items", [])
    ]


def list_departments_for_designation_form(db, company_id: int | None = None) -> list[dict[str, Any]]:
    from department_master_service import list_departments_master

    listing = list_departments_master(db, company_id=company_id, status="Active", per_page=1000)
    return [
        {
            "id": d["id"],
            "department_code": d.get("department_code"),
            "department_name": d.get("department_name"),
            "company_id": d.get("company_id"),
        }
        for d in listing.get("items", [])
    ]


def seed_default_designations(db, company_id: int | None = None) -> None:
    """Seed standard designations for a company (idempotent)."""
    if not company_id:
        row = db.execute(
            "SELECT id FROM companies WHERE COALESCE(is_deleted,0)=0 ORDER BY id LIMIT 1"
        ).fetchone()
        if not row:
            return
        company_id = int(row[0])
    for name in DEFAULT_DESIGNATION_NAMES:
        existing = db.execute(
            "SELECT id FROM designations WHERE designation_name=? AND COALESCE(is_deleted,0)=0",
            (name,),
        ).fetchone()
        if existing:
            db.execute(
                "UPDATE designations SET company_id=? WHERE id=? AND company_id IS NULL",
                (company_id, existing[0]),
            )
            continue
        code = re.sub(r"[^A-Za-z0-9]+", "-", name).strip("-").upper()[:16]
        try:
            db.execute(
                """
                INSERT INTO designations(
                    company_id, designation_code, designation_name, status, approval_status
                ) VALUES(?,?,?,?, 'Approved')
                """,
                (company_id, f"DG-{code}", name, "Active"),
            )
        except Exception:
            pass
