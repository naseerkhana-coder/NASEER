"""User Management (MODULE-004) — ERP users, security, and access."""

from __future__ import annotations

import csv
import io
import json
import os
import re
import secrets
import uuid
from datetime import datetime, timedelta
from io import BytesIO
from typing import Any, Callable

import bcrypt

from branch_master_service import ensure_branch_master_schema, get_branch_master
from company_master_service import (
    COMPANY_CURRENCIES,
    COMPANY_TIMEZONES,
    _ensure_column,
    _now_ts,
    _table_exists,
    ensure_company_master_schema,
    get_company,
    list_companies,
    validate_email,
    validate_phone,
)
from department_master_service import ensure_department_master_schema, get_department_master

USER_STATUSES = ("Active", "Inactive")
USER_SYSTEM_ROLES = (
    "User",
    "Guest",
    "Admin",
    "Customer Admin",
    "Manager",
    "Cashier",
    "Accountant",
    "Store Keeper",
    "HR User",
)
USER_WORKFLOW_ROLES = ("Maker", "Checker", "Approver", "Administrator")
USER_SORT_COLUMNS = (
    "username",
    "employee_name",
    "email",
    "mobile",
    "role",
    "status",
    "last_login",
    "created_at",
    "company_id",
)
USER_EXPORT_COLUMNS = (
    "username",
    "display_name",
    "first_name",
    "last_name",
    "email",
    "mobile",
    "company_code",
    "branch_code",
    "department_name",
    "role",
    "workflow_role",
    "designation_name",
    "status",
    "account_locked",
    "last_login",
    "password_expiry",
)
USER_AUDIT_FIELDS = (
    "username",
    "employee_name",
    "first_name",
    "middle_name",
    "last_name",
    "email",
    "mobile",
    "alternative_mobile",
    "company_id",
    "branch_id",
    "department_id",
    "staff_id",
    "role",
    "workflow_role",
    "designation_id",
    "reporting_manager",
    "status",
    "account_locked",
    "language",
    "timezone",
    "date_format",
    "currency",
)
PASSWORD_MIN_LENGTH = 8
PASSWORD_HISTORY_LIMIT = 5
MAX_LOGIN_ATTEMPTS = 5
PASSWORD_EXPIRY_DAYS = 90
PROFILE_PHOTO_EXTENSIONS = frozenset({".jpg", ".jpeg", ".png", ".webp"})
DATE_FORMATS = ("DD/MM/YYYY", "MM/DD/YYYY", "YYYY-MM-DD")
LANGUAGES = ("en", "hi", "ar")
EMAIL_RE = re.compile(r"^[^@\s]+@[^@\s]+\.[^@\s]+$")
MOBILE_RE = re.compile(r"^\+?[\d\s\-()]{7,20}$")


def hash_user_password(plain_password: str) -> str:
    return bcrypt.hashpw(plain_password.encode("utf-8"), bcrypt.gensalt()).decode("utf-8")


def verify_user_password(stored: str, provided: str) -> bool:
    if not stored or not provided:
        return False
    if stored.startswith("$2a$") or stored.startswith("$2b$") or stored.startswith("$2y$"):
        try:
            return bcrypt.checkpw(provided.encode("utf-8"), stored.encode("utf-8"))
        except (ValueError, TypeError):
            return False
    return stored == provided


def validate_password_policy(password: str) -> None:
    if len(password or "") < PASSWORD_MIN_LENGTH:
        raise ValueError(f"Password must be at least {PASSWORD_MIN_LENGTH} characters.")
    if not re.search(r"[A-Z]", password):
        raise ValueError("Password must include an uppercase letter.")
    if not re.search(r"[a-z]", password):
        raise ValueError("Password must include a lowercase letter.")
    if not re.search(r"\d", password):
        raise ValueError("Password must include a number.")
    if not re.search(r"[^A-Za-z0-9]", password):
        raise ValueError("Password must include a special character.")


def validate_user_mobile(value: str) -> None:
    text = (value or "").strip()
    if not text:
        return
    if not MOBILE_RE.match(text):
        raise ValueError("Enter a valid mobile number.")
    validate_phone(text)


def ensure_user_master_schema(db) -> None:
    """Extend users table for MODULE-004 (idempotent)."""
    ensure_company_master_schema(db)
    ensure_branch_master_schema(db)
    ensure_department_master_schema(db)
    db.execute(
        """
        CREATE TABLE IF NOT EXISTS users(
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            username TEXT UNIQUE,
            password TEXT,
            role TEXT,
            status TEXT DEFAULT 'Active'
        )
        """
    )
    for col, ctype in (
        ("company_id", "INTEGER"),
        ("branch_id", "INTEGER"),
        ("department_id", "INTEGER"),
        ("staff_id", "INTEGER"),
        ("first_name", "TEXT"),
        ("middle_name", "TEXT"),
        ("last_name", "TEXT"),
        ("employee_name", "TEXT"),
        ("email", "TEXT"),
        ("mobile", "TEXT"),
        ("alternative_mobile", "TEXT"),
        ("profile_photo", "TEXT"),
        ("digital_signature", "TEXT"),
        ("role_id", "INTEGER"),
        ("designation_id", "INTEGER"),
        ("workflow_role", "TEXT"),
        ("reporting_manager", "TEXT"),
        ("default_dashboard", "TEXT"),
        ("language", "TEXT DEFAULT 'en'"),
        ("timezone", "TEXT DEFAULT 'Asia/Kolkata'"),
        ("date_format", "TEXT DEFAULT 'DD/MM/YYYY'"),
        ("currency", "TEXT DEFAULT 'INR'"),
        ("last_login", "TEXT"),
        ("last_password_change", "TEXT"),
        ("password_expiry", "TEXT"),
        ("login_attempt_count", "INTEGER DEFAULT 0"),
        ("account_locked", "INTEGER DEFAULT 0"),
        ("account_active", "INTEGER DEFAULT 1"),
        ("email_verified", "INTEGER DEFAULT 0"),
        ("mobile_verified", "INTEGER DEFAULT 0"),
        ("mfa_enabled", "INTEGER DEFAULT 0"),
        ("department", "TEXT"),
        ("customer_id", "INTEGER"),
        ("created_by", "TEXT"),
        ("created_at", "TEXT"),
        ("modified_by", "TEXT"),
        ("modified_at", "TEXT"),
        ("is_deleted", "INTEGER DEFAULT 0"),
        ("deleted_by", "TEXT"),
        ("deleted_at", "TEXT"),
    ):
        _ensure_column(db, "users", col, ctype)
    db.execute(
        """
        CREATE TABLE IF NOT EXISTS user_password_history(
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            user_id INTEGER NOT NULL,
            password_hash TEXT NOT NULL,
            changed_at TEXT NOT NULL
        )
        """
    )
    db.execute(
        """
        CREATE TABLE IF NOT EXISTS user_password_reset_tokens(
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            user_id INTEGER NOT NULL,
            token TEXT UNIQUE NOT NULL,
            expires_at TEXT NOT NULL,
            used INTEGER DEFAULT 0,
            created_at TEXT NOT NULL
        )
        """
    )
    try:
        from audit_trail_service import ensure_audit_schema

        ensure_audit_schema(db)
    except Exception:
        pass
    db.execute(
        """
        CREATE TABLE IF NOT EXISTS designations(
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            designation_name TEXT,
            status TEXT DEFAULT 'Active'
        )
        """
    )


def _display_name(data: dict[str, Any]) -> str:
    explicit = (data.get("employee_name") or data.get("display_name") or "").strip()
    if explicit:
        return explicit
    parts = [
        (data.get("first_name") or "").strip(),
        (data.get("middle_name") or "").strip(),
        (data.get("last_name") or "").strip(),
    ]
    joined = " ".join(p for p in parts if p)
    return joined or (data.get("username") or "").strip()


def validate_user_uniqueness(
    db,
    *,
    username: str,
    email: str = "",
    mobile: str = "",
    customer_id: int | None = None,
    user_id: int | None = None,
) -> None:
    uname = (username or "").strip()
    if not uname:
        raise ValueError("Username is required.")
    row = db.execute(
        """
        SELECT id FROM users
        WHERE LOWER(username)=LOWER(?) AND COALESCE(is_deleted,0)=0
        AND (customer_id IS ? OR (customer_id IS NULL AND ? IS NULL))
        """,
        (uname, customer_id, customer_id),
    ).fetchone()
    if row and (not user_id or int(row[0]) != int(user_id)):
        raise ValueError(f"Username '{uname}' is already in use.")
    mail = (email or "").strip().lower()
    if mail:
        row = db.execute(
            """
            SELECT id, username FROM users
            WHERE LOWER(email)=? AND COALESCE(is_deleted,0)=0
            AND (customer_id IS ? OR (customer_id IS NULL AND ? IS NULL))
            """,
            (mail, customer_id, customer_id),
        ).fetchone()
        if row and (not user_id or int(row[0]) != int(user_id)):
            raise ValueError(f"Email '{email}' is already registered to user {row[1]}.")
    mob = (mobile or "").strip()
    if mob:
        row = db.execute(
            """
            SELECT id, username FROM users
            WHERE mobile=? AND COALESCE(is_deleted,0)=0
            AND (customer_id IS ? OR (customer_id IS NULL AND ? IS NULL))
            """,
            (mob, customer_id, customer_id),
        ).fetchone()
        if row and (not user_id or int(row[0]) != int(user_id)):
            raise ValueError(f"Mobile '{mobile}' is already registered to user {row[1]}.")


def _password_expiry_from_now() -> str:
    return (datetime.now() + timedelta(days=PASSWORD_EXPIRY_DAYS)).strftime("%Y-%m-%d %H:%M:%S")


def _check_password_history(db, user_id: int, password_hash: str) -> None:
    rows = db.execute(
        "SELECT password_hash FROM user_password_history WHERE user_id=? ORDER BY id DESC LIMIT ?",
        (user_id, PASSWORD_HISTORY_LIMIT),
    ).fetchall()
    for row in rows:
        if row[0] == password_hash:
            raise ValueError("Cannot reuse a recent password. Choose a new password.")


def _store_password_history(db, user_id: int, password_hash: str) -> None:
    db.execute(
        "INSERT INTO user_password_history(user_id, password_hash, changed_at) VALUES(?,?,?)",
        (user_id, password_hash, _now_ts()),
    )


def log_user_audit(
    db,
    user_id: int,
    action: str,
    username: str,
    *,
    field_name: str | None = None,
    old_value: str | None = None,
    new_value: str | None = None,
    remarks: str | None = None,
) -> None:
    try:
        from audit_trail_service import log_audit_event

        log_audit_event(
            db,
            record_table="users",
            record_id=user_id,
            action=action,
            changed_by=username,
            field_name=field_name,
            old_value=old_value,
            new_value=new_value,
            remarks=remarks,
        )
    except Exception:
        pass


def log_user_field_changes(
    db,
    before: dict[str, Any] | None,
    after: dict[str, Any] | None,
    actor: str,
) -> None:
    if not before or not after:
        return
    uid = int(after.get("id") or before.get("id") or 0)
    if not uid:
        return
    for field in USER_AUDIT_FIELDS:
        old_val = before.get(field)
        new_val = after.get(field)
        if str(old_val or "") != str(new_val or ""):
            log_user_audit(
                db,
                uid,
                "update",
                actor,
                field_name=field,
                old_value=str(old_val or ""),
                new_value=str(new_val or ""),
            )


def list_user_audit_trail(db, user_id: int, limit: int = 100) -> list[dict[str, Any]]:
    try:
        from audit_trail_service import list_audit_trail

        return list_audit_trail(db, "users", user_id, limit=limit)
    except Exception:
        return []


def _user_row_to_dict(row) -> dict[str, Any]:
    item = dict(row)
    item["display_name"] = item.get("employee_name") or _display_name(item)
    item["account_active"] = 0 if str(item.get("status", "")).strip() != "Active" else 1
    return item


def list_users_master(
    db,
    *,
    search: str = "",
    company_id: int | None = None,
    branch_id: int | None = None,
    department_id: int | None = None,
    status: str = "",
    locked: str = "",
    customer_id: int | None = None,
    include_deleted: bool = False,
    page: int = 1,
    per_page: int = 25,
    sort_by: str = "username",
    sort_dir: str = "asc",
) -> dict[str, Any]:
    if not _table_exists(db, "users"):
        return {"items": [], "total": 0, "page": 1, "per_page": per_page, "pages": 0}
    sql = (
        "SELECT u.*, c.company_code, c.company_name, b.branch_code, b.branch_name, "
        "d.department_name AS dept_master_name, des.designation_name "
        "FROM users u "
        "LEFT JOIN companies c ON u.company_id = c.id "
        "LEFT JOIN company_branches b ON u.branch_id = b.id "
        "LEFT JOIN departments d ON u.department_id = d.id "
        "LEFT JOIN designations des ON u.designation_id = des.id "
        "WHERE 1=1"
    )
    count_sql = "SELECT COUNT(*) FROM users u WHERE 1=1"
    params: list[Any] = []
    if not include_deleted:
        sql += " AND COALESCE(u.is_deleted,0)=0"
        count_sql += " AND COALESCE(u.is_deleted,0)=0"
    if customer_id is not None:
        sql += " AND (u.customer_id IS ? OR (u.customer_id IS NULL AND ? IS NULL))"
        count_sql += " AND (u.customer_id IS ? OR (u.customer_id IS NULL AND ? IS NULL))"
        params.extend([customer_id, customer_id])
    if company_id:
        sql += " AND u.company_id=?"
        count_sql += " AND u.company_id=?"
        params.append(company_id)
    if branch_id:
        sql += " AND u.branch_id=?"
        count_sql += " AND u.branch_id=?"
        params.append(branch_id)
    if department_id:
        sql += " AND u.department_id=?"
        count_sql += " AND u.department_id=?"
        params.append(department_id)
    if status:
        sql += " AND u.status=?"
        count_sql += " AND u.status=?"
        params.append(status)
    if locked == "1":
        sql += " AND COALESCE(u.account_locked,0)=1"
        count_sql += " AND COALESCE(u.account_locked,0)=1"
    elif locked == "0":
        sql += " AND COALESCE(u.account_locked,0)=0"
        count_sql += " AND COALESCE(u.account_locked,0)=0"
    if search:
        clause = (
            " AND (u.username LIKE ? OR u.employee_name LIKE ? OR u.email LIKE ? "
            "OR u.mobile LIKE ? OR u.first_name LIKE ? OR u.last_name LIKE ?)"
        )
        sql += clause
        count_sql += clause
        like = f"%{search}%"
        params.extend([like, like, like, like, like, like])
    sort_col = sort_by if sort_by in USER_SORT_COLUMNS else "username"
    sort_col = f"u.{sort_col}"
    direction = "DESC" if str(sort_dir).lower() == "desc" else "ASC"
    sql += f" ORDER BY {sort_col} {direction}, u.id DESC"
    per_page = max(1, min(int(per_page or 25), 10000))
    page = max(1, int(page or 1))
    offset = (page - 1) * per_page
    sql += " LIMIT ? OFFSET ?"
    total = int(db.execute(count_sql, params).fetchone()[0])
    rows = db.execute(sql, [*params, per_page, offset]).fetchall()
    items = [_user_row_to_dict(r) for r in rows]
    for item in items:
        item["user_type"] = "Staff-linked" if item.get("staff_id") else "Standalone"
        item["department_name"] = item.get("dept_master_name") or item.get("department") or ""
    pages = (total + per_page - 1) // per_page if total else 0
    return {"items": items, "total": total, "page": page, "per_page": per_page, "pages": pages}


def get_user_master(db, user_id: int, *, include_deleted: bool = False) -> dict[str, Any] | None:
    if not user_id:
        return None
    row = db.execute(
        """
        SELECT u.*, c.company_code, c.company_name, b.branch_code, b.branch_name,
        d.department_name AS dept_master_name, des.designation_name
        FROM users u
        LEFT JOIN companies c ON u.company_id = c.id
        LEFT JOIN company_branches b ON u.branch_id = b.id
        LEFT JOIN departments d ON u.department_id = d.id
        LEFT JOIN designations des ON u.designation_id = des.id
        WHERE u.id=?
        """
        + ("" if include_deleted else " AND COALESCE(u.is_deleted,0)=0"),
        (user_id,),
    ).fetchone()
    if not row:
        return None
    item = _user_row_to_dict(row)
    item["user_type"] = "Staff-linked" if item.get("staff_id") else "Standalone"
    item["department_name"] = item.get("dept_master_name") or item.get("department") or ""
    return item


def _parse_user_form(form) -> dict[str, Any]:
    dept_raw = form.get("department_id")
    branch_raw = form.get("branch_id")
    staff_raw = form.get("staff_id")
    desig_raw = form.get("designation_id")
    return {
        "company_id": int(form.get("company_id") or 0),
        "branch_id": int(branch_raw) if branch_raw not in (None, "", "0") else None,
        "department_id": int(dept_raw) if dept_raw not in (None, "", "0") else None,
        "staff_id": int(staff_raw) if staff_raw not in (None, "", "0") else None,
        "username": (form.get("username") or "").strip(),
        "first_name": (form.get("first_name") or "").strip(),
        "middle_name": (form.get("middle_name") or "").strip(),
        "last_name": (form.get("last_name") or "").strip(),
        "employee_name": (form.get("display_name") or form.get("employee_name") or "").strip(),
        "email": (form.get("email") or "").strip(),
        "mobile": (form.get("mobile") or "").strip(),
        "alternative_mobile": (form.get("alternative_mobile") or "").strip(),
        "role": (form.get("system_role") or form.get("role") or "User").strip(),
        "workflow_role": (form.get("workflow_role") or "Maker").strip(),
        "designation_id": int(desig_raw) if desig_raw not in (None, "", "0") else None,
        "reporting_manager": (form.get("reporting_manager") or "").strip(),
        "default_dashboard": (form.get("default_dashboard") or "").strip(),
        "language": (form.get("language") or "en").strip(),
        "timezone": (form.get("timezone") or "Asia/Kolkata").strip(),
        "date_format": (form.get("date_format") or "DD/MM/YYYY").strip(),
        "currency": (form.get("currency") or "INR").strip(),
        "status": (form.get("status") or "Active").strip(),
    }


def _validate_org_links(db, data: dict[str, Any]) -> None:
    company_id = data["company_id"]
    if not company_id:
        raise ValueError("Company is required.")
    if not get_company(db, company_id):
        raise ValueError("Selected company was not found.")
    if not data["branch_id"]:
        raise ValueError("Branch is required.")
    branch = get_branch_master(db, data["branch_id"])
    if not branch:
        raise ValueError("Selected branch was not found.")
    if int(branch.get("company_id") or 0) != company_id:
        raise ValueError("Branch must belong to the selected company.")
    if not data["department_id"]:
        raise ValueError("Department is required.")
    dept = get_department_master(db, data["department_id"])
    if not dept:
        raise ValueError("Selected department was not found.")
    if int(dept.get("company_id") or 0) != company_id:
        raise ValueError("Department must belong to the selected company.")


def _notify_user_event(db, user_id: int, message: str, notification_type: str) -> None:
    try:
        from workflow_service import create_notification

        create_notification(db, user_id, message, notification_type, record_table="users", record_id=user_id)
    except Exception:
        pass


def save_user_master(
    db,
    form,
    actor: str,
    user_id: int | None = None,
    *,
    customer_id: int | None = None,
    password: str | None = None,
    assert_user_limit_fn: Callable | None = None,
) -> int:
    data = _parse_user_form(form)
    _validate_org_links(db, data)
    if not data["username"]:
        raise ValueError("Username is required.")
    display = _display_name(data)
    if not display:
        raise ValueError("Display name or name fields are required.")
    if data["status"] not in USER_STATUSES:
        raise ValueError("Select a valid status.")
    if data["workflow_role"] not in USER_WORKFLOW_ROLES:
        raise ValueError("Select a valid workflow role.")
    if data["workflow_role"] == "Administrator":
        data["role"] = "Admin"
    if data["email"]:
        validate_email(data["email"])
    validate_user_mobile(data["mobile"])
    validate_user_mobile(data["alternative_mobile"])
    if data["timezone"] not in COMPANY_TIMEZONES:
        raise ValueError("Select a valid timezone.")
    if data["currency"] not in COMPANY_CURRENCIES:
        raise ValueError("Select a valid currency.")
    if data["date_format"] not in DATE_FORMATS:
        raise ValueError("Select a valid date format.")
    validate_user_uniqueness(
        db,
        username=data["username"],
        email=data["email"],
        mobile=data["mobile"],
        customer_id=customer_id,
        user_id=user_id,
    )
    dept = get_department_master(db, data["department_id"])
    department_text = dept.get("department_name") if dept else ""
    now = _now_ts()
    account_active = 1 if data["status"] == "Active" else 0
    core = (
        data["company_id"],
        data["branch_id"],
        data["department_id"],
        data["staff_id"],
        data["username"],
        data["first_name"],
        data["middle_name"],
        data["last_name"],
        display,
        data["email"],
        data["mobile"],
        data["alternative_mobile"],
        data["role"],
        data["workflow_role"],
        data["designation_id"],
        data["reporting_manager"],
        data["default_dashboard"],
        data["language"],
        data["timezone"],
        data["date_format"],
        data["currency"],
        department_text,
        data["status"],
        account_active,
    )
    if user_id:
        existing = get_user_master(db, user_id, include_deleted=True)
        if not existing:
            raise ValueError("User not found.")
        db.execute(
            """
            UPDATE users SET company_id=?, branch_id=?, department_id=?, staff_id=?, username=?,
            first_name=?, middle_name=?, last_name=?, employee_name=?, email=?, mobile=?,
            alternative_mobile=?, role=?, workflow_role=?, designation_id=?, reporting_manager=?,
            default_dashboard=?, language=?, timezone=?, date_format=?, currency=?, department=?,
            status=?, account_active=?, modified_by=?, modified_at=? WHERE id=?
            """,
            (*core, actor, now, user_id),
        )
        if customer_id is not None:
            db.execute("UPDATE users SET customer_id=? WHERE id=?", (customer_id, user_id))
        if password:
            validate_password_policy(password)
            pwd_hash = hash_user_password(password)
            _check_password_history(db, user_id, pwd_hash)
            db.execute(
                """
                UPDATE users SET password=?, last_password_change=?, password_expiry=?,
                login_attempt_count=0, account_locked=0 WHERE id=?
                """,
                (pwd_hash, now, _password_expiry_from_now(), user_id),
            )
            _store_password_history(db, user_id, pwd_hash)
            log_user_audit(db, user_id, "reset_password", actor, remarks="Password updated")
        log_user_field_changes(db, existing, get_user_master(db, user_id, include_deleted=True), actor)
        return user_id
    if not password:
        raise ValueError("Password is required for new users.")
    validate_password_policy(password)
    if customer_id and assert_user_limit_fn:
        assert_user_limit_fn(db, customer_id)
    pwd_hash = hash_user_password(password)
    insert_cols = (
        "company_id, branch_id, department_id, staff_id, username, password, first_name, middle_name, "
        "last_name, employee_name, email, mobile, alternative_mobile, role, workflow_role, designation_id, "
        "reporting_manager, default_dashboard, language, timezone, date_format, currency, department, "
        "status, account_active, last_password_change, password_expiry, created_by, created_at, "
        "modified_by, modified_at"
    )
    placeholders = ",".join(["?"] * 31)
    vals = (
        *core[:4],
        data["username"],
        pwd_hash,
        *core[5:],
        now,
        _password_expiry_from_now(),
        actor,
        now,
        actor,
        now,
    )
    if customer_id is not None:
        cur = db.execute(
            f"INSERT INTO users({insert_cols}, customer_id) VALUES({placeholders},?)",
            (*vals, customer_id),
        )
    else:
        cur = db.execute(f"INSERT INTO users({insert_cols}) VALUES({placeholders})", vals)
    new_id = int(cur.lastrowid)
    _store_password_history(db, new_id, pwd_hash)
    log_user_audit(db, new_id, "create", actor, remarks=f"Created user {data['username']}")
    _notify_user_event(db, new_id, "Welcome to MAXEK ERP. Your account has been created.", "user_welcome")
    return new_id


def soft_delete_user_master(db, user_id: int, actor: str) -> None:
    row = get_user_master(db, user_id, include_deleted=True)
    if not row:
        raise ValueError("User not found.")
    if row.get("is_deleted"):
        return
    now = _now_ts()
    db.execute(
        """
        UPDATE users SET is_deleted=1, deleted_by=?, deleted_at=?, status='Inactive',
        account_active=0, modified_by=?, modified_at=? WHERE id=?
        """,
        (actor, now, actor, now, user_id),
    )
    log_user_audit(db, user_id, "soft_delete", actor, remarks=f"Soft-deleted user {row.get('username')}")


def activate_user_master(db, user_id: int, actor: str) -> None:
    if not get_user_master(db, user_id):
        raise ValueError("User not found.")
    now = _now_ts()
    db.execute(
        """
        UPDATE users SET status='Active', account_active=1, modified_by=?, modified_at=? WHERE id=?
        """,
        (actor, now, user_id),
    )
    log_user_audit(db, user_id, "activate", actor, field_name="status", old_value="Inactive", new_value="Active")
    _notify_user_event(db, user_id, "Your MAXEK ERP account has been activated.", "user_activated")


def deactivate_user_master(db, user_id: int, actor: str) -> None:
    if not get_user_master(db, user_id):
        raise ValueError("User not found.")
    now = _now_ts()
    db.execute(
        """
        UPDATE users SET status='Inactive', account_active=0, modified_by=?, modified_at=? WHERE id=?
        """,
        (actor, now, user_id),
    )
    log_user_audit(db, user_id, "deactivate", actor, field_name="status", old_value="Active", new_value="Inactive")


def lock_user_master(db, user_id: int, actor: str, *, reason: str = "") -> None:
    if not get_user_master(db, user_id):
        raise ValueError("User not found.")
    now = _now_ts()
    db.execute(
        "UPDATE users SET account_locked=1, modified_by=?, modified_at=? WHERE id=?",
        (actor, now, user_id),
    )
    log_user_audit(db, user_id, "lock", actor, remarks=reason or "Account locked")
    _notify_user_event(db, user_id, "Your account has been locked. Contact your administrator.", "user_locked")


def unlock_user_master(db, user_id: int, actor: str) -> None:
    if not get_user_master(db, user_id):
        raise ValueError("User not found.")
    now = _now_ts()
    db.execute(
        """
        UPDATE users SET account_locked=0, login_attempt_count=0, modified_by=?, modified_at=? WHERE id=?
        """,
        (actor, now, user_id),
    )
    log_user_audit(db, user_id, "unlock", actor, remarks="Account unlocked")


def reset_user_password(
    db,
    user_id: int,
    new_password: str,
    actor: str,
    *,
    notify: bool = True,
) -> None:
    if not get_user_master(db, user_id):
        raise ValueError("User not found.")
    validate_password_policy(new_password)
    pwd_hash = hash_user_password(new_password)
    _check_password_history(db, user_id, pwd_hash)
    now = _now_ts()
    db.execute(
        """
        UPDATE users SET password=?, last_password_change=?, password_expiry=?,
        login_attempt_count=0, account_locked=0, modified_by=?, modified_at=? WHERE id=?
        """,
        (pwd_hash, now, _password_expiry_from_now(), actor, now, user_id),
    )
    _store_password_history(db, user_id, pwd_hash)
    log_user_audit(db, user_id, "reset_password", actor, remarks="Password reset by administrator")
    if notify:
        _notify_user_event(db, user_id, "Your password has been reset by an administrator.", "password_reset")


def create_password_reset_token(db, user_id: int) -> str:
    token = secrets.token_urlsafe(32)
    expires = (datetime.now() + timedelta(hours=24)).strftime("%Y-%m-%d %H:%M:%S")
    db.execute(
        """
        INSERT INTO user_password_reset_tokens(user_id, token, expires_at, used, created_at)
        VALUES(?,?,?,0,?)
        """,
        (user_id, token, expires, _now_ts()),
    )
    return token


def record_failed_login(db, username: str) -> None:
    row = db.execute(
        "SELECT id, login_attempt_count FROM users WHERE username=? AND COALESCE(is_deleted,0)=0",
        (username.strip(),),
    ).fetchone()
    if not row:
        return
    count = int(row[1] or 0) + 1
    locked = 1 if count >= MAX_LOGIN_ATTEMPTS else 0
    db.execute(
        "UPDATE users SET login_attempt_count=?, account_locked=? WHERE id=?",
        (count, locked, row[0]),
    )
    if locked:
        log_user_audit(db, int(row[0]), "lock", "system", remarks="Auto-locked after failed login attempts")
        _notify_user_event(db, int(row[0]), "Account locked due to failed login attempts.", "user_locked")


def record_successful_login(db, user_id: int) -> None:
    now = _now_ts()
    db.execute(
        """
        UPDATE users SET last_login=?, login_attempt_count=0, account_locked=0 WHERE id=?
        """,
        (now, user_id),
    )
    log_user_audit(db, user_id, "login", "system", remarks="Successful login")


def record_user_logout(db, user_id: int, actor: str = "system") -> None:
    log_user_audit(db, user_id, "logout", actor, remarks="User logged out")


def user_is_login_allowed(user_row) -> bool:
    if user_row is None:
        return False
    keys = user_row.keys() if hasattr(user_row, "keys") else []
    if "is_deleted" in keys and user_row["is_deleted"]:
        return False
    if "account_locked" in keys and int(user_row["account_locked"] or 0) == 1:
        return False
    status = str(user_row["status"] if "status" in keys else "Active").strip()
    if status.lower() != "active":
        return False
    if "password_expiry" in keys and user_row["password_expiry"]:
        try:
            expiry = datetime.strptime(str(user_row["password_expiry"]), "%Y-%m-%d %H:%M:%S")
            if datetime.now() > expiry:
                return False
        except ValueError:
            pass
    return True


def save_user_profile_photo(db, user_id: int, filename: str, actor: str) -> None:
    before = get_user_master(db, user_id, include_deleted=True)
    db.execute(
        "UPDATE users SET profile_photo=?, modified_by=?, modified_at=? WHERE id=?",
        (filename, actor, _now_ts(), user_id),
    )
    log_user_field_changes(db, before, get_user_master(db, user_id, include_deleted=True), actor)


def user_can_user_management(
    db,
    user_id: int | None,
    action: str,
    *,
    is_admin: bool = False,
    is_customer_admin: bool = False,
) -> bool:
    if is_admin or is_customer_admin:
        return True
    if not user_id:
        return False
    action_map = {
        "deactivate": "edit",
        "activate": "edit",
        "lock": "edit",
        "unlock": "edit",
        "reset_password": "edit",
    }
    check = action_map.get(action, action)
    try:
        from user_permission_service import (
            empty_permission_actions,
            ensure_user_tab_permissions_schema,
            normalize_permission_actions,
        )

        ensure_user_tab_permissions_schema(db)
        row = db.execute(
            """
            SELECT granted, action_flags FROM user_tab_permissions
            WHERE user_id=? AND granted=1 AND endpoint='user_management'
            LIMIT 1
            """,
            (user_id,),
        ).fetchone()
        if not row:
            return False
        raw_flags = row["action_flags"] if hasattr(row, "keys") else row[1]
        actions = normalize_permission_actions(
            json.loads(raw_flags) if raw_flags else empty_permission_actions()
        )
        if check == "import":
            return bool(actions.get("import") or actions.get("create"))
        if check == "delete":
            return bool(actions.get("delete") or actions.get("edit"))
        if check == "reset_password":
            return bool(actions.get("reset_password") or actions.get("edit"))
        return bool(actions.get(check))
    except Exception:
        return False


def users_for_export(db, *, include_deleted: bool = False, **filters) -> list[dict[str, Any]]:
    listing = list_users_master(db, include_deleted=include_deleted, per_page=10000, **filters)
    rows = []
    for item in listing["items"]:
        rows.append(
            {
                "username": item.get("username"),
                "display_name": item.get("display_name"),
                "first_name": item.get("first_name"),
                "last_name": item.get("last_name"),
                "email": item.get("email"),
                "mobile": item.get("mobile"),
                "company_code": item.get("company_code"),
                "branch_code": item.get("branch_code"),
                "department_name": item.get("department_name"),
                "role": item.get("role"),
                "workflow_role": item.get("workflow_role"),
                "designation_name": item.get("designation_name"),
                "status": item.get("status"),
                "account_locked": "Yes" if item.get("account_locked") else "No",
                "last_login": item.get("last_login"),
                "password_expiry": item.get("password_expiry"),
            }
        )
    return rows


def export_users_excel(db, **filters) -> BytesIO:
    from openpyxl import Workbook

    rows = users_for_export(db, **filters)
    wb = Workbook()
    ws = wb.active
    ws.title = "Users"
    headers = list(USER_EXPORT_COLUMNS)
    ws.append(headers)
    for row in rows:
        ws.append([row.get(h, "") for h in headers])
    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def export_users_csv(db, **filters) -> str:
    rows = users_for_export(db, **filters)
    si = io.StringIO()
    writer = csv.writer(si)
    headers = list(USER_EXPORT_COLUMNS)
    writer.writerow(headers)
    for row in rows:
        writer.writerow([row.get(h, "") for h in headers])
    return si.getvalue()


def export_users_pdf(db, *, report_title: str = "User Management Report", **filters) -> BytesIO:
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.pdfgen import canvas

    rows = users_for_export(db, **filters)
    buf = BytesIO()
    page_size = landscape(A4)
    c = canvas.Canvas(buf, pagesize=page_size)
    _, height = page_size
    y = height - 40
    c.setFont("Helvetica-Bold", 14)
    c.drawString(40, y, f"MAXEK ERP — {report_title}")
    y -= 24
    c.setFont("Helvetica", 9)
    for row in rows[:250]:
        line = (
            f"{row.get('username')} | {row.get('display_name')} | {row.get('company_code')} | "
            f"{row.get('status')} | Locked: {row.get('account_locked')}"
        )
        if y < 40:
            c.showPage()
            y = height - 40
            c.setFont("Helvetica", 9)
        c.drawString(40, y, line[:130])
        y -= 14
    c.save()
    buf.seek(0)
    return buf


def user_report(db, report_key: str, **filters) -> list[dict[str, Any]]:
    key = (report_key or "list").lower().strip()
    if key == "active":
        filters["status"] = "Active"
    elif key == "inactive":
        filters["status"] = "Inactive"
    elif key == "locked":
        filters["locked"] = "1"
    elif key == "last_login":
        listing = list_users_master(db, per_page=5000, sort_by="last_login", sort_dir="desc", **filters)
        return [r for r in listing["items"] if r.get("last_login")]
    elif key == "password_expiry":
        listing = list_users_master(db, per_page=5000, **filters)
        now = datetime.now()
        out = []
        for row in listing["items"]:
            exp = row.get("password_expiry")
            if not exp:
                continue
            try:
                if datetime.strptime(str(exp), "%Y-%m-%d %H:%M:%S") <= now + timedelta(days=14):
                    out.append(row)
            except ValueError:
                continue
        return out
    listing = list_users_master(db, per_page=5000, **filters)
    return listing["items"]


def user_import_template() -> BytesIO:
    from bulk_import_service import build_xlsx_template

    return build_xlsx_template(
        [
            "Company Code",
            "Branch Code",
            "Department Code",
            "Username",
            "Password",
            "First Name",
            "Last Name",
            "Display Name",
            "Email",
            "Mobile",
            "System Role",
            "Workflow Role",
            "Status",
        ],
        [
            "CO-2026-0001",
            "BR-HO-01",
            "DEPT-HR-01",
            "jdoe",
            "TempPass1!",
            "John",
            "Doe",
            "John Doe",
            "john@example.com",
            "+91 9876543210",
            "User",
            "Maker",
            "Active",
        ],
    )


def ai_validate_user(db, user_id: int | None = None, form: dict | None = None) -> dict[str, Any]:
    data = dict(form or {})
    if user_id and not form:
        row = get_user_master(db, user_id)
        if row:
            data = dict(row)
    issues: list[str] = []
    suggestions: list[str] = []
    duplicates: list[str] = []
    missing: list[str] = []
    if not int(data.get("company_id") or 0):
        issues.append("Company is required.")
        missing.append("company_id")
    if not int(data.get("branch_id") or 0):
        issues.append("Branch is required.")
        missing.append("branch_id")
    if not int(data.get("department_id") or 0):
        issues.append("Department is required.")
        missing.append("department_id")
    username = (data.get("username") or "").strip()
    if not username:
        issues.append("Username is required.")
        missing.append("username")
    else:
        try:
            validate_user_uniqueness(
                db,
                username=username,
                email=(data.get("email") or "").strip(),
                mobile=(data.get("mobile") or "").strip(),
                user_id=user_id,
            )
        except ValueError as exc:
            duplicates.append(str(exc))
    email = (data.get("email") or "").strip()
    if email and not EMAIL_RE.match(email):
        issues.append("Email format is invalid.")
    if email:
        try:
            validate_email(email)
        except ValueError as exc:
            issues.append(str(exc))
    mobile = (data.get("mobile") or "").strip()
    if mobile:
        try:
            validate_user_mobile(mobile)
        except ValueError as exc:
            issues.append(str(exc))
    if not _display_name(data):
        suggestions.append("Add first/last name or display name.")
    if not (data.get("reporting_manager") or "").strip():
        suggestions.append("Assign a reporting manager for workflow routing.")
    result = {
        "ok": not issues and not duplicates,
        "issues": issues,
        "duplicates": duplicates,
        "suggestions": suggestions,
        "missing": missing,
    }
    try:
        from ai_service import chat_completion_json

        ai = chat_completion_json(
            "Validate ERP user records. Return JSON with keys: issues (array), suggestions (array).",
            json.dumps({"user": data, "rule_findings": result}, ensure_ascii=False),
            max_tokens=500,
        )
        for key in ("issues", "suggestions"):
            extra = ai.get(key) or []
            if isinstance(extra, list):
                result[key].extend(str(x) for x in extra if x)
        result["ok"] = not result["issues"] and not result["duplicates"]
        result["ai_enriched"] = True
    except Exception:
        result["ai_enriched"] = False
    return result


def list_companies_for_user_form(db) -> list[dict[str, Any]]:
    listing = list_companies(db, per_page=1000)
    return [
        {"id": c["id"], "company_code": c.get("company_code"), "company_name": c.get("company_name") or c.get("legal_name")}
        for c in listing.get("items", [])
    ]


def list_branches_for_user_form(db, company_id: int | None = None) -> list[dict[str, Any]]:
    from branch_master_service import list_branches_master

    listing = list_branches_master(db, company_id=company_id, status="Active", per_page=1000)
    return [
        {
            "id": b["id"],
            "branch_code": b.get("branch_code"),
            "branch_name": b.get("branch_name"),
            "company_id": b.get("company_id"),
        }
        for b in listing.get("items", [])
    ]


def list_departments_for_user_form(db, company_id: int | None = None) -> list[dict[str, Any]]:
    from department_master_service import list_departments_master

    listing = list_departments_master(db, company_id=company_id, status="Active", per_page=1000)
    return [
        {
            "id": d["id"],
            "department_code": d.get("department_code"),
            "department_name": d.get("department_name"),
            "company_id": d.get("company_id"),
        }
        for d in listing.get("items", [])
    ]


def user_profile_photo_dir(base_dir: str) -> str:
    path = os.path.join(base_dir, "static", "uploads", "users", "profiles")
    os.makedirs(path, exist_ok=True)
    return path
