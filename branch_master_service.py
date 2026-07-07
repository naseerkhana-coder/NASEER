"""Branch Master (MODULE-002) — branches linked to Company Master."""

from __future__ import annotations

import csv
import io
import json
import re
from datetime import datetime
from io import BytesIO
from typing import Any

from company_master_service import (
    COMPANY_COUNTRIES,
    COMPANY_CURRENCIES,
    COMPANY_TIMEZONES,
    EMAIL_RE,
    GSTIN_RE,
    PAN_RE,
    PHONE_RE,
    _ensure_column,
    _json_dump,
    _json_load,
    _now_ts,
    _table_exists,
    ensure_company_master_schema,
    get_company,
    list_companies,
    validate_email,
    validate_gst_number,
    validate_pan_number,
    validate_phone,
)

BRANCH_STATUSES = ("Active", "Inactive")
BRANCH_TYPES = (
    "Head Office",
    "Regional Office",
    "Site Office",
    "Site",
    "Warehouse",
    "Sales Office",
    "Other",
)
APPROVAL_STATUSES = ("Draft", "Pending", "Approved", "Rejected")
BRANCH_SORT_COLUMNS = (
    "branch_code",
    "branch_name",
    "country",
    "city",
    "status",
    "created_at",
    "company_id",
)
BRANCH_EXPORT_COLUMNS = (
    "company_code",
    "company_name",
    "branch_code",
    "branch_name",
    "branch_type",
    "gst_number",
    "pan_number",
    "country",
    "state_region",
    "district",
    "city",
    "postal_code",
    "latitude",
    "longitude",
    "phone",
    "email",
    "branch_manager",
    "opening_date",
    "closing_date",
    "working_hours",
    "currency",
    "timezone",
    "status",
    "approval_status",
)
BRANCH_AUDIT_FIELDS = (
    "branch_code",
    "branch_name",
    "branch_type",
    "company_id",
    "gst_number",
    "pan_number",
    "country",
    "state_region",
    "district",
    "city",
    "postal_code",
    "latitude",
    "longitude",
    "phone",
    "email",
    "branch_manager",
    "opening_date",
    "closing_date",
    "working_hours",
    "currency",
    "timezone",
    "status",
    "approval_status",
)
PIN_RE = re.compile(r"^\d{6}$")
BRANCH_TRANSACTION_TABLES = (
    "material_requests",
    "purchase_requests",
    "purchase_orders",
    "store_receipts",
    "store_issues",
    "payroll_records",
    "account_transactions",
    "petty_cash_requests",
    "projects",
    "staff",
    "clients",
)


def ensure_branch_master_schema(db) -> None:
    """Extend company_branches for MODULE-002 (idempotent)."""
    ensure_company_master_schema(db)
    if not _table_exists(db, "company_branches"):
        return
    for col, ctype in (
        ("branch_type", "TEXT"),
        ("gst_number", "TEXT"),
        ("pan_number", "TEXT"),
        ("district", "TEXT"),
        ("latitude", "REAL"),
        ("longitude", "REAL"),
        ("branch_manager", "TEXT"),
        ("opening_date", "TEXT"),
        ("closing_date", "TEXT"),
        ("working_hours", "TEXT"),
        ("currency", "TEXT DEFAULT 'INR'"),
        ("timezone", "TEXT DEFAULT 'Asia/Kolkata'"),
        ("modified_by", "TEXT"),
        ("approval_status", "TEXT DEFAULT 'Approved'"),
        ("approved_by", "TEXT"),
        ("approved_at", "TEXT"),
        ("is_deleted", "INTEGER DEFAULT 0"),
        ("deleted_by", "TEXT"),
        ("deleted_at", "TEXT"),
    ):
        _ensure_column(db, "company_branches", col, ctype)
    try:
        from audit_trail_service import ensure_audit_schema

        ensure_audit_schema(db)
    except Exception:
        pass


def _next_branch_code(db, company_id: int) -> str:
    prefix = f"BR-{company_id}-"
    row = db.execute(
        "SELECT branch_code FROM company_branches WHERE company_id=? AND branch_code LIKE ? "
        "ORDER BY id DESC LIMIT 1",
        (company_id, f"{prefix}%"),
    ).fetchone()
    seq = 1
    if row and row[0]:
        m = re.search(r"-(\d+)$", str(row[0]))
        if m:
            seq = int(m.group(1)) + 1
    return f"{prefix}{seq:03d}"


def validate_pin_code(value: str, *, country: str = "India") -> None:
    text = (value or "").strip()
    if not text:
        return
    if country == "India" and not PIN_RE.match(text):
        raise ValueError("Enter a valid 6-digit PIN code.")


def validate_branch_contact(*, email: str = "", phone: str = "") -> None:
    validate_email(email)
    validate_phone(phone)


def validate_branch_uniqueness(
    db,
    *,
    company_id: int,
    branch_code: str,
    gst_number: str | None = None,
    branch_id: int | None = None,
) -> None:
    code = (branch_code or "").strip().upper()
    if not code:
        raise ValueError("Branch code is required.")
    row = db.execute(
        """
        SELECT id FROM company_branches
        WHERE company_id=? AND UPPER(branch_code)=? AND COALESCE(is_deleted,0)=0
        """,
        (company_id, code),
    ).fetchone()
    if row and (not branch_id or int(row[0]) != int(branch_id)):
        raise ValueError(f"Branch code '{code}' already exists for this company.")
    gst = (gst_number or "").strip().upper()
    if gst:
        row = db.execute(
            """
            SELECT id, branch_code FROM company_branches
            WHERE UPPER(gst_number)=? AND COALESCE(is_deleted,0)=0
            """,
            (gst,),
        ).fetchone()
        if row and (not branch_id or int(row[0]) != int(branch_id)):
            raise ValueError(f"GST number '{gst}' is already registered to branch {row[1]}.")


def branch_has_transactions(db, branch_id: int) -> bool:
    for table in BRANCH_TRANSACTION_TABLES:
        if not _table_exists(db, table):
            continue
        cols = {r[1] for r in db.execute(f"PRAGMA table_info({table})").fetchall()}
        if "branch_id" not in cols:
            continue
        row = db.execute(
            f"SELECT 1 FROM {table} WHERE branch_id=? LIMIT 1",
            (branch_id,),
        ).fetchone()
        if row:
            return True
    return False


def log_branch_audit(
    db,
    branch_id: int,
    action: str,
    username: str,
    *,
    field_name: str | None = None,
    old_value: str | None = None,
    new_value: str | None = None,
    remarks: str | None = None,
) -> None:
    try:
        from audit_trail_service import log_audit_event

        log_audit_event(
            db,
            record_table="company_branches",
            record_id=branch_id,
            action=action,
            changed_by=username,
            field_name=field_name,
            old_value=old_value,
            new_value=new_value,
            remarks=remarks,
        )
    except Exception:
        pass


def log_branch_field_changes(
    db,
    before: dict[str, Any] | None,
    after: dict[str, Any] | None,
    username: str,
) -> None:
    if not before or not after:
        return
    branch_id = int(after.get("id") or before.get("id") or 0)
    if not branch_id:
        return
    for field in BRANCH_AUDIT_FIELDS:
        old_val = before.get(field)
        new_val = after.get(field)
        if str(old_val or "") != str(new_val or ""):
            log_branch_audit(
                db,
                branch_id,
                "update",
                username,
                field_name=field,
                old_value=str(old_val or ""),
                new_value=str(new_val or ""),
            )


def list_branch_audit_trail(db, branch_id: int, limit: int = 100) -> list[dict[str, Any]]:
    try:
        from audit_trail_service import list_audit_trail

        return list_audit_trail(db, "company_branches", branch_id, limit=limit)
    except Exception:
        return []


def _branch_row_to_dict(row) -> dict[str, Any]:
    item = dict(row)
    item["country_fields"] = _json_load(item.get("country_fields_json"))
    if not item.get("gst_number") and item.get("tax_registration"):
        item["gst_number"] = item["tax_registration"]
    return item


def list_branches_master(
    db,
    *,
    search: str = "",
    company_id: int | None = None,
    country: str = "",
    status: str = "",
    branch_type: str = "",
    include_deleted: bool = False,
    page: int = 1,
    per_page: int = 25,
    sort_by: str = "branch_name",
    sort_dir: str = "asc",
) -> dict[str, Any]:
    if not _table_exists(db, "company_branches"):
        return {"items": [], "total": 0, "page": 1, "per_page": per_page, "pages": 0}
    sql = (
        "SELECT b.*, c.company_code, c.company_name, c.legal_name AS company_legal_name "
        "FROM company_branches b "
        "JOIN companies c ON b.company_id = c.id "
        "WHERE COALESCE(c.is_deleted,0)=0"
    )
    count_sql = (
        "SELECT COUNT(*) FROM company_branches b "
        "JOIN companies c ON b.company_id = c.id "
        "WHERE COALESCE(c.is_deleted,0)=0"
    )
    params: list[Any] = []
    if not include_deleted:
        sql += " AND COALESCE(b.is_deleted,0)=0"
        count_sql += " AND COALESCE(b.is_deleted,0)=0"
    if company_id:
        sql += " AND b.company_id=?"
        count_sql += " AND b.company_id=?"
        params.append(company_id)
    if country:
        sql += " AND b.country=?"
        count_sql += " AND b.country=?"
        params.append(country)
    if status:
        sql += " AND b.status=?"
        count_sql += " AND b.status=?"
        params.append(status)
    if branch_type:
        sql += " AND b.branch_type=?"
        count_sql += " AND b.branch_type=?"
        params.append(branch_type)
    if search:
        clause = (
            " AND (b.branch_name LIKE ? OR b.branch_code LIKE ? OR b.city LIKE ? "
            "OR c.company_name LIKE ? OR c.company_code LIKE ? OR b.gst_number LIKE ?)"
        )
        sql += clause
        count_sql += clause
        like = f"%{search}%"
        params.extend([like, like, like, like, like, like])
    sort_col = sort_by if sort_by in BRANCH_SORT_COLUMNS else "branch_name"
    if sort_col == "company_id":
        sort_col = "b.company_id"
    else:
        sort_col = f"b.{sort_col}"
    direction = "DESC" if str(sort_dir).lower() == "desc" else "ASC"
    sql += f" ORDER BY {sort_col} {direction}, b.id DESC"
    per_page = max(1, min(int(per_page or 25), 10000))
    page = max(1, int(page or 1))
    offset = (page - 1) * per_page
    sql += " LIMIT ? OFFSET ?"
    total = int(db.execute(count_sql, params).fetchone()[0])
    rows = db.execute(sql, [*params, per_page, offset]).fetchall()
    items = [_branch_row_to_dict(r) for r in rows]
    pages = (total + per_page - 1) // per_page if total else 0
    return {"items": items, "total": total, "page": page, "per_page": per_page, "pages": pages}


def get_branch_master(db, branch_id: int, *, include_deleted: bool = False) -> dict[str, Any] | None:
    if not branch_id or not _table_exists(db, "company_branches"):
        return None
    sql = (
        "SELECT b.*, c.company_code, c.company_name, c.legal_name AS company_legal_name "
        "FROM company_branches b "
        "JOIN companies c ON b.company_id = c.id "
        "WHERE b.id=?"
    )
    if not include_deleted:
        sql += " AND COALESCE(b.is_deleted,0)=0"
    row = db.execute(sql, (branch_id,)).fetchone()
    return _branch_row_to_dict(row) if row else None


def _parse_branch_form(form) -> dict[str, Any]:
    company_id = int(form.get("company_id") or 0)
    branch_code = (form.get("branch_code") or "").strip().upper()
    branch_name = (form.get("branch_name") or "").strip()
    branch_type = (form.get("branch_type") or "").strip()
    country = (form.get("country") or form.get("branch_country") or "India").strip()
    gst_number = (
        (form.get("gst_number") or form.get("tax_registration") or "").strip().upper()
    )
    pan_number = (form.get("pan_number") or "").strip().upper()
    email = (form.get("email") or form.get("branch_email") or "").strip()
    phone = (form.get("phone") or form.get("branch_phone") or "").strip()
    postal_code = (form.get("postal_code") or form.get("branch_postal_code") or "").strip()
    lat_raw = (form.get("latitude") or "").strip()
    lon_raw = (form.get("longitude") or "").strip()
    latitude = float(lat_raw) if lat_raw else None
    longitude = float(lon_raw) if lon_raw else None
    currency = (form.get("currency") or "INR").strip()
    timezone = (form.get("timezone") or "Asia/Kolkata").strip()
    status = (form.get("status") or form.get("branch_status") or "Active").strip()
    is_ho = 1 if form.get("is_head_office") == "on" else 0
    return {
        "company_id": company_id,
        "branch_code": branch_code,
        "branch_name": branch_name,
        "branch_type": branch_type,
        "country": country,
        "address_line1": (form.get("address_line1") or form.get("branch_address_line1") or "").strip(),
        "address_line2": (form.get("address_line2") or form.get("branch_address_line2") or "").strip(),
        "city": (form.get("city") or form.get("branch_city") or "").strip(),
        "state_region": (form.get("state_region") or form.get("branch_state_region") or "").strip(),
        "district": (form.get("district") or "").strip(),
        "postal_code": postal_code,
        "latitude": latitude,
        "longitude": longitude,
        "phone": phone,
        "email": email,
        "gst_number": gst_number,
        "pan_number": pan_number,
        "tax_registration": gst_number,
        "branch_manager": (form.get("branch_manager") or "").strip(),
        "opening_date": (form.get("opening_date") or "").strip(),
        "closing_date": (form.get("closing_date") or "").strip(),
        "working_hours": (form.get("working_hours") or "").strip(),
        "currency": currency,
        "timezone": timezone,
        "status": status,
        "is_head_office": is_ho,
        "country_fields_json": _json_dump({}),
    }


def save_branch_master(
    db,
    form,
    username: str,
    branch_id: int | None = None,
    *,
    customer_id: int | None = None,
) -> int:
    data = _parse_branch_form(form)
    company_id = data["company_id"]
    if not company_id:
        raise ValueError("Company is required.")
    if not get_company(db, company_id):
        raise ValueError("Selected company was not found.")
    if not data["branch_code"]:
        data["branch_code"] = _next_branch_code(db, company_id)
    if not data["branch_name"]:
        raise ValueError("Branch name is required.")
    if data["branch_type"] and data["branch_type"] not in BRANCH_TYPES:
        raise ValueError("Select a valid branch type.")
    if data["country"] not in COMPANY_COUNTRIES:
        raise ValueError("Select a valid country.")
    if data["status"] not in BRANCH_STATUSES:
        raise ValueError("Select a valid status.")
    validate_branch_contact(email=data["email"], phone=data["phone"])
    validate_pin_code(data["postal_code"], country=data["country"])
    if data["gst_number"]:
        validate_gst_number(data["gst_number"])
    if data["pan_number"]:
        validate_pan_number(data["pan_number"])
    if data["currency"] not in COMPANY_CURRENCIES:
        raise ValueError("Select a valid currency.")
    if data["timezone"] not in COMPANY_TIMEZONES:
        raise ValueError("Select a valid timezone.")
    validate_branch_uniqueness(
        db,
        company_id=company_id,
        branch_code=data["branch_code"],
        gst_number=data["gst_number"],
        branch_id=branch_id,
    )
    if not branch_id and customer_id:
        from super_admin_service import assert_branch_limit_not_exceeded

        assert_branch_limit_not_exceeded(db, customer_id)
    now = _now_ts()
    core = (
        company_id,
        data["branch_code"],
        data["branch_name"],
        data["branch_type"],
        data["country"],
        data["address_line1"],
        data["address_line2"],
        data["city"],
        data["state_region"],
        data["district"],
        data["postal_code"],
        data["latitude"],
        data["longitude"],
        data["phone"],
        data["email"],
        data["gst_number"],
        data["pan_number"],
        data["tax_registration"],
        data["country_fields_json"],
        data["branch_manager"],
        data["opening_date"],
        data["closing_date"],
        data["working_hours"],
        data["currency"],
        data["timezone"],
        data["is_head_office"],
        data["status"],
    )
    if branch_id:
        existing = get_branch_master(db, branch_id, include_deleted=True)
        if not existing:
            raise ValueError("Branch not found.")
        db.execute(
            """
            UPDATE company_branches SET company_id=?, branch_code=?, branch_name=?, branch_type=?,
            country=?, address_line1=?, address_line2=?, city=?, state_region=?, district=?,
            postal_code=?, latitude=?, longitude=?, phone=?, email=?, gst_number=?, pan_number=?,
            tax_registration=?, country_fields_json=?, branch_manager=?, opening_date=?,
            closing_date=?, working_hours=?, currency=?, timezone=?, is_head_office=?, status=?,
            modified_by=?, modified_at=? WHERE id=?
            """,
            (*core, username, now, branch_id),
        )
        if customer_id is not None:
            db.execute(
                "UPDATE company_branches SET customer_id=? WHERE id=?",
                (customer_id, branch_id),
            )
        log_branch_field_changes(db, existing, get_branch_master(db, branch_id, include_deleted=True), username)
        return branch_id
    approval_status = (form.get("approval_status") or "Draft").strip()
    if approval_status not in APPROVAL_STATUSES:
        approval_status = "Draft"
    insert_cols = (
        "company_id, branch_code, branch_name, branch_type, country, address_line1, address_line2, "
        "city, state_region, district, postal_code, latitude, longitude, phone, email, gst_number, "
        "pan_number, tax_registration, country_fields_json, branch_manager, opening_date, closing_date, "
        "working_hours, currency, timezone, is_head_office, status, approval_status, created_by, "
        "created_at, modified_by, modified_at"
    )
    placeholders = ",".join(["?"] * 32)
    vals = (*core, approval_status, username, now, username, now)
    if customer_id is not None:
        cur = db.execute(
            f"INSERT INTO company_branches({insert_cols}, customer_id) VALUES({placeholders},?)",
            (*vals, customer_id),
        )
    else:
        cur = db.execute(
            f"INSERT INTO company_branches({insert_cols}) VALUES({placeholders})",
            vals,
        )
    new_id = int(cur.lastrowid)
    if customer_id is not None:
        from super_admin_service import sync_customer_usage_counts

        sync_customer_usage_counts(db, customer_id)
    log_branch_audit(db, new_id, "create", username, remarks=f"Created branch {data['branch_code']}")
    return new_id


def soft_delete_branch_master(db, branch_id: int, username: str) -> None:
    if not branch_id:
        raise ValueError("Invalid branch.")
    row = get_branch_master(db, branch_id, include_deleted=True)
    if not row:
        raise ValueError("Branch not found.")
    if row.get("is_deleted"):
        return
    if branch_has_transactions(db, branch_id):
        raise ValueError("Branch cannot be deleted because transactions exist. Deactivate instead.")
    now = _now_ts()
    db.execute(
        """
        UPDATE company_branches SET is_deleted=1, deleted_by=?, deleted_at=?,
        modified_by=?, modified_at=? WHERE id=?
        """,
        (username, now, username, now, branch_id),
    )
    log_branch_audit(
        db,
        branch_id,
        "soft_delete",
        username,
        remarks=f"Soft-deleted branch {row.get('branch_code')}",
    )


def activate_branch_master(db, branch_id: int, username: str) -> None:
    row = get_branch_master(db, branch_id)
    if not row:
        raise ValueError("Branch not found.")
    now = _now_ts()
    db.execute(
        "UPDATE company_branches SET status='Active', modified_by=?, modified_at=? WHERE id=?",
        (username, now, branch_id),
    )
    log_branch_audit(db, branch_id, "activate", username, field_name="status", old_value="Inactive", new_value="Active")


def deactivate_branch_master(db, branch_id: int, username: str) -> None:
    row = get_branch_master(db, branch_id)
    if not row:
        raise ValueError("Branch not found.")
    now = _now_ts()
    db.execute(
        "UPDATE company_branches SET status='Inactive', modified_by=?, modified_at=? WHERE id=?",
        (username, now, branch_id),
    )
    log_branch_audit(db, branch_id, "deactivate", username, field_name="status", old_value="Active", new_value="Inactive")


def approve_branch_master(db, branch_id: int, username: str) -> None:
    row = get_branch_master(db, branch_id)
    if not row:
        raise ValueError("Branch not found.")
    now = _now_ts()
    db.execute(
        """
        UPDATE company_branches SET approval_status='Approved', approved_by=?, approved_at=?,
        modified_by=?, modified_at=? WHERE id=?
        """,
        (username, now, username, now, branch_id),
    )
    log_branch_audit(db, branch_id, "approve", username, remarks="Branch approved")


def reject_branch_master(db, branch_id: int, username: str, remarks: str = "") -> None:
    row = get_branch_master(db, branch_id)
    if not row:
        raise ValueError("Branch not found.")
    now = _now_ts()
    db.execute(
        "UPDATE company_branches SET approval_status='Rejected', modified_by=?, modified_at=? WHERE id=?",
        (username, now, branch_id),
    )
    log_branch_audit(db, branch_id, "reject", username, remarks=remarks or "Branch rejected")


def branch_allows_transactions(db, branch_id: int) -> bool:
    row = get_branch_master(db, branch_id)
    if not row:
        return False
    if row.get("is_deleted"):
        return False
    if (row.get("status") or "").strip() != "Active":
        return False
    if (row.get("approval_status") or "Approved") not in ("Approved",):
        return False
    return True


def user_can_branch_master(
    db,
    user_id: int | None,
    action: str,
    *,
    is_admin: bool = False,
) -> bool:
    if is_admin:
        return True
    if not user_id:
        return False
    action_map = {
        "deactivate": "edit",
        "activate": "edit",
    }
    check = action_map.get(action, action)
    try:
        from user_permission_service import (
            empty_permission_actions,
            ensure_user_tab_permissions_schema,
            normalize_permission_actions,
        )

        ensure_user_tab_permissions_schema(db)
        row = db.execute(
            """
            SELECT granted, action_flags FROM user_tab_permissions
            WHERE user_id=? AND granted=1 AND endpoint='branch_master'
            LIMIT 1
            """,
            (user_id,),
        ).fetchone()
        if not row:
            return False
        raw_flags = row["action_flags"] if hasattr(row, "keys") else row[1]
        actions = normalize_permission_actions(
            json.loads(raw_flags) if raw_flags else empty_permission_actions()
        )
        if check == "import":
            return bool(actions.get("import") or actions.get("create"))
        if check == "delete":
            return bool(actions.get("delete") or actions.get("edit"))
        return bool(actions.get(check))
    except Exception:
        return False


def branches_for_export(db, *, include_deleted: bool = False, **filters) -> list[dict[str, Any]]:
    listing = list_branches_master(db, include_deleted=include_deleted, per_page=10000, **filters)
    return [{col: item.get(col, "") for col in BRANCH_EXPORT_COLUMNS} for item in listing["items"]]


def export_branches_excel(db, **filters) -> BytesIO:
    from openpyxl import Workbook

    rows = branches_for_export(db, **filters)
    wb = Workbook()
    ws = wb.active
    ws.title = "Branches"
    headers = list(BRANCH_EXPORT_COLUMNS)
    ws.append(headers)
    for row in rows:
        ws.append([row.get(h, "") for h in headers])
    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def export_branches_csv(db, **filters) -> str:
    rows = branches_for_export(db, **filters)
    si = io.StringIO()
    writer = csv.writer(si)
    headers = list(BRANCH_EXPORT_COLUMNS)
    writer.writerow(headers)
    for row in rows:
        writer.writerow([row.get(h, "") for h in headers])
    return si.getvalue()


def export_branches_pdf(db, *, report_title: str = "Branch Master Report", **filters) -> BytesIO:
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.pdfgen import canvas

    rows = branches_for_export(db, **filters)
    buf = BytesIO()
    page_size = landscape(A4)
    c = canvas.Canvas(buf, pagesize=page_size)
    _, height = page_size
    y = height - 40
    c.setFont("Helvetica-Bold", 14)
    c.drawString(40, y, f"MAXEK ERP — {report_title}")
    y -= 24
    c.setFont("Helvetica", 9)
    for row in rows[:250]:
        line = (
            f"{row.get('company_code')} | {row.get('branch_code')} | {row.get('branch_name')} | "
            f"{row.get('city')} | {row.get('status')}"
        )
        if y < 40:
            c.showPage()
            y = height - 40
            c.setFont("Helvetica", 9)
        c.drawString(40, y, line[:130])
        y -= 14
    c.save()
    buf.seek(0)
    return buf


def branch_report(db, report_key: str, **filters) -> list[dict[str, Any]] | dict[str, Any]:
    key = (report_key or "master").lower().strip()
    if key == "active":
        filters["status"] = "Active"
    elif key == "inactive":
        filters["status"] = "Inactive"
    elif key == "directory":
        filters["status"] = "Active"
    elif key in ("company_wise", "by_company", "branch_company"):
        listing = list_branches_master(db, per_page=5000, **filters)
        groups: dict[int, dict[str, Any]] = {}
        for item in listing["items"]:
            cid = int(item.get("company_id") or 0)
            if cid not in groups:
                groups[cid] = {
                    "company_id": cid,
                    "company_code": item.get("company_code"),
                    "company_name": item.get("company_name"),
                    "company_legal_name": item.get("company_legal_name"),
                    "branches": [],
                }
            groups[cid]["branches"].append(
                {
                    "id": item.get("id"),
                    "branch_code": item.get("branch_code"),
                    "branch_name": item.get("branch_name"),
                    "branch_type": item.get("branch_type"),
                    "city": item.get("city"),
                    "country": item.get("country"),
                    "status": item.get("status"),
                    "approval_status": item.get("approval_status"),
                }
            )
        companies = list(groups.values())
        for group in companies:
            group["branch_count"] = len(group["branches"])
        return {
            "report": "company_wise",
            "company_count": len(companies),
            "branch_count": sum(g["branch_count"] for g in companies),
            "companies": companies,
        }
    listing = list_branches_master(db, per_page=5000, **filters)
    return listing["items"]


def branch_import_template() -> BytesIO:
    from bulk_import_service import build_xlsx_template

    return build_xlsx_template(
        [
            "Company Code",
            "Branch Code",
            "Branch Name",
            "Branch Type",
            "GST Number",
            "PAN Number",
            "Country",
            "State",
            "District",
            "City",
            "PIN Code",
            "Latitude",
            "Longitude",
            "Phone",
            "Email",
            "Branch Manager",
            "Opening Date",
            "Closing Date",
            "Working Hours",
            "Currency",
            "Timezone",
            "Status",
        ],
        [
            "CO-2026-0001",
            "BR-HO-01",
            "Head Office Mumbai",
            "Head Office",
            "",
            "",
            "India",
            "Maharashtra",
            "Mumbai",
            "Mumbai",
            "400001",
            "",
            "",
            "+91 9876543210",
            "branch@example.com",
            "Manager Name",
            "2026-01-01",
            "",
            "09:00-18:00",
            "INR",
            "Asia/Kolkata",
            "Active",
        ],
    )


def ai_validate_branch(db, branch_id: int | None = None, form: dict | None = None) -> dict[str, Any]:
    """Rule-based validation + optional OpenAI enrichment."""
    data = form or {}
    if branch_id and not form:
        row = get_branch_master(db, branch_id)
        if row:
            data = dict(row)
    issues: list[str] = []
    suggestions: list[str] = []
    duplicates: list[str] = []
    missing: list[str] = []
    address_issues: list[str] = []
    company_id = int(data.get("company_id") or 0)
    branch_code = (data.get("branch_code") or "").strip()
    if not company_id:
        issues.append("Company is not selected.")
        missing.append("company_id")
    if not branch_code:
        issues.append("Branch code is missing.")
        missing.append("branch_code")
    elif company_id:
        try:
            validate_branch_uniqueness(
                db,
                company_id=company_id,
                branch_code=branch_code,
                gst_number=data.get("gst_number"),
                branch_id=branch_id,
            )
        except ValueError as exc:
            duplicates.append(str(exc))
    if not (data.get("branch_name") or "").strip():
        issues.append("Branch name is missing.")
        missing.append("branch_name")
    if not (data.get("address_line1") or "").strip():
        suggestions.append("Add address line 1 for complete branch directory.")
        missing.append("address_line1")
    if not (data.get("phone") or "").strip():
        suggestions.append("Add a contact phone number.")
    if not (data.get("email") or "").strip():
        suggestions.append("Add a branch email address.")
    if data.get("country") == "India" and not (data.get("postal_code") or "").strip():
        suggestions.append("PIN code is recommended for Indian branches.")
    if data.get("latitude") and data.get("longitude"):
        try:
            lat = float(data["latitude"])
            lon = float(data["longitude"])
            if not (-90 <= lat <= 90 and -180 <= lon <= 180):
                issues.append("Latitude/longitude values are out of valid range.")
        except (TypeError, ValueError):
            issues.append("Latitude/longitude must be numeric.")
    country = (data.get("country") or "").strip()
    city = (data.get("city") or "").strip()
    state = (data.get("state_region") or "").strip()
    address_consistent = bool(country and city)
    if country and not city:
        address_issues.append("City is missing for the selected country.")
        address_consistent = False
    if city and not country:
        address_issues.append("Country is missing while city is provided.")
        address_consistent = False
    if country == "India" and city and state and city.lower() not in state.lower() and state.lower() not in city.lower():
        address_issues.append("City and state may be inconsistent — verify the address.")
    result = {
        "ok": not issues and not duplicates and not address_issues,
        "issues": issues,
        "duplicates": duplicates,
        "suggestions": suggestions,
        "missing": missing,
        "address_issues": address_issues,
        "address_consistent": address_consistent,
    }
    try:
        from ai_service import OpenAIConfigurationError, chat_completion_json

        prompt = json.dumps(
            {
                "branch": {
                    k: data.get(k)
                    for k in (
                        "branch_code",
                        "branch_name",
                        "country",
                        "state_region",
                        "city",
                        "postal_code",
                        "address_line1",
                    )
                },
                "rule_findings": result,
            },
            ensure_ascii=False,
        )
        ai = chat_completion_json(
            "You validate ERP branch master records. Return JSON with keys: issues (array), suggestions (array), address_consistent (boolean).",
            prompt,
            max_tokens=600,
        )
        for key in ("issues", "suggestions"):
            extra = ai.get(key) or []
            if isinstance(extra, list):
                result[key].extend(str(x) for x in extra if x)
        if "address_consistent" in ai:
            result["address_consistent"] = bool(ai["address_consistent"])
        result["ok"] = not result["issues"] and not result["duplicates"] and not result["address_issues"]
        result["ai_enriched"] = True
    except Exception:
        result["ai_enriched"] = False
    return result


def list_companies_for_branch_form(db) -> list[dict[str, Any]]:
    listing = list_companies(db, per_page=1000)
    return [
        {
            "id": c["id"],
            "company_code": c.get("company_code"),
            "company_name": c.get("company_name") or c.get("legal_name"),
        }
        for c in listing.get("items", [])
    ]
